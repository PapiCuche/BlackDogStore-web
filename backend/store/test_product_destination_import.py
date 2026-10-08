"""
PRODUCT-DESTINATION-01 · la carga masiva dice a dónde va cada producto.

Un producto puede estar en el catálogo de la empresa y NO venderse por la web:
un repuesto, un equipo que sólo se despacha en tienda. Eso no es «inactivo»
—se vende en caja, se cuenta en inventario, se usa en el taller—, es otro dato:
`Product.is_published_online`.

Lo que estas pruebas fijan del importador y de su plantilla:

  * la columna «Destino» —y «Destino (pendiente web)», como la escribió quien
    preparó la plantilla— se reconoce sola, se valida y se ve en la vista previa;
  * vacía no cambia nada: un producto nuevo se publica, uno existente se queda
    como estaba. Un archivo de antes de esta columna se importa igual;
  * la plantilla trae las categorías DE ESA EMPRESA y sólo las suyas, y sus
    hojas de ayuda nunca se importan;
  * una categoría se valida contra las de la empresa. Ninguna está escrita en
    el código.

Quién ve un producto «sólo stock interno» está en `test_product_visibility.py`.
"""
import io
import json

import openpyxl
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient

from store import import_exports
from store.models import BulkImportRow, Category, Product
from store.tests import _p2d_member, _p3_company

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PREVIEW = '/api/admin/products/import/preview/'
TEMPLATE = '/api/admin/products/import/template/'
INSPECT = '/api/admin/imports/inspect/'
ONLINE, INTERNAL = 'Publicar en e-commerce', 'Solo stock interno'
HEADERS = ['Código de barras', 'Código', 'Nombre', 'Descripción', 'Precio de venta', 'Categoría',
           'Imagen principal', 'Imágenes']


def workbook_bytes(rows, *, destination_header='Destino', headers=None, extra_sheets=()):
    """A product file as somebody would fill it: header row, then `rows`."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Productos'
    sheet.append(list(headers) if headers is not None else [*HEADERS, destination_header])
    for row in rows:
        sheet.append(list(row))
    for name, lines in extra_sheets:
        other = workbook.create_sheet(name)
        for line in lines:
            other.append(list(line))
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def product_row(code, name, *, price=100, category=None, destination=None):
    return [None, code, name, None, price, category, None, None, destination]


class _Base(TestCase):
    def setUp(self):
        cache.clear()
        self.company = _p3_company('destino-tienda', 'Tienda Destino')
        self.manager, _ = _p2d_member(
            self.company, 'destino_gestor', ['company.view', 'products.view', 'products.manage'])
        self.client = APIClient()
        self.client.force_authenticate(user=self.manager)
        self.phones = Category.objects.create(company=self.company, name='Teléfonos', slug='telefonos')
        self.parts = Category.objects.create(company=self.company, name='Repuestos', slug='repuestos')

    def preview(self, data, *, company=None, client=None, options=None, **fields):
        return (client or self.client).post(
            f'{PREVIEW}?company={(company or self.company).pk}',
            {'file': SimpleUploadedFile('productos.xlsx', data, content_type=XLSX),
             'options': json.dumps(options or {}), **fields},
            format='multipart',
        )

    def apply(self, job_id, *, company=None, client=None):
        return (client or self.client).post(
            f'/api/admin/products/import/{job_id}/apply/?company={(company or self.company).pk}')

    def imported(self, rows, **kwargs):
        body = self.preview(workbook_bytes(rows, **kwargs)).json()
        self.assertEqual(body['counts']['error'], 0, body['rows'])
        res = self.apply(body['id'])
        self.assertEqual(res.status_code, 200, res.content)
        return body

    def product(self, name):
        return Product.objects.get(company=self.company, name=name)


class DestinationColumnTest(_Base):
    def test_each_row_goes_where_its_cell_says(self):
        body = self.imported([
            product_row('WEB-1', 'Funda de silicona', destination=ONLINE),
            product_row('INT-1', 'Flex de carga', destination=INTERNAL),
        ])

        shown = {row['data']['name']: row['data']['destination'] for row in body['rows'] if row['action'] == 'create'}
        self.assertEqual(shown, {'Funda de silicona': 'online', 'Flex de carga': 'internal'})
        self.assertTrue(self.product('Funda de silicona').is_published_online)
        self.assertFalse(self.product('Flex de carga').is_published_online)
        # Otro dato que «activo»: lo de sólo stock interno sigue siendo un producto en uso.
        self.assertTrue(self.product('Flex de carga').is_active)

    def test_the_header_is_recognised_as_the_template_wrote_it_and_as_it_is_now(self):
        for header in ('Destino', 'Destino (pendiente web)', 'DESTINO', ' destino  (Pendiente Web) '):
            with self.subTest(header):
                body = self.preview(workbook_bytes(
                    [product_row(f'CODH-{len(header)}', f'Producto {header.strip()}', destination=INTERNAL)],
                    destination_header=header)).json()
                self.assertEqual(body['counts'], {**body['counts'], 'create': 1, 'error': 0}, body['rows'])
                self.assertEqual(body['rows'][0]['data']['destination'], 'internal')

    def test_it_is_read_as_people_write_it(self):
        for position, (written, meaning) in enumerate((
            ('Publicar en e-commerce', 'online'), ('publicar en e-commerce', 'online'), (' PUBLICAR EN E-COMMERCE ', 'online'),
            ('Publicar en ecommerce', 'online'), ('E-commerce', 'online'),
            ('Solo stock interno', 'internal'), ('Sólo stock interno', 'internal'), ('SOLO STOCK INTERNO', 'internal'),
            ('  solo   stock  interno ', 'internal'), ('Stock interno', 'internal'),
        )):
            with self.subTest(written):
                body = self.preview(workbook_bytes(
                    [product_row(f'CODW-{position}', f'Producto {position}', destination=written)])).json()
                self.assertEqual(body['counts']['error'], 0, body['rows'])
                self.assertEqual(body['rows'][0]['data']['destination'], meaning)

    def test_a_value_that_is_neither_is_an_error_that_names_both(self):
        for written in ('Quizás', 'Sí', 'No', '1', 'interno o web'):
            with self.subTest(written):
                res = self.preview(workbook_bytes([product_row('CODX-1', 'Producto dudoso', destination=written)]))
                body = res.json()
                self.assertEqual(body['counts']['error'], 1, body['rows'])
                message = ' '.join(body['rows'][0]['errors'])
                self.assertIn(ONLINE, message)
                self.assertIn(INTERNAL, message)
                self.assertIn(written, message)
                self.assertIn(self.apply(body['id']).status_code, (400, 409))        # nothing with an error is applied
        self.assertFalse(Product.objects.filter(company=self.company).exists())

    def test_an_empty_cell_publishes_what_is_new_and_leaves_alone_what_exists(self):
        self.imported([product_row('CODE-1', 'Cargador', destination=INTERNAL)])
        self.assertFalse(self.product('Cargador').is_published_online)

        body = self.imported([
            product_row('CODE-1', 'Cargador', price=120),          # existe, celda vacía
            product_row('CODE-2', 'Cable'),                        # nuevo, celda vacía
        ])

        by_name = {row['data']['name']: row for row in body['rows']}
        self.assertEqual(by_name['Cargador']['action'], 'update')
        self.assertEqual(by_name['Cargador']['data']['destination'], '')          # no dice nada: no cambia nada
        self.assertEqual(by_name['Cable']['data']['destination'], 'online')       # lo que va a pasar, a la vista
        self.assertFalse(self.product('Cargador').is_published_online)
        self.assertEqual(self.product('Cargador').price, 120)
        self.assertTrue(self.product('Cable').is_published_online)

    def test_an_explicit_cell_changes_an_existing_product_and_says_so(self):
        self.imported([product_row('CODC-1', 'Batería', destination=ONLINE)])

        body = self.imported([product_row('CODC-1', 'Batería', destination=INTERNAL)])
        self.assertEqual(body['rows'][0]['action'], 'update')
        self.assertFalse(self.product('Batería').is_published_online)

        body = self.imported([product_row('CODC-1', 'Batería', destination=ONLINE)])
        self.assertTrue(self.product('Batería').is_published_online)

    def test_a_file_from_before_the_column_imports_as_it_did(self):
        Product.objects.create(company=self.company, name='Pantalla', slug='pantalla', price=50, is_published_online=False)
        old_file = workbook_bytes(
            [[None, 'CODV-1', 'Pantalla', None, 60, None, None, None], [None, 'CODV-2', 'Mica', None, 10, None, None, None]],
            headers=HEADERS)

        body = self.preview(old_file).json()
        self.assertEqual(body['counts']['error'], 0, body['rows'])
        self.assertEqual(self.apply(body['id']).status_code, 200)

        self.assertFalse(self.product('Pantalla').is_published_online)             # nadie lo pidió: no se toca
        self.assertEqual(self.product('Pantalla').price, 60)
        self.assertTrue(self.product('Mica').is_published_online)

    def test_a_preview_made_before_this_version_still_applies(self):
        body = self.preview(workbook_bytes([product_row('CODP-1', 'Soporte', destination=INTERNAL)])).json()
        for row in BulkImportRow.objects.filter(job_id=body['id']):
            row.normalized_data.pop('destination', None)
            row.save(update_fields=['normalized_data'])

        self.assertEqual(self.apply(body['id']).status_code, 200)
        self.assertTrue(self.product('Soporte').is_published_online)

    def test_the_column_can_be_assigned_by_hand(self):
        data = workbook_bytes([product_row('CODM-1', 'Lápiz', destination=INTERNAL)], destination_header='A dónde va')
        body = self.preview(data, header_row='1', mapping=json.dumps({'code': 1, 'name': 2, 'price': 4, 'destination': 8})).json()
        self.assertEqual(body['counts']['error'], 0, body)
        self.assertEqual(body['rows'][0]['data']['destination'], 'internal')


class TheColumnIsNeverSilentlyIgnoredTest(_Base):
    """
    A cell that says «Solo stock interno» and is not read publishes the product.

    So the column is found by its header whatever mapping arrives: one the
    screen remembered from before this column existed, one of a format that
    never had it, one somebody assigned by hand and forgot it in.
    """

    def test_a_mapping_remembered_from_before_the_column_still_reads_it(self):
        for header in ('Destino', 'Destino (pendiente web)'):
            with self.subTest(header):
                data = workbook_bytes([product_row(f'CODR-{len(header)}', f'Flex {header}', destination=INTERNAL)],
                                      destination_header=header)
                # What the wizard sends when it prefers its stored profile: no «destination».
                old_profile = {'code': 1, 'name': 2, 'price': 4, 'category': 5}
                body = self.preview(data, header_row='1', mapping=json.dumps(old_profile)).json()

                self.assertEqual(body['counts']['error'], 0, body['rows'])
                self.assertEqual(body['rows'][0]['data']['destination'], 'internal')
                self.assertEqual(self.apply(body['id']).status_code, 200)
                self.assertFalse(self.product(f'Flex {header}').is_published_online)

    def test_a_header_that_starts_like_it_is_read_too(self):
        for position, header in enumerate(('Destino web', 'Destino del producto', 'Destino (e-commerce)', 'destino:')):
            with self.subTest(header):
                body = self.preview(workbook_bytes(
                    [product_row(f'CODS-{position}', f'Repuesto {position}', destination=INTERNAL)],
                    destination_header=header)).json()
                self.assertEqual(body['counts']['error'], 0, body['rows'])
                self.assertEqual(body['rows'][0]['data']['destination'], 'internal', header)

    def test_the_preview_says_which_column_it_took_it_from(self):
        data = workbook_bytes([product_row('CODN-1', 'Repuesto', destination=INTERNAL)], destination_header='Destino web')
        body = self.preview(data, header_row='1', mapping=json.dumps({'code': 1, 'name': 2, 'price': 4})).json()
        self.assertIn('Destino web', ' '.join(body['summary']['format_notes']))

    def test_a_column_somebody_assigned_by_hand_is_the_one_used(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Productos'
        sheet.append([*HEADERS, 'Destino', 'A dónde va de verdad'])
        sheet.append([None, 'CODM-9', 'Lápiz óptico', None, 100, None, None, None, ONLINE, INTERNAL])
        out = io.BytesIO()
        workbook.save(out)
        body = self.preview(out.getvalue(), header_row='1',
                            mapping=json.dumps({'code': 1, 'name': 2, 'price': 4, 'destination': 9})).json()
        self.assertEqual(body['rows'][0]['data']['destination'], 'internal')

    def test_a_help_row_is_only_the_row_under_the_headers(self):
        """A product that really is called «Obligatorio» further down is not swallowed."""
        body = self.preview(workbook_bytes([
            product_row('CODO-1', 'Cable'),
            product_row('CODO-2', 'Obligatorio'),
        ])).json()
        self.assertEqual(body['counts']['create'], 2, body['rows'])


class OwnersWorkbookTest(_Base):
    """El libro tal como lo preparó el propietario: ayuda en mayúsculas, hoja «Categorías», columna I."""

    def book(self, *rows):
        return workbook_bytes(
            [
                ['EAN / UPC; como TEXTO', 'Código interno; como TEXTO', 'OBLIGATORIO', 'Opcional', 'OBLIGATORIO para nuevos',
                 'Selecciona una categoría de la hoja Categorías', 'Opcional: nombre exacto de archivo',
                 'Opcional: archivo1.webp|archivo2.webp', 'CONTROL EN EXCEL: hoy NO se importa'],
                ['195949806575', 'IP16PM-256-W', 'iPhone 16 Pro Max 256GB White', 'Pantalla de 6,9 pulgadas, 256 GB.', 4999,
                 'Teléfonos', None, None, ONLINE],
                *rows,
            ],
            destination_header='Destino (pendiente web)',
            extra_sheets=(
                ('Categorías', [['CATEGORÍAS DE TU EMPRESA', 'INSTRUCCIONES'], ['Nombre exacto de la categoría', 'Copia aquí…'], ['Teléfonos']]),
                ('Instrucciones', [['PLANTILLA DE PRODUCTOS'], ['Estado de funciones', 'Cómo funciona']]),
                ('Ejemplo', [HEADERS, ['195949806575', 'IP16PM-256-W', 'iPhone 16 Pro Max 256GB White', 'x', 4999, 'Teléfonos']]),
            ),
        )

    def test_as_it_is_it_creates_nothing_and_complains_about_nothing(self):
        body = self.preview(self.book()).json()

        self.assertEqual(body['counts']['create'], 0, body['rows'])
        self.assertEqual(body['counts']['error'], 0, body['rows'])
        self.assertFalse(Product.objects.filter(company=self.company).exists())

    def test_its_real_rows_are_imported_with_their_destination(self):
        body = self.imported_book(
            ['', 'FLEX-IP13', 'Flex de carga iPhone 13', None, 35, 'Repuestos', None, None, INTERNAL],
            ['', 'FUNDA-IP16', 'Funda iPhone 16', None, 59, 'Teléfonos', None, None, ONLINE],
        )

        self.assertEqual(body['counts']['create'], 2, body['rows'])
        self.assertFalse(self.product('Flex de carga iPhone 13').is_published_online)
        self.assertEqual(self.product('Flex de carga iPhone 13').category, self.parts)
        self.assertTrue(self.product('Funda iPhone 16').is_published_online)

    def imported_book(self, *rows):
        body = self.preview(self.book(*rows)).json()
        self.assertEqual(body['counts']['error'], 0, body['rows'])
        self.assertEqual(self.apply(body['id']).status_code, 200)
        return body

    def test_its_categories_sheet_is_help_and_never_products(self):
        res = self.client.post(
            f'{INSPECT}?company={self.company.pk}',
            {'file': SimpleUploadedFile('p.xlsx', self.book(), content_type=XLSX)}, format='multipart')
        self.assertEqual(res.status_code, 200, res.content)
        sheets = {sheet['name']: sheet for sheet in res.json()['sheets']}
        self.assertTrue(sheets['Categorías']['help'])
        self.assertFalse(sheets['Productos']['help'])

        chosen = self.preview(self.book(), sheet_name='Categorías')
        self.assertEqual(chosen.status_code, 400, chosen.content)
        self.assertFalse(Product.objects.filter(company=self.company).exists())


class CategoriesTest(_Base):
    def test_a_category_is_checked_against_this_companys_own(self):
        other = _p3_company('destino-otra', 'Otra Tienda')
        Category.objects.create(company=other, name='Sólo de la otra', slug='solo-de-la-otra')

        body = self.preview(workbook_bytes([
            product_row('CODK-1', 'Con la mía', category='Teléfonos'),
            product_row('CODK-2', 'Con la ajena', category='Sólo de la otra'),
        ])).json()

        by_name = {row['data'].get('name'): row for row in body['rows']}
        self.assertEqual(by_name['Con la mía']['action'], 'create')
        self.assertEqual(by_name['Con la mía']['data']['category_id'], self.phones.pk)
        self.assertEqual(by_name['Con la ajena']['action'], 'error')
        self.assertIn('no existe', ' '.join(by_name['Con la ajena']['errors']))

    def test_a_category_can_be_chosen_for_the_rows_that_name_none(self):
        body = self.preview(
            workbook_bytes([product_row('CODD-1', 'Sin categoría'), product_row('CODD-2', 'Con categoría', category='Teléfonos')]),
            options={'default_category_id': self.parts.pk}).json()
        self.assertEqual(body['counts']['error'], 0, body)
        self.assertEqual(self.apply(body['id']).status_code, 200)

        self.assertEqual(self.product('Sin categoría').category, self.parts)
        self.assertEqual(self.product('Con categoría').category, self.phones)       # la de la fila manda

    def test_the_chosen_category_has_to_be_this_companys(self):
        other = _p3_company('destino-otra', 'Otra Tienda')
        foreign = Category.objects.create(company=other, name='Ajena', slug='ajena')

        for bad in (foreign.pk, 999999, 'abc', -1):
            with self.subTest(bad):
                res = self.preview(workbook_bytes([product_row('CODF-1', 'Producto')]), options={'default_category_id': bad})
                self.assertEqual(res.status_code, 400, res.content)
                self.assertNotIn('Ajena', res.content.decode())
        self.assertFalse(Product.objects.filter(company=self.company).exists())

    def test_the_chosen_category_does_not_move_a_product_that_already_has_one(self):
        Product.objects.create(company=self.company, name='Ya ordenado', slug='ya-ordenado', price=10, category=self.phones)
        Product.objects.create(company=self.company, name='Sin ordenar', slug='sin-ordenar', price=10)
        body = self.preview(workbook_bytes([product_row(None, 'Ya ordenado', price=12), product_row(None, 'Sin ordenar', price=12)]),
                            options={'default_category_id': self.parts.pk}).json()
        # The preview says what will happen, not what was asked for.
        by_name = {row['data']['name']: row['data'] for row in body['rows']}
        self.assertEqual(by_name['Ya ordenado']['category_default'], '')
        self.assertEqual(by_name['Sin ordenar']['category_default'], 'Repuestos')

        self.assertEqual(self.apply(body['id']).status_code, 200)
        self.assertEqual(self.product('Ya ordenado').category, self.phones)
        self.assertEqual(self.product('Sin ordenar').category, self.parts)

    def test_a_category_that_was_switched_off_cannot_be_the_chosen_one(self):
        Category.objects.filter(pk=self.parts.pk).update(is_active=False)
        res = self.preview(workbook_bytes([product_row('CODI-1', 'Producto')]), options={'default_category_id': self.parts.pk})
        self.assertEqual(res.status_code, 400, res.content)

    def test_a_number_no_database_holds_is_one_more_refusal(self):
        for bad in (1e999, 10 ** 40, 1.5, True, [1], {'id': 1}):
            with self.subTest(repr(bad)):
                res = self.client.post(
                    f'{PREVIEW}?company={self.company.pk}',
                    {'file': SimpleUploadedFile('p.xlsx', workbook_bytes([product_row('CODB-1', 'Producto')]), content_type=XLSX),
                     'options': json.dumps({'default_category_id': bad}).replace('Infinity', '1e999')},
                    format='multipart')
                self.assertEqual(res.status_code, 400, res.content)


class TemplateTest(_Base):
    def book(self, company=None):
        return openpyxl.load_workbook(io.BytesIO(import_exports.product_template_bytes(company or self.company)))

    def test_its_sheets_and_columns_are_the_contract(self):
        workbook = self.book()

        self.assertEqual(workbook.sheetnames, ['Productos', 'Categorías', 'Instrucciones', 'Ejemplo'])
        self.assertEqual([c.value for c in workbook['Productos'][1]], [*HEADERS, 'Destino'])
        self.assertEqual(import_exports.PRODUCT_TEMPLATE_HEADERS, [*HEADERS, 'Destino'])

    def test_it_lists_this_companys_categories_and_nobody_elses(self):
        other = _p3_company('destino-otra', 'Otra Tienda')
        Category.objects.create(company=other, name='Sólo de la otra', slug='solo-de-la-otra')

        names = [row[0].value for row in self.book()['Categorías'].iter_rows(min_row=3) if row[0].value]

        self.assertEqual(names, ['Repuestos', 'Teléfonos'])
        self.assertEqual(
            [row[0].value for row in self.book(other)['Categorías'].iter_rows(min_row=3) if row[0].value],
            ['Sólo de la otra'])

    def test_the_category_and_the_destination_are_chosen_from_a_list(self):
        sheet = self.book()['Productos']
        rules = {str(rule.sqref): rule for rule in sheet.data_validations.dataValidation}

        destination = next(rule for ref, rule in rules.items() if ref.startswith('I4'))
        self.assertEqual(destination.type, 'list')
        self.assertEqual(destination.formula1, f'"{ONLINE},{INTERNAL}"')
        category = next(rule for ref, rule in rules.items() if ref.startswith('F4'))
        self.assertEqual(category.type, 'list')
        self.assertEqual(category.formula1, "'Categorías'!$A$3:$A$4")
        # Una lista ayuda; no impide escribir una categoría nueva para «crear las que falten».
        self.assertFalse(category.showErrorMessage)

    def test_a_category_name_is_text_whatever_it_looks_like(self):
        """«=1+1» is a category somebody typed, not a formula for Excel to run."""
        Category.objects.create(company=self.company, name='=1+1', slug='uno-mas-uno')
        Category.objects.create(company=self.company, name='+Ofertas, "nuevas"', slug='ofertas-nuevas')
        import zipfile

        data = import_exports.product_template_bytes(self.company)
        xml = zipfile.ZipFile(io.BytesIO(data)).read('xl/worksheets/sheet2.xml').decode()
        self.assertNotIn('<f>', xml)
        names = [row[0].value for row in openpyxl.load_workbook(io.BytesIO(data))['Categorías'].iter_rows(min_row=3) if row[0].value]
        self.assertIn('=1+1', names)
        self.assertIn('+Ofertas, "nuevas"', names)

    def test_a_company_without_categories_gets_a_template_that_still_opens(self):
        empty = _p3_company('destino-vacia', 'Tienda Vacía')
        workbook = self.book(empty)

        self.assertEqual(workbook.sheetnames, ['Productos', 'Categorías', 'Instrucciones', 'Ejemplo'])
        refs = [str(rule.sqref) for rule in workbook['Productos'].data_validations.dataValidation]
        self.assertFalse([ref for ref in refs if ref.startswith('F')])
        self.assertTrue([ref for ref in refs if ref.startswith('I')])

    def test_codes_stay_text_and_images_stay_optional(self):
        sheet = self.book()['Productos']

        for column in ('A', 'B'):
            self.assertEqual(sheet.column_dimensions[column].number_format, '@', column)
        # …and without five hundred formatted empty rows for the importer to wade through.
        self.assertLessEqual(sheet.max_row, 3)
        help_row = dict(zip(import_exports.PRODUCT_TEMPLATE_HEADERS, [c.value for c in sheet[2]]))
        self.assertTrue(help_row['Imagen principal'].startswith('Opcional'))
        self.assertTrue(help_row['Imágenes'].startswith('Opcional'))
        self.assertTrue(help_row['Destino'].startswith('Opcional'))

    def test_the_instructions_say_what_each_destination_means(self):
        text = '\n'.join(str(c.value) for row in self.book()['Instrucciones'].iter_rows() for c in row if c.value)

        self.assertIn(ONLINE, text)
        self.assertIn(INTERNAL, text)
        self.assertIn('caja', text)                     # sigue vendiéndose en tienda

    def test_downloaded_and_uploaded_as_it_is_it_creates_nothing(self):
        res = self.client.get(f'{TEMPLATE}?company={self.company.pk}')
        self.assertEqual(res.status_code, 200)
        data = b''.join(res.streaming_content) if hasattr(res, 'streaming_content') else res.content

        body = self.preview(data).json()
        self.assertEqual((body['counts']['create'], body['counts']['error']), (0, 0), body['rows'])

        names = [row[0].value for row in openpyxl.load_workbook(io.BytesIO(data))['Categorías'].iter_rows(min_row=3) if row[0].value]
        self.assertEqual(names, ['Repuestos', 'Teléfonos'])

    def test_nobody_downloads_another_companys_categories(self):
        other = _p3_company('destino-otra', 'Otra Tienda')
        Category.objects.create(company=other, name='Sólo de la otra', slug='solo-de-la-otra')

        res = self.client.get(f'{TEMPLATE}?company={other.pk}')

        self.assertIn(res.status_code, (403, 404))
        self.assertNotIn('spreadsheetml', res.get('Content-Type', ''))                # not a workbook at all


class AuthorityTest(_Base):
    def test_who_cannot_manage_products_cannot_send_a_destination(self):
        viewer, _ = _p2d_member(self.company, 'destino_lector', ['company.view', 'products.view'])
        client = APIClient()
        client.force_authenticate(user=viewer)

        res = self.preview(workbook_bytes([product_row('CODA-1', 'Producto', destination=INTERNAL)]), client=client)

        self.assertEqual(res.status_code, 403)
        self.assertFalse(Product.objects.filter(company=self.company).exists())

    def test_an_import_writes_only_in_its_own_company(self):
        other = _p3_company('destino-otra', 'Otra Tienda')
        theirs = Product.objects.create(company=other, name='Flex de carga', slug='flex-de-carga', price=30)

        self.imported([product_row('CODT-1', 'Flex de carga', destination=INTERNAL)])

        theirs.refresh_from_db()
        self.assertTrue(theirs.is_published_online)
        self.assertFalse(self.product('Flex de carga').is_published_online)

    def test_an_import_that_moves_a_product_off_the_web_or_onto_it_is_written_down(self):
        from store.models import AdminAuditLog

        self.imported([product_row('CODA-1', 'Batería', destination=ONLINE), product_row('CODA-2', 'Tapa')])
        before = AdminAuditLog.objects.filter(action='product_destination_changed').count()
        self.assertEqual(before, 0)                                # being born somewhere is not a change

        body = self.imported([product_row('CODA-1', 'Batería', destination=INTERNAL), product_row('CODA-2', 'Tapa', price=5)])

        entries = AdminAuditLog.objects.filter(action='product_destination_changed')
        self.assertEqual(entries.count(), 1)
        entry = entries.get()
        self.assertEqual(entry.target_id, str(self.product('Batería').pk))
        self.assertEqual(entry.metadata['is_published_online'], False)
        self.assertEqual(entry.metadata['import_job_id'], body['id'])
        self.assertEqual(entry.company_id, self.company.pk)

    def test_a_job_of_another_company_is_not_there_to_apply(self):
        other = _p3_company('destino-otra', 'Otra Tienda')
        outsider, _ = _p2d_member(other, 'destino_ajeno', ['company.view', 'products.view', 'products.manage'])
        client = APIClient()
        client.force_authenticate(user=outsider)
        body = self.preview(workbook_bytes([product_row('CODJ-1', 'Producto', destination=INTERNAL)])).json()

        self.assertIn(self.apply(body['id'], client=client).status_code, (403, 404))
        self.assertIn(self.apply(body['id'], client=client, company=other).status_code, (403, 404))
        self.assertFalse(Product.objects.filter(name='Producto').exists())
