"""
IMPORT-TEMPLATE-HELP — la plantilla de productos explica cómo van las imágenes.

La duda de siempre: «¿pego la foto en el Excel?». No. El Excel guarda NOMBRES
de archivo, y los archivos se adjuntan después.

IMPORT-SPLIT-01: la plantilla es UNA hoja con sus cabeceras. Lo que antes
decían una fila de ayuda, una fila de ejemplo y las hojas «Instrucciones» y
«Ejemplo» lo dicen ahora la pantalla y una nota en cada cabecera.

    NADA DE LA AYUDA SE IMPORTA COMO PRODUCTO.

Y los archivos hechos con la plantilla anterior —que la gente todavía tiene—
se siguen leyendo igual: ni su fila de ayuda, ni su fila de ejemplo, ni sus
hojas de instrucciones crean nada.
"""
import io
import json

import openpyxl
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient

from store import import_exports
from store.models import BulkImportJob, Product
from store.tests import _p2d_member, _p3_company

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PREVIEW = '/api/admin/products/import/preview/'
INSPECT = '/api/admin/imports/inspect/'


def book(data=None):
    return openpyxl.load_workbook(io.BytesIO(data or import_exports.product_template_bytes()))


def notes_of(sheet) -> dict:
    return {cell.value: cell.comment.text for cell in sheet[1] if cell.comment}


def legacy_template() -> bytes:
    """
    The template as it was before IMPORT-SPLIT-01: a help row and an example
    row under the headers, and three sheets that only explain.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Productos'
    sheet.append(import_exports.PRODUCT_TEMPLATE_HEADERS)
    sheet.append(import_exports.PRODUCT_TEMPLATE_HELP)
    sheet.append(import_exports.PRODUCT_TEMPLATE_EXAMPLE)
    categories = workbook.create_sheet('Categorías')
    categories.append(['Categorías de tu empresa'])
    categories.append(['Nombre exacto'])
    categories.append(['Teléfonos'])
    notes = workbook.create_sheet('Instrucciones')
    notes.append([None, 'Cómo llenar esta plantilla'])
    example = workbook.create_sheet('Ejemplo')
    example.append(import_exports.PRODUCT_TEMPLATE_HEADERS)
    example.append(import_exports.PRODUCT_TEMPLATE_EXAMPLE)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


class TemplateContentTest(TestCase):
    def test_it_is_one_sheet_and_its_first_row_is_the_headers(self):
        workbook = book()

        self.assertEqual(workbook.sheetnames, ['Productos'])
        sheet = workbook['Productos']
        self.assertEqual([c.value for c in sheet[1]], import_exports.PRODUCT_TEMPLATE_HEADERS)
        self.assertEqual(sheet.max_row, 1)

    def test_the_headers_say_how_images_travel(self):
        notes = notes_of(book()['Productos'])

        self.assertIn('NOMBRE del archivo', notes['Imagen principal'])
        self.assertIn('no se pega', notes['Imagen principal'])
        self.assertIn('separados por |', notes['Imágenes'])
        self.assertIn('se adjuntan', notes['Imágenes'])
        for extension in ('PNG', 'JPG', 'WEBP'):
            self.assertIn(extension, notes['Imágenes'])

    @override_settings(PRODUCT_IMAGE_MAX_PER_PRODUCT=5, STOREFRONT_IMAGE_MAX_UPLOAD_BYTES=3 * 1024 * 1024,
                       IMPORT_IMAGES_MAX_FILES=40)
    def test_the_limits_it_states_are_the_ones_the_server_applies(self):
        text = notes_of(book()['Productos'])['Imágenes']

        self.assertIn('5 imágenes por producto', text)
        self.assertIn('3 MB', text)
        self.assertIn('40 archivos', text)


class TemplateIsNeverImportedTest(TestCase):
    def setUp(self):
        cache.clear()
        self.company = _p3_company('plantilla-tienda', 'Tienda Plantilla')
        self.manager, _ = _p2d_member(
            self.company, 'plantilla_gestor', ['company.view', 'products.view', 'products.manage'])
        self.client = APIClient()
        self.client.force_authenticate(user=self.manager)

    def upload(self, data=None, name='carga-masiva-productos.xlsx'):
        return SimpleUploadedFile(name, data or import_exports.product_template_bytes(), content_type=XLSX)

    def preview(self, data=None, **fields):
        return self.client.post(
            f'{PREVIEW}?company={self.company.pk}',
            {'file': self.upload(data), 'options': json.dumps({'create_missing_categories': True}), **fields},
            format='multipart',
        )

    def test_inspecting_it_offers_its_one_sheet(self):
        res = self.client.post(
            f'{INSPECT}?company={self.company.pk}',
            {'file': self.upload(), 'import_type': 'products'}, format='multipart')

        self.assertEqual(res.status_code, 200, res.content)
        [sheet] = res.json()['sheets']
        self.assertEqual((sheet['name'], sheet['preset'], sheet['help']), ('Productos', 'products_platform', False))

    def test_inspecting_the_previous_template_marks_its_help_sheets(self):
        res = self.client.post(
            f'{INSPECT}?company={self.company.pk}',
            {'file': self.upload(legacy_template()), 'import_type': 'products'}, format='multipart')

        self.assertEqual(res.status_code, 200, res.content)
        sheets = {s['name']: s for s in res.json()['sheets']}
        self.assertEqual(sheets['Productos']['preset'], 'products_platform')
        self.assertFalse(sheets['Productos']['help'])
        for name in ('Categorías', 'Instrucciones', 'Ejemplo'):
            self.assertTrue(sheets[name]['help'], name)
            self.assertEqual(sheets[name]['preset'], '', name)
            self.assertEqual(sheets[name]['detected'], '', name)

    def test_inspecting_tells_the_screen_the_image_rules(self):
        rules = self.client.post(
            f'{INSPECT}?company={self.company.pk}',
            {'file': self.upload(), 'import_type': 'products'}, format='multipart').json()['image_rules']

        self.assertEqual(rules['separator'], '|')
        self.assertEqual(rules['formats'], ['PNG', 'JPG', 'JPEG', 'WEBP'])
        self.assertEqual(rules['max_per_product'], 12)
        self.assertEqual(rules['max_file_mb'], 8)
        self.assertEqual(rules['max_files'], 200)

    def test_uploading_the_template_as_downloaded_creates_nothing_and_complains_about_nothing(self):
        res = self.preview()

        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        self.assertEqual((body['counts']['total'], body['counts']['create'], body['counts']['error']), (0, 0, 0))
        self.assertEqual(Product.objects.filter(company=self.company).count(), 0)

    def test_the_previous_template_as_it_was_downloaded_still_creates_nothing(self):
        res = self.preview(legacy_template())

        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        self.assertEqual(body['counts']['create'], 0, body['rows'])
        self.assertEqual(body['counts']['error'], 0, body['rows'])
        warnings = ' '.join(w for row in body['rows'] for w in row['warnings'])
        self.assertIn('Fila de ayuda de la plantilla', warnings)
        self.assertIn('Fila de ejemplo de la plantilla', warnings)
        self.assertEqual(Product.objects.filter(company=self.company).count(), 0)

    def test_a_real_row_under_the_example_is_imported_and_the_example_is_not(self):
        workbook = book(legacy_template())
        sheet = workbook['Productos']
        sheet.append([None, 'COD-REAL-1', 'Funda de silicona', 'Negra', 39.9, 'Accesorios', None, None])
        out = io.BytesIO()
        workbook.save(out)

        body = self.preview(out.getvalue()).json()

        self.assertEqual(body['counts']['create'], 1, body['rows'])
        self.assertEqual(body['counts']['error'], 0, body['rows'])
        created = [row for row in body['rows'] if row['action'] == 'create']
        self.assertEqual(created[0]['data']['name'], 'Funda de silicona')

    def test_a_shop_that_really_sells_that_phone_is_not_silenced(self):
        """Sólo se salta la fila EXACTA de la plantilla anterior, no cualquier iPhone."""
        workbook = book(legacy_template())
        sheet = workbook['Productos']
        sheet.append([None, 'IP16PM-256-N', 'iPhone 16 Pro Max 256GB White', None, 4500, 'Teléfonos', None, None])
        out = io.BytesIO()
        workbook.save(out)

        body = self.preview(out.getvalue()).json()

        self.assertEqual(body['counts']['create'], 1, body['rows'])

    def test_a_help_sheet_cannot_be_chosen_as_the_products(self):
        for name in ('Ejemplo', 'Instrucciones'):
            with self.subTest(name):
                res = self.preview(legacy_template(), sheet_name=name)
                self.assertEqual(res.status_code, 400, res.content)
                self.assertIn('hoja de ayuda', res.json()['detail'])
        self.assertFalse(BulkImportJob.objects.filter(company=self.company).exists())
