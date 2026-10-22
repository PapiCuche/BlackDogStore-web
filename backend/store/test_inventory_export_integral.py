"""
INV-EXP-02 — the integral report: ONE workbook, the shop's own stock and
devices and what suppliers can deliver, each on its own sheet. Quantities are
never added across them: a supplier's units are not stock. A quantity the
supplier did not state stays «Desconocida», never a number. The agreed cost
is only for `costs.view`; the report is for reading, not for loading back.
"""
import io
from datetime import timedelta
from decimal import Decimal

from django.utils import timezone
from openpyxl import load_workbook

from store.models import Supplier, SupplierOffer
from store.test_inventory_exports import _Base

SHEETS = ['Stock propio', 'Equipos', 'Disponibilidad de proveedores']


class _Integral(_Base):
    def setUp(self):
        super().setUp()
        today = timezone.localdate()
        self.supplier = Supplier.objects.create(company=self.company, name='Mayorista Lima')
        self.known = SupplierOffer.objects.create(
            company=self.company, supplier=self.supplier, product=self.phone, agreed_cost=Decimal('3500.00'),
            commercial_price=Decimal('4400.00'), confirmed_quantity=5, committed_quantity=2,
            valid_from=today - timedelta(days=1), valid_until=today + timedelta(days=6), lead_time_days=3)
        self.unknown = SupplierOffer.objects.create(
            company=self.company, supplier=self.supplier, product=self.case, agreed_cost=Decimal('20.00'),
            commercial_price=Decimal('59.90'), confirmed_quantity=None,
            valid_from=today, valid_until=today + timedelta(days=6), lead_time_days=2)
        SupplierOffer.objects.create(  # expired: not something to sell today
            company=self.company, supplier=self.supplier, product=self.cable, commercial_price=Decimal('25.00'),
            confirmed_quantity=9, valid_from=today - timedelta(days=20), valid_until=today - timedelta(days=10),
            lead_time_days=2)
        other_supplier = Supplier.objects.create(company=self.other, name='Ajeno')
        SupplierOffer.objects.create(
            company=self.other, supplier=other_supplier, product=self.foreign, commercial_price=Decimal('10.00'),
            confirmed_quantity=50, valid_from=today, valid_until=today + timedelta(days=6), lead_time_days=1)

    def book(self, client=None, **query):
        response = self.get('integral', client=client, mode='report', **query)
        self.assertEqual(response.status_code, 200, response.content[:300])
        return load_workbook(io.BytesIO(response.content))

    def records(self, book, sheet):
        rows = [list(row) for row in book[sheet].iter_rows(values_only=True)]
        return [dict(zip(rows[0], row)) for row in rows[1:]]


class IntegralReportTest(_Integral):
    def test_three_sheets_in_one_workbook(self):
        self.assertEqual(self.book().sheetnames, SHEETS)

    def test_own_stock_is_the_stock_report(self):
        rows = self.records(self.book(), 'Stock propio')
        held = {(row['Producto'], row['Sucursal']): row['Cantidad'] for row in rows}
        self.assertEqual(held[('Funda', 'Centro')], 4)
        self.assertEqual(held[('Cable', 'Norte')], 7)

    def test_devices_are_the_devices_report(self):
        self.assertEqual(len(self.records(self.book(), 'Equipos')), 3)

    def test_supplier_availability_is_apart_and_never_added_to_stock(self):
        book = self.book()
        offers = {row['Producto']: row for row in self.records(book, 'Disponibilidad de proveedores')}
        self.assertEqual(sorted(offers), ['Funda', 'Teléfono X'])
        self.assertEqual(offers['Teléfono X']['Disponible'], 3)
        self.assertEqual(offers['Funda']['Disponible'], 'Desconocida')
        # Unknown is not «available»: the shop does not sell it on order (valid_offer_q).
        self.assertEqual((offers['Teléfono X']['Se puede vender'], offers['Funda']['Se puede vender']), ('Sí', 'No'))
        self.assertEqual(offers['Teléfono X']['Proveedor'], 'Mayorista Lima')
        self.assertEqual(offers['Teléfono X']['Línea'], 'Nuevo')
        held = {(row['Producto'], row['Sucursal']): row['Cantidad'] for row in self.records(book, 'Stock propio')}
        self.assertEqual(held[('Funda', 'Centro')], 4)  # the offer's units are not here

    def test_the_agreed_cost_needs_costs_view(self):
        offers = self.records(self.book(), 'Disponibilidad de proveedores')
        self.assertEqual({row['Costo pactado'] for row in offers}, {3500, 20})
        seller = self.member('exp_no_cost', 'products.view', 'inventory.view')
        offers = self.records(self.book(client=seller), 'Disponibilidad de proveedores')
        self.assertEqual({row['Costo pactado'] for row in offers}, {None})

    def test_the_line_filter_applies_to_every_sheet(self):
        book = self.book(line='accessory')
        self.assertEqual({row['Producto'] for row in self.records(book, 'Disponibilidad de proveedores')}, {'Funda'})
        self.assertEqual(self.records(book, 'Equipos'), [])

    def test_only_a_report(self):
        self.assertEqual(self.get('integral', mode='reimport').status_code, 400)

    def test_it_needs_catalogue_and_warehouse(self):
        only_catalogue = self.member('exp_catalogue', 'products.view')
        self.assertEqual(self.get('integral', client=only_catalogue, mode='report').status_code, 403)
        only_warehouse = self.member('exp_warehouse', 'inventory.view')
        self.assertEqual(self.get('integral', client=only_warehouse, mode='report').status_code, 403)

    def test_another_companys_offers_never_appear(self):
        names = {row['Producto'] for row in self.records(self.book(), 'Disponibilidad de proveedores')}
        self.assertNotIn('Producto ajeno', names)

    def test_it_is_audited_with_the_rows_of_each_sheet(self):
        self.book()
        entry = self.exported().latest('pk')
        self.assertEqual(entry.metadata['kind'], 'integral')
        self.assertEqual(entry.metadata['rows']['offers'], 2)
