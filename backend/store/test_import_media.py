"""
BULK-MEDIA — la carga masiva de productos, con sus imágenes.

El importador ya inspeccionaba, previsualizaba y aplicaba. Lo que llega aquí es
que una fila pueda decir QUÉ ARCHIVOS son sus imágenes, y que esos archivos
viajen con el libro: varios a la vez, o dentro de un ZIP.

Lo que no cambia, y por eso se fija otra vez:

  * nada comercial se escribe en la previsualización;
  * aplicar es todo o nada;
  * un libro sin columnas de imagen se importa exactamente como antes;
  * otra empresa no ve, no aplica y no puede citar nada de esto.
"""
import io
import json
import shutil
import tempfile
import zipfile

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from rest_framework.test import APIClient

from store import import_exports, import_media, storefront_media
from store.models import (
    AdminAuditLog, BulkImportJob, Product, ProductBarcode, ProductImage, StorefrontImage,
)
from store.test_storefront_media import jpeg, png_with_transparency, webp
from store.tests import _p2d_member, _p3_company

_ROOT = tempfile.mkdtemp(prefix='import-media-')

PREVIEW = '/api/admin/products/import/preview/'
HEADERS = ['Código', 'Nombre', 'Precio de venta', 'Imagen principal', 'Imágenes']
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def tearDownModule():
    shutil.rmtree(_ROOT, ignore_errors=True)


def workbook(rows, headers=HEADERS):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Productos'
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=header)
    for offset, values in enumerate(rows):
        for column, header in enumerate(headers, start=1):
            if header in values:
                sheet.cell(row=2 + offset, column=column, value=values[header])
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile('productos.xlsx', out.getvalue(), content_type=XLSX)


def picture(name, content=None, content_type='image/png'):
    return SimpleUploadedFile(name, content or png_with_transparency(), content_type=content_type)


def archive(members, name='imagenes.zip'):
    """`members`: nombre dentro del ZIP -> bytes, o un `ZipInfo` ya preparado."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as bundle:
        for member, content in members:
            bundle.writestr(member, content)
    return SimpleUploadedFile(name, out.getvalue(), content_type='application/zip')


@override_settings(EVIDENCE_STORAGE_BACKEND='filesystem', EVIDENCE_STORAGE_ROOT=_ROOT)
class BulkMediaBase(TestCase):
    def setUp(self):
        cache.clear()
        self.company = _p3_company('masiva-tienda', 'Tienda Masiva')
        self.other = _p3_company('masiva-otra', 'Otra Tienda')
        self.manager, _ = _p2d_member(
            self.company, 'masiva_gestor', ['company.view', 'products.view', 'products.manage'])
        self.viewer, _ = _p2d_member(self.company, 'masiva_lector', ['company.view', 'products.view'])
        self.outsider, _ = _p2d_member(
            self.other, 'masiva_ajeno', ['company.view', 'products.view', 'products.manage'])

    @staticmethod
    def _as(user):
        client = APIClient()
        client.force_authenticate(user=user)
        return client

    def _preview(self, rows, *, images=(), images_zip=None, user=None, company=None, headers=HEADERS):
        data = {'file': workbook(rows, headers)}
        if images:
            data['images'] = list(images)
        if images_zip is not None:
            data['images_zip'] = images_zip
        return self._as(user or self.manager).post(
            f'{PREVIEW}?company={(company or self.company).pk}', data, format='multipart',
        )

    def _apply(self, job_id, *, user=None, company=None):
        return self._as(user or self.manager).post(
            f'/api/admin/products/import/{job_id}/apply/?company={(company or self.company).pk}',
        )

    @staticmethod
    def _row(body, number):
        return next(row for row in body['rows'] if row['row'] == number)


class BulkMediaPreviewTest(BulkMediaBase):
    def test_a_workbook_without_image_columns_imports_as_before(self):
        res = self._preview(
            [{'Código': 'COD-A1', 'Nombre': 'Cable', 'Precio de venta': 19.9}],
            headers=['Código', 'Nombre', 'Precio de venta'],
        )
        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        self.assertEqual(body['counts']['create'], 1)
        self.assertEqual(body['counts']['error'], 0)
        self.assertEqual(body['summary']['media']['referenced'], 0)
        self.assertEqual(StorefrontImage.objects.count(), 0)

        self.assertEqual(self._apply(body['id']).status_code, 200)
        product = Product.objects.get(company=self.company, name='Cable')
        self.assertEqual(product.image_url, '')
        self.assertEqual(product.images.count(), 0)

    def test_the_preview_matches_each_row_with_its_files(self):
        res = self._preview(
            [
                {'Código': 'COD-A1', 'Nombre': 'Teléfono', 'Precio de venta': 2999,
                 'Imagen principal': 'telefono-frontal.png',
                 'Imágenes': 'telefono-frontal.png|telefono-trasera.jpg'},
                {'Código': 'COD-A2', 'Nombre': 'Cable', 'Precio de venta': 49.9,
                 'Imágenes': 'Cable.WEBP'},
                {'Código': 'COD-A3', 'Nombre': 'Funda', 'Precio de venta': 39},
            ],
            images=[
                picture('telefono-frontal.png'),
                picture('telefono-trasera.jpg', jpeg(), 'image/jpeg'),
                picture('cable.webp', webp(), 'image/webp'),
                picture('sobrante.png', png_with_transparency((60, 40))),
            ],
        )
        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        self.assertEqual(body['counts'], {
            'total': 3, 'create': 3, 'update': 0, 'no_change': 0, 'skip': 0, 'error': 0,
        })
        self.assertTrue(body['is_applicable'])

        phone = self._row(body, 2)['data']['images']
        self.assertEqual([(i['name'], i['primary']) for i in phone], [
            ('telefono-frontal.png', True), ('telefono-trasera.jpg', False),
        ])
        for item in phone:
            self.assertRegex(item['address'], r'^/api/storefront/images/[0-9a-f]{32}$')
        # El nombre del archivo se compara sin mayúsculas.
        self.assertEqual([i['name'] for i in self._row(body, 3)['data']['images']], ['Cable.WEBP'])
        self.assertEqual(self._row(body, 4)['data']['images'], [])

        media = body['summary']['media']
        self.assertEqual(media['attached'], 4)
        self.assertEqual(media['referenced'], 3)
        self.assertEqual(media['valid'], 3)
        self.assertEqual(media['invalid'], 0)
        self.assertEqual(media['missing'], 0)
        self.assertEqual(media['orphans'], 1)
        self.assertEqual(media['orphan_names'], ['sobrante.png'])

        # Nada comercial se escribió: ni productos ni galerías.
        self.assertEqual(Product.objects.filter(company=self.company).count(), 0)
        self.assertEqual(ProductImage.objects.count(), 0)
        # Las imágenes que sí se usarán esperan guardadas, en la empresa; la sobrante no.
        self.assertEqual(StorefrontImage.objects.filter(company=self.company).count(), 3)

    def test_what_is_wrong_with_an_image_is_said_on_its_row(self):
        res = self._preview(
            [
                {'Código': 'COD-A1', 'Nombre': 'Sin archivo', 'Precio de venta': 10,
                 'Imagen principal': 'no-vino.png'},
                {'Código': 'COD-A2', 'Nombre': 'Dañada', 'Precio de venta': 10,
                 'Imágenes': 'rota.png'},
                {'Código': 'COD-A3', 'Nombre': 'Falsa', 'Precio de venta': 10,
                 'Imágenes': 'programa.png'},
                {'Código': 'COD-A4', 'Nombre': 'Buena', 'Precio de venta': 10,
                 'Imágenes': 'buena.png'},
            ],
            images=[
                picture('rota.png', png_with_transparency()[:70]),
                picture('programa.png', b'MZ\x90\x00' + b'\x00' * 300),
                picture('buena.png'),
            ],
        )
        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        self.assertEqual(body['counts']['error'], 3)
        self.assertFalse(body['is_applicable'])

        self.assertIn('«no-vino.png»', self._row(body, 2)['errors'][0])
        self.assertIn('no está entre los archivos adjuntos', self._row(body, 2)['errors'][0])
        self.assertIn('«rota.png»', self._row(body, 3)['errors'][0])
        self.assertIn('«programa.png»', self._row(body, 4)['errors'][0])
        self.assertEqual(self._row(body, 5)['errors'], [])

        media = body['summary']['media']
        self.assertEqual((media['missing'], media['invalid'], media['valid']), (1, 2, 1))
        # Un trabajo con errores no se puede aplicar: no se deja nada guardado.
        self.assertEqual(StorefrontImage.objects.count(), 0)
        self.assertIsNone(self._row(body, 5)['data']['images'][0]['address'])

        report = self._as(self.manager).get(
            f'/api/admin/imports/{body["id"]}/errors.csv/?company={self.company.pk}')
        self.assertEqual(report.status_code, 200)
        self.assertIn('no-vino.png', report.content.decode('utf-8'))

    def test_the_same_picture_twice_is_stored_once(self):
        same = png_with_transparency()
        res = self._preview(
            [
                {'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imágenes': 'a.png|copia.png'},
                {'Código': 'COD-A2', 'Nombre': 'Dos', 'Precio de venta': 10, 'Imágenes': 'copia.png'},
            ],
            images=[picture('a.png', same), picture('copia.png', same)],
        )
        body = res.json()
        self.assertEqual(body['counts']['error'], 0, body)
        self.assertEqual(StorefrontImage.objects.count(), 1)
        self.assertEqual(body['summary']['media']['duplicates'], 1)
        # En la misma fila, la repetida no entra dos veces a la galería.
        row = self._row(body, 2)
        self.assertEqual([i['name'] for i in row['data']['images']], ['a.png'])
        self.assertTrue(any('copia.png' in warning for warning in row['warnings']))

    def test_two_different_files_with_the_same_name_are_ambiguous(self):
        res = self._preview(
            [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imágenes': 'foto.png'}],
            images_zip=archive([
                ('frente/foto.png', png_with_transparency()),
                ('atras/foto.png', png_with_transparency((60, 40))),
            ]),
        )
        body = res.json()
        self.assertEqual(body['counts']['error'], 1)
        self.assertIn('más de un archivo', self._row(body, 2)['errors'][0])

    @override_settings(PRODUCT_IMAGE_MAX_PER_PRODUCT=2)
    def test_a_row_cannot_exceed_the_gallery_limit(self):
        res = self._preview(
            [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imágenes': 'a.png|b.png|c.png'}],
            images=[
                picture('a.png', png_with_transparency((30, 30))),
                picture('b.png', png_with_transparency((40, 30))),
                picture('c.png', png_with_transparency((50, 30))),
            ],
        )
        body = res.json()
        self.assertEqual(body['counts']['error'], 1)
        self.assertIn('hasta 2 imágenes', self._row(body, 2)['errors'][0])

    def test_files_and_a_typed_address_on_the_same_row_keep_the_files(self):
        res = self._preview(
            [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imágenes': 'a.png',
              'URL de imagen': 'https://cdn.example.com/a.jpg'}],
            images=[picture('a.png')],
            headers=HEADERS + ['URL de imagen'],
        )
        body = res.json()
        row = self._row(body, 2)
        self.assertEqual(row['errors'], [])
        self.assertEqual(row['data']['image_url'], '')
        self.assertTrue(any('URL de imagen' in warning for warning in row['warnings']))

    def test_the_downloadable_template_previews_cleanly_and_knows_the_image_columns(self):
        self.assertIn('Imagen principal', import_exports.PRODUCT_TEMPLATE_HEADERS)
        self.assertIn('Imágenes', import_exports.PRODUCT_TEMPLATE_HEADERS)
        template = SimpleUploadedFile(
            'plantilla.xlsx', import_exports.product_template_bytes(), content_type=XLSX)

        inspected = self._as(self.manager).post(
            f'/api/admin/imports/inspect/?company={self.company.pk}',
            {'file': template, 'import_type': 'products'}, format='multipart',
        )
        self.assertEqual(inspected.status_code, 200, inspected.content)
        sheet = inspected.json()['sheets'][0]
        self.assertEqual(sheet['preset'], 'products_platform')
        self.assertEqual(set(sheet['mapping']) & {'image_main', 'image_files'},
                         {'image_main', 'image_files'})
        fields = inspected.json()['fields']
        self.assertIn('image_main', fields)
        self.assertIn('image_files', fields)

        template.seek(0)
        res = self._as(self.manager).post(
            f'{PREVIEW}?company={self.company.pk}',
            {'file': template, 'options': json.dumps({'create_missing_categories': True})},
            format='multipart',
        )
        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        # La fila de ayuda de la plantilla no es un producto.
        self.assertEqual(body['counts']['error'], 0, body['rows'])
        self.assertEqual(body['counts']['create'], 1)


class BulkMediaApplyTest(BulkMediaBase):
    def _import(self, rows, images):
        res = self._preview(rows, images=images)
        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        self.assertEqual(body['counts']['error'], 0, body)
        return body

    def test_apply_creates_the_products_with_their_galleries(self):
        body = self._import(
            [
                {'Código': 'COD-A1', 'Nombre': 'Teléfono', 'Precio de venta': 2999,
                 'Imagen principal': 'trasera.jpg', 'Imágenes': 'frontal.png|trasera.jpg'},
                {'Código': 'COD-A2', 'Nombre': 'Cable', 'Precio de venta': 49.9, 'Imágenes': 'cable.webp'},
            ],
            [picture('frontal.png'), picture('trasera.jpg', jpeg(), 'image/jpeg'),
             picture('cable.webp', webp(), 'image/webp')],
        )

        res = self._apply(body['id'])

        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()['status'], 'applied')
        self.assertEqual(res.json()['summary']['applied']['images_added'], 3)

        phone = Product.objects.get(company=self.company, name='Teléfono')
        gallery = list(phone.images.all())
        self.assertEqual(len(gallery), 2)
        primary = next(image for image in gallery if image.is_primary)
        # «Imagen principal» manda, esté donde esté en la lista.
        self.assertEqual(phone.image_url, primary.image_url)
        self.assertEqual(
            StorefrontImage.objects.get(
                public_id=storefront_media.managed_public_id(primary.image_url)).mime_type,
            'image/jpeg',
        )
        self.assertTrue(all(image.company_id == self.company.pk for image in gallery))
        self.assertTrue(all(image.uploaded_by == self.manager for image in gallery))

        cable = Product.objects.get(company=self.company, name='Cable')
        self.assertEqual(cable.images.count(), 1)
        self.assertEqual(cable.image_url, cable.images.get().image_url)

        # La tienda las sirve.
        self.assertEqual(APIClient().get(cable.image_url + '/').status_code, 200)

        audit = AdminAuditLog.objects.filter(company=self.company, action='product_import_applied').get()
        self.assertEqual(audit.metadata['images_added'], 3)

    def test_applying_twice_does_not_duplicate_anything(self):
        body = self._import(
            [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imágenes': 'a.png'}],
            [picture('a.png')],
        )
        self.assertEqual(self._apply(body['id']).status_code, 200)
        self.assertEqual(self._apply(body['id']).status_code, 200)
        self.assertEqual(ProductImage.objects.count(), 1)

    def test_importing_the_same_file_again_does_not_duplicate_the_gallery(self):
        rows = [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10,
                 'Imagen principal': 'a.png', 'Imágenes': 'a.png|b.png'}]
        files = lambda: [picture('a.png'), picture('b.png', png_with_transparency((60, 40)))]  # noqa: E731
        self._apply(self._import(rows, files())['id'])

        again = self._preview(rows, images=files()).json()

        self.assertEqual(again['counts']['error'], 0, again)
        row = self._row(again, 2)
        self.assertEqual(row['action'], 'update')
        self.assertTrue(all(item['existing'] for item in row['data']['images']))
        self.assertEqual(again['summary']['media']['already_present'], 2)
        self.assertEqual(self._apply(again['id']).status_code, 200)
        self.assertEqual(ProductImage.objects.count(), 2)

    def test_an_update_adds_to_the_gallery_and_can_move_the_primary(self):
        product = Product.objects.create(
            company=self.company, name='Uno', slug='uno', price='10.00',
        )
        ProductBarcode.objects.create(company=self.company, product=product, code='COD-A1')
        first = self._as(self.manager).post(
            f'/api/admin/products/{product.pk}/images/?company={self.company.pk}',
            {'file': picture('vieja.png', png_with_transparency((30, 30)))}, format='multipart',
        ).json()

        body = self._import(
            [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imagen principal': 'nueva.png'}],
            [picture('nueva.png')],
        )
        self.assertEqual(self._row(body, 2)['action'], 'update')
        self.assertEqual(self._apply(body['id']).status_code, 200)

        product.refresh_from_db()
        self.assertEqual(product.images.count(), 2)
        self.assertNotEqual(product.image_url, first['url'])
        self.assertEqual(product.images.get(is_primary=True).image_url, product.image_url)

    def test_a_typed_address_does_not_overwrite_a_gallery(self):
        product = Product.objects.create(
            company=self.company, name='Uno', slug='uno', price='10.00',
        )
        ProductBarcode.objects.create(company=self.company, product=product, code='COD-A1')
        placed = self._as(self.manager).post(
            f'/api/admin/products/{product.pk}/images/?company={self.company.pk}',
            {'file': picture('a.png')}, format='multipart',
        ).json()

        res = self._preview(
            [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10,
              'URL de imagen': 'https://cdn.example.com/otra.jpg'}],
            headers=['Código', 'Nombre', 'Precio de venta', 'URL de imagen'],
        )
        body = res.json()
        self.assertTrue(any('galería' in w for w in self._row(body, 2)['warnings']))
        self.assertEqual(self._apply(body['id']).status_code, 200)
        product.refresh_from_db()
        self.assertEqual(product.image_url, placed['url'])

    def test_an_image_that_expired_before_apply_aborts_everything(self):
        body = self._import(
            [
                {'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10},
                {'Código': 'COD-A2', 'Nombre': 'Dos', 'Precio de venta': 10, 'Imágenes': 'a.png'},
            ],
            [picture('a.png')],
        )
        # La limpieza de imágenes sin colocar pasó entre la previsualización y el aplicar.
        StorefrontImage.objects.update(created_at='2020-01-01T00:00:00Z')
        call_command('cleanup_storefront_images', '--older-than-hours', '1', verbosity=0)
        self.assertEqual(StorefrontImage.objects.count(), 0)

        res = self._apply(body['id'])

        self.assertEqual(res.status_code, 400, res.content)
        self.assertIn('a.png', res.json()['detail'])
        self.assertIn('previsualizar', res.json()['detail'])
        # Todo o nada: tampoco se creó el producto que no llevaba imagen.
        self.assertEqual(Product.objects.filter(company=self.company).count(), 0)
        self.assertEqual(BulkImportJob.objects.get(pk=body['id']).status, 'previewed')

    def test_a_row_cannot_be_argued_into_placing_another_company_image(self):
        foreign = storefront_media.upload(
            company=self.other, actor=self.outsider, uploaded=picture('ajena.png'),
        )
        body = self._import(
            [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imágenes': 'a.png'}],
            [picture('a.png')],
        )
        row = BulkImportJob.objects.get(pk=body['id']).rows.get(row_number=2)
        data = row.normalized_data
        data['images'][0]['address'] = foreign.url
        row.normalized_data = data
        row.save(update_fields=['normalized_data'])

        res = self._apply(body['id'])

        self.assertEqual(res.status_code, 400)
        self.assertEqual(ProductImage.objects.count(), 0)
        self.assertEqual(Product.objects.filter(company=self.company).count(), 0)
        self.assertTrue(StorefrontImage.objects.filter(pk=foreign.pk).exists())


class BulkMediaAuthorityTest(BulkMediaBase):
    ROWS = [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imágenes': 'a.png'}]

    def test_without_the_capability_nothing_is_previewed_or_stored(self):
        res = self._preview(self.ROWS, images=[picture('a.png')], user=self.viewer)
        self.assertEqual(res.status_code, 403)
        self.assertEqual(StorefrontImage.objects.count(), 0)
        self.assertEqual(BulkImportJob.objects.count(), 0)

    def test_the_capability_is_checked_again_at_apply(self):
        body = self._preview(self.ROWS, images=[picture('a.png')]).json()
        self.assertEqual(self._apply(body['id'], user=self.viewer).status_code, 403)
        self.assertEqual(ProductImage.objects.count(), 0)

    def test_another_company_cannot_apply_or_read_the_job(self):
        body = self._preview(self.ROWS, images=[picture('a.png')]).json()

        self.assertEqual(self._apply(body['id'], user=self.outsider, company=self.other).status_code, 404)
        self.assertIn(self._apply(body['id'], user=self.outsider).status_code, (403, 404))
        read = self._as(self.outsider).get(f'/api/admin/imports/{body["id"]}/?company={self.other.pk}')
        self.assertEqual(read.status_code, 404)
        self.assertEqual(ProductImage.objects.count(), 0)
        self.assertEqual(Product.objects.filter(company=self.other).count(), 0)

    def test_staged_images_belong_to_the_company_of_the_job(self):
        self._preview(self.ROWS, images=[picture('a.png')])
        self.assertEqual(
            set(StorefrontImage.objects.values_list('company_id', flat=True)), {self.company.pk})
        key = StorefrontImage.objects.get().storage_key
        self.assertRegex(key, rf'^companies/{self.company.pk}/storefront/[0-9a-f]{{32}}\.png$')

    def test_an_abandoned_preview_is_swept_by_the_cleanup(self):
        self._preview(self.ROWS, images=[picture('a.png')])
        self.assertEqual(StorefrontImage.objects.count(), 1)
        StorefrontImage.objects.update(created_at='2020-01-01T00:00:00Z')

        with self.captureOnCommitCallbacks(execute=True):
            call_command('cleanup_storefront_images', '--older-than-hours', '24', verbosity=0)

        self.assertEqual(StorefrontImage.objects.count(), 0)


class BulkMediaLimitsTest(BulkMediaBase):
    ROWS = [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imágenes': 'a.png'}]

    @override_settings(IMPORT_IMAGES_MAX_FILES=2)
    def test_too_many_files_are_refused_before_anything_is_read(self):
        res = self._preview(self.ROWS, images=[picture(f'{n}.png') for n in range(3)])
        self.assertEqual(res.status_code, 400)
        self.assertIn('2', res.json()['detail'])
        self.assertEqual(BulkImportJob.objects.count(), 0)

    @override_settings(IMPORT_IMAGES_MAX_TOTAL_BYTES=200)
    def test_too_many_bytes_are_refused_as_too_large(self):
        res = self._preview(self.ROWS, images=[picture('a.png'), picture('b.png')])
        self.assertEqual(res.status_code, 413)
        self.assertEqual(BulkImportJob.objects.count(), 0)

    @override_settings(STOREFRONT_IMAGE_MAX_UPLOAD_BYTES=300)
    def test_one_oversized_image_is_an_error_on_its_row(self):
        res = self._preview(self.ROWS, images=[picture('a.png', jpeg((600, 600)), 'image/jpeg')])
        self.assertEqual(res.status_code, 201, res.content)
        self.assertIn('«a.png»', self._row(res.json(), 2)['errors'][0])

    def test_files_that_are_not_images_are_ignored_and_reported(self):
        res = self._preview(
            self.ROWS,
            images=[picture('a.png'), SimpleUploadedFile('leeme.txt', b'hola', content_type='text/plain')],
        )
        body = res.json()
        self.assertEqual(body['counts']['error'], 0, body)
        self.assertEqual(body['summary']['media']['ignored'], 1)
        self.assertEqual(body['summary']['media']['attached'], 1)


class ZipSafetyTest(BulkMediaBase):
    ROWS = [{'Código': 'COD-A1', 'Nombre': 'Uno', 'Precio de venta': 10, 'Imágenes': 'a.png'}]

    def test_a_zip_of_images_works_like_loose_files(self):
        res = self._preview(self.ROWS, images_zip=archive([
            ('images/a.png', png_with_transparency()),
            ('images/', b''),
            ('__MACOSX/images/._a.png', b'basura'),
            ('images/.DS_Store', b'basura'),
            ('notas.txt', b'hola'),
        ]))
        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        self.assertEqual(body['counts']['error'], 0, body)
        media = body['summary']['media']
        self.assertEqual((media['attached'], media['valid'], media['ignored']), (1, 1, 1))
        self.assertEqual(self._apply(body['id']).status_code, 200)
        self.assertEqual(ProductImage.objects.count(), 1)

    def test_paths_that_leave_the_archive_refuse_the_whole_zip(self):
        for member in ('../evil.png', 'images/../../evil.png', '/etc/evil.png',
                       'C:\\Windows\\evil.png', 'images\\..\\..\\evil.png'):
            with self.subTest(member):
                res = self._preview(self.ROWS, images_zip=archive([
                    ('a.png', png_with_transparency()), (member, png_with_transparency()),
                ]))
                self.assertEqual(res.status_code, 400, res.content)
                self.assertIn('ZIP', res.json()['detail'])
        self.assertEqual(BulkImportJob.objects.count(), 0)
        self.assertEqual(StorefrontImage.objects.count(), 0)

    def test_a_symbolic_link_refuses_the_whole_zip(self):
        link = zipfile.ZipInfo('a.png')
        link.external_attr = (0o120777 << 16)
        res = self._preview(self.ROWS, images_zip=archive([(link, '/etc/passwd')]))
        self.assertEqual(res.status_code, 400)
        self.assertEqual(BulkImportJob.objects.count(), 0)

    @override_settings(IMPORT_IMAGES_ZIP_MAX_ENTRIES=3)
    def test_too_many_entries_are_refused(self):
        res = self._preview(self.ROWS, images_zip=archive(
            [(f'{n}.png', png_with_transparency()) for n in range(4)]))
        self.assertEqual(res.status_code, 400)
        self.assertIn('3', res.json()['detail'])

    @override_settings(IMPORT_IMAGES_MAX_TOTAL_BYTES=4000)
    def test_what_expands_beyond_the_limit_is_refused_without_expanding_it(self):
        # 2 MB de ceros: unos pocos KB comprimidos.
        res = self._preview(self.ROWS, images_zip=archive([('a.png', b'\x00' * 2_000_000)]))
        self.assertEqual(res.status_code, 413, res.content)

    def test_a_disproportionate_compression_ratio_is_refused(self):
        res = self._preview(self.ROWS, images_zip=archive([('a.png', b'\x00' * 6_000_000)]))
        self.assertIn(res.status_code, (400, 413), res.content)
        self.assertEqual(BulkImportJob.objects.count(), 0)

    def test_an_entry_that_lies_about_its_size_is_not_trusted(self):
        raw = png_with_transparency()
        bundle = import_media.collect([], archive([('a.png', raw)]))
        source = bundle.sources('a.png')[0]
        source.declared_size = 10  # la cabecera dice menos de lo que hay
        with self.assertRaises(import_media.ImportMediaError):
            bundle.read('a.png')

    def test_what_is_not_a_zip_is_refused(self):
        fake = SimpleUploadedFile('imagenes.zip', b'esto no es un zip', content_type='application/zip')
        res = self._preview(self.ROWS, images_zip=fake)
        self.assertEqual(res.status_code, 400)
        self.assertIn('ZIP', res.json()['detail'])

    def test_an_encrypted_zip_is_refused(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w') as bundle:
            bundle.writestr('a.png', png_with_transparency())
        data = bytearray(out.getvalue())
        # Bit 0 de las banderas generales = cifrado, en la cabecera local y en la central.
        local = data.find(b'PK\x03\x04')
        central = data.find(b'PK\x01\x02')
        data[local + 6] |= 1
        data[central + 8] |= 1
        res = self._preview(self.ROWS, images_zip=SimpleUploadedFile(
            'imagenes.zip', bytes(data), content_type='application/zip'))
        self.assertEqual(res.status_code, 400)
