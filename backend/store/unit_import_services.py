"""
UNIT-IMPORT — «Equipos serializados.xlsx»: many devices in one upload.

    ONE ROW = ONE PHYSICAL UNIT.

The stock importer writes QUANTITIES, and a serialized product has no quantity
anybody writes: it has units. That importer therefore refuses those products,
and until now the only way in was one device at a time on screen.

THIS IS NOT ANOTHER WRITER. Every row ends in
`stock_unit_services.receive_units` — the same function behind «Registrar
equipo». The serial and IMEI rules, the Kardex line and the invariant
(`BranchStock.quantity == available units`) are that function's, so a device
that arrives in a spreadsheet is exactly a device that arrived at the counter.

What this module adds is the rehearsal: stage every row, say what is wrong with
each, write nothing — and at apply time, all of the file or none of it.

NO QUANTITY COLUMN, ON PURPOSE. A row saying "iPhone 16, quantity 2" names no
device. A file that carries one is refused whole rather than read as two units
with one serial.
"""

from __future__ import annotations

import io
from collections import OrderedDict

from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from . import stock_unit_services, xlsx_reader
from .import_exports import _autosize, _header_row, _workbook
from .import_formats import normalize_header
from .import_services import CatalogueIndex, _safe_filename, _sha256
from .models import (
    AdminAuditLog, Branch, BulkImportJob, BulkImportRow, Product, StockMovement, StockUnit,
)

TEMPLATE_FILENAME = 'Equipos serializados.xlsx'
SHEET = 'Equipos'
HELP_SHEET = 'Instrucciones'

#: One file, one sitting. Above this the preview stops being something a person
#: reads, and the reception is better split by purchase document anyway.
MAX_UNITS = 1000

HEADERS = ['Código', 'Producto', 'Sucursal', 'Número de serie', 'IMEI', 'IMEI 2', 'Condición', 'Costo',
           'Motivo / referencia']

#: field -> the headers it answers to, already folded by `normalize_header`.
_ALIASES = {
    'code': ('codigo', 'codigo de barras', 'codigo del producto', 'sku', 'codigo sku', 'codigo producto'),
    'name': ('producto', 'producto modelo', 'modelo', 'nombre', 'nombre del producto'),
    'branch': ('sucursal', 'tienda', 'almacen', 'local'),
    'serial_number': ('numero de serie', 'n de serie', 'serie', 'serial', 'ns', 'numero serie', 'n serie'),
    'imei': ('imei', 'imei 1', 'imei1'),
    'imei2': ('imei 2', 'imei2', 'segundo imei'),
    'condition': ('condicion', 'estado'),
    'cost': ('costo', 'costo unitario', 'costo de compra'),
    'reason': ('motivo referencia', 'motivo', 'referencia', 'documento', 'motivo o documento de ingreso',
               'documento de ingreso'),
}
_QUANTITY_HEADERS = {'cantidad', 'cant', 'stock', 'unidades', 'qty'}

#: What marks the template's own example row, so that uploading the template
#: exactly as it was downloaded registers nothing.
EXAMPLE_SERIAL_PREFIX = 'EJEMPLO'
_EXAMPLE_ROW = ['IP16PM-256-W', 'iPhone 16 Pro Max 256GB White', 'Tienda principal', 'EJEMPLO-F2LXK0PQN7',
                '490154203237518', '', 'Nuevo', 3999.00, 'Compra factura F001-123']


class UnitImportError(Exception):
    """The file, or the job, cannot be used. The message is for the operator."""


# ---------------------------------------------------------------------------
# Template
# ---------------------------------------------------------------------------

def _conditions() -> dict[str, str]:
    """What a person may write in «Condición» -> the stored value."""
    known = {}
    for value, label in StockUnit.Condition.choices:
        known[normalize_header(label)] = value
        known[normalize_header(value)] = value
    return known


def template_bytes() -> bytes:
    """
    A blank template: the data sheet, with one example row that is recognised
    and skipped, and a sheet of instructions that is never read as data.
    """
    from openpyxl.styles import Alignment, Font

    workbook = _workbook()
    sheet = workbook.active
    sheet.title = SHEET
    _header_row(sheet, HEADERS)
    for column, value in enumerate(_EXAMPLE_ROW, start=1):
        cell = sheet.cell(row=2, column=column, value=value)
        cell.font = Font(italic=True, color='6B7280')
    # Identifiers as TEXT. A 15-digit IMEI typed into a general cell becomes
    # 4.90154E+14 on screen, and a serial that is all digits loses its zeros.
    for row in range(2, 502):
        for column in (1, 4, 5, 6):
            sheet.cell(row=row, column=column).number_format = '@'
    sheet.freeze_panes = 'A2'
    _autosize(sheet, [18, 34, 22, 22, 20, 20, 16, 12, 34])

    notes = workbook.create_sheet(HELP_SHEET)
    notes.sheet_view.showGridLines = False
    _autosize(notes, [4, 30, 82])
    title = notes.cell(row=1, column=2, value='Cómo cargar equipos con número de serie')
    title.font = Font(bold=True, size=14)
    rule = notes.cell(row=2, column=2, value='UNA FILA = UNA UNIDAD FÍSICA')
    rule.font = Font(bold=True, size=12)
    labels = ', '.join(label for _value, label in StockUnit.Condition.choices)
    steps = [
        ('Una fila por equipo',
         'Dos teléfonos iguales son dos filas, cada una con su número de serie y su IMEI. '
         'No hay columna de cantidad: una cantidad no dice qué equipos son.'),
        ('Código o Producto',
         'El código del producto (o su código de barras). Si no tiene, el nombre exacto del producto. '
         'El producto debe existir y estar marcado «con número de serie» en Inventario › Equipos.'),
        ('Sucursal', 'El nombre de la sucursal donde entra el equipo, tal como aparece en el sistema.'),
        ('Número de serie', 'Obligatorio en cada fila. No se repite dentro del mismo producto.'),
        ('IMEI',
         '15 dígitos. Obligatorio si el producto lleva línea celular; en los demás se deja vacío. '
         'No se repite en ningún equipo de la empresa.'),
        ('IMEI 2', 'Opcional, sólo para equipos con dos líneas. Distinto del primero.'),
        ('Lo que no tiene, vacío', 'No escribas «N/A», «no tiene» ni ceros: deja la celda vacía.'),
        ('Condición', f'{labels}. Vacío = Nuevo.'),
        ('Costo', 'Opcional. Lo que costó ese equipo, sin símbolo de moneda.'),
        ('Motivo / referencia',
         'El documento con el que entró: factura de compra, guía, inventario inicial. '
         'Queda en el Kardex de cada equipo.'),
        ('Antes de cargar',
         f'La fila de ejemplo (serie «{EXAMPLE_SERIAL_PREFIX}…») no se carga. El sistema muestra primero qué '
         f'haría con cada fila; si una tiene un error, no entra ninguna. Hasta {MAX_UNITS} equipos por archivo.'),
    ]
    for offset, (heading, body) in enumerate(steps, start=4):
        notes.cell(row=offset, column=1, value=offset - 3).font = Font(bold=True)
        notes.cell(row=offset, column=2, value=heading).font = Font(bold=True)
        cell = notes.cell(row=offset, column=3, value=body)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        notes.cell(row=offset, column=2).alignment = Alignment(vertical='top')
        notes.row_dimensions[offset].height = max(18, 15 * (1 + len(body) // 80))

    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Preview
# ---------------------------------------------------------------------------

def _columns(headers) -> dict[str, int]:
    """field -> column index, from the header row. Refuses a file that is not this template."""
    folded = [normalize_header(header) for header in headers]
    columns = {}
    for field, aliases in _ALIASES.items():
        for index, key in enumerate(folded):
            if key in aliases:
                columns[field] = index
                break
    if 'serial_number' not in columns:
        raise UnitImportError(
            'No se encontró la columna «Número de serie». Descarga la plantilla «Equipos serializados» '
            'y úsala tal cual.'
        )
    if 'code' not in columns and 'name' not in columns:
        raise UnitImportError('No se encontró la columna «Código» ni «Producto».')
    extra = sorted({str(headers[i]) for i, key in enumerate(folded) if key in _QUANTITY_HEADERS})
    if extra:
        raise UnitImportError(
            f'Esta plantilla no lleva la columna «{extra[0]}»: una fila es un equipo. '
            f'Para dos equipos iguales escribe dos filas, cada una con su número de serie.'
        )
    return columns


def _pick_sheet(workbook) -> str:
    names = list(workbook.sheetnames)
    for name in names:
        if normalize_header(name) == normalize_header(SHEET):
            return name
    for name in names:
        if normalize_header(name) != normalize_header(HELP_SHEET):
            return name
    raise UnitImportError('El archivo no tiene una hoja de equipos.')


@transaction.atomic
def preview_units(*, company, actor, upload, filename, reachable_branch_ids,
                  default_branch=None, default_reason: str = '') -> BulkImportJob:
    """
    Stage a workbook of devices. Writes NO unit and NO movement.

    `reachable_branch_ids` are the branches this caller may write to. A row for
    any other branch of the company is an error on that row; a branch of another
    company is not found at all, because the lookup starts from this company.
    """
    data = xlsx_reader.check_upload(upload, filename=filename)
    workbook, _notes = xlsx_reader.load_workbook(data)
    sheet_name = _pick_sheet(workbook)
    try:
        headers, rows, _truncated = xlsx_reader.read_rows(
            workbook, sheet_name, header_row=1, limit=MAX_UNITS + 1, mode=xlsx_reader.FULL_IMPORT,
        )
    except xlsx_reader.XlsxError:
        raise UnitImportError(f'Se cargan hasta {MAX_UNITS} equipos por archivo. Divide el archivo.') from None
    columns = _columns(headers)

    reachable = set(reachable_branch_ids)
    branches = {
        ' '.join(branch.name.split()).lower(): branch
        for branch in Branch.objects.filter(company=company)
    }
    fallback = None
    if default_branch is not None:
        fallback = next((b for b in branches.values() if b.pk == default_branch), None)
        if fallback is None:
            raise UnitImportError('La sucursal indicada no pertenece a esta empresa.')
    default_reason = ' '.join(str(default_reason or '').split())

    index = CatalogueIndex(company)
    products = {product.pk: product for product in Product.objects.filter(company=company)}
    conditions = _conditions()
    in_stock_serials = set(StockUnit.objects.filter(company=company).values_list('product_id', 'serial_number'))
    in_stock_imeis = set()
    for imei, imei2 in StockUnit.objects.filter(company=company).filter(
        Q(imei__isnull=False) | Q(imei2__isnull=False),
    ).values_list('imei', 'imei2'):
        in_stock_imeis.update(value for value in (imei, imei2) if value)

    job = BulkImportJob.objects.create(
        company=company, import_type=BulkImportJob.UNITS,
        original_filename=_safe_filename(filename), file_sha256=_sha256(data),
        mapping_snapshot={'sheet_name': sheet_name, 'mapping': columns,
                          'headers': [str(header) for header in headers]},
        options_snapshot={'default_branch': getattr(fallback, 'pk', None), 'default_reason': default_reason},
        created_by=actor,
    )

    staged, counts = [], dict(create=0, skip=0, error=0)
    seen_serials: dict[tuple[int, str], int] = {}
    seen_imeis: dict[str, int] = {}
    branch_ids = set()

    def stage(row_number, action, *, key='', data=None, errors=None, warnings=None):
        staged.append(BulkImportRow(
            job=job, sheet_name=sheet_name, row_number=row_number, action=action, match_key=key[:120],
            normalized_data=data or {}, errors=errors or [], warnings=warnings or [],
        ))
        counts[action] += 1

    for row_number, values, _numeric in rows:
        def cell(field):
            raw = values.get(columns.get(field)) if field in columns else None
            if raw is xlsx_reader.FORMULA:
                return raw
            return '' if raw is None else str(raw).strip()

        fields = {field: cell(field) for field in _ALIASES}
        if not any(value for value in fields.values() if value is not xlsx_reader.FORMULA):
            if any(value is xlsx_reader.FORMULA for value in fields.values()):
                stage(row_number, BulkImportRow.ERROR, errors=['La fila tiene fórmulas: escribe los valores.'])
            else:
                stage(row_number, BulkImportRow.SKIP, warnings=['Fila vacía.'])
            continue
        if any(value is xlsx_reader.FORMULA for value in fields.values()):
            stage(row_number, BulkImportRow.ERROR, errors=['La fila tiene fórmulas: escribe los valores.'])
            continue
        if fields['serial_number'].upper().startswith(EXAMPLE_SERIAL_PREFIX):
            stage(row_number, BulkImportRow.SKIP, key=fields['serial_number'],
                  warnings=['Fila de ejemplo de la plantilla: no se carga. Reemplázala por tus equipos.'])
            continue

        errors, warnings = [], []
        data = {'name': fields['name'] or fields['code']}

        # -- which product ----------------------------------------------------
        product_id, match_error, _matched_by, match_warnings = index.match(
            code=fields['code'], name=fields['name'],
        )
        warnings.extend(match_warnings)
        product = products.get(product_id)
        if match_error:
            errors.append(match_error)
        elif product is None:
            errors.append('No existe un producto con ese código o nombre en esta empresa.')
        elif not product.is_serialized:
            errors.append(
                f'«{product.name}» no se controla por número de serie. Actívalo en Inventario › Equipos '
                f'antes de cargar sus equipos.'
            )
            product = None
        if product is not None:
            data.update(product_id=product.pk, name=product.name)

        # -- which branch -----------------------------------------------------
        branch = None
        if fields['branch']:
            branch = branches.get(' '.join(fields['branch'].split()).lower())
            if branch is None:
                errors.append(f'No existe la sucursal «{fields["branch"][:60]}» en esta empresa.')
        elif fallback is not None:
            branch = fallback
        else:
            errors.append('Indica la sucursal del equipo.')
        if branch is not None:
            if not branch.is_active:
                errors.append(f'La sucursal {branch.name} está inactiva.')
                branch = None
            elif branch.pk not in reachable:
                errors.append(f'No tienes acceso a la sucursal {branch.name}.')
                branch = None
        if branch is not None:
            data.update(branch_id=branch.pk, branch=branch.name)

        # -- the device itself ------------------------------------------------
        condition = ''
        if fields['condition']:
            condition = conditions.get(normalize_header(fields['condition']), '')
            if not condition:
                labels = ', '.join(label for _value, label in StockUnit.Condition.choices)
                errors.append(f'Condición desconocida: «{fields["condition"][:40]}». Usa {labels}.')

        cleaned = None
        if product is not None:
            try:
                cleaned = stock_unit_services.clean_unit(product, {
                    'serial_number': fields['serial_number'], 'imei': fields['imei'], 'imei2': fields['imei2'],
                    'condition': condition or StockUnit.Condition.NEW, 'cost': fields['cost'],
                })
            except stock_unit_services.StockUnitError as exc:
                errors.append(str(exc))
        elif not fields['serial_number']:
            errors.append('Cada equipo lleva su número de serie.')

        if cleaned is not None:
            serial_key = (product.pk, cleaned['serial_number'])
            imeis = [value for value in (cleaned['imei'], cleaned['imei2']) if value]
            if serial_key in seen_serials:
                errors.append(
                    f'Este número de serie está repetido en el archivo (fila {seen_serials[serial_key]}).')
            elif serial_key in in_stock_serials:
                errors.append(
                    f'El número de serie {cleaned["serial_number"]} ya está registrado para este producto.')
            for value in imeis:
                if value in seen_imeis:
                    errors.append(f'El IMEI {value} está repetido en el archivo (fila {seen_imeis[value]}).')
                elif value in in_stock_imeis:
                    errors.append(f'El IMEI {value} ya está registrado en otro equipo de la empresa.')
            seen_serials.setdefault(serial_key, row_number)
            for value in imeis:
                seen_imeis.setdefault(value, row_number)
            data.update(
                serial_number=cleaned['serial_number'], imei=cleaned['imei'] or '',
                imei2=cleaned['imei2'] or '', condition=cleaned['condition'],
                cost=str(cleaned['cost']) if cleaned['cost'] is not None else '',
            )
        else:
            data.update(serial_number=fields['serial_number'][:64])

        # -- why it came in ---------------------------------------------------
        try:
            reason = stock_unit_services.clean_reason(fields['reason'] or default_reason)
            data['reason'] = reason
        except stock_unit_services.StockUnitError:
            errors.append('Indica el motivo o el documento de ingreso (en la columna o al subir el archivo).')

        if errors:
            stage(row_number, BulkImportRow.ERROR, key=fields['serial_number'], data=data,
                  errors=errors, warnings=warnings)
            continue
        branch_ids.add(branch.pk)
        stage(row_number, BulkImportRow.CREATE, key=cleaned['serial_number'], data=data, warnings=warnings)

    BulkImportRow.objects.bulk_create(staged, batch_size=500)
    job.rows_total = len(staged)
    job.rows_create = counts['create']
    job.rows_skip = counts['skip']
    job.rows_error = counts['error']
    job.mapping_snapshot = {**job.mapping_snapshot, 'branch_ids': sorted(branch_ids)}
    job.summary = {'units': counts['create']}
    job.save(update_fields=['rows_total', 'rows_create', 'rows_skip', 'rows_error',
                            'mapping_snapshot', 'summary'])
    return job


# ---------------------------------------------------------------------------
# Apply
# ---------------------------------------------------------------------------

@transaction.atomic
def apply_units(*, job, actor, request=None):
    """
    Register every staged device. ALL OF THE FILE OR NONE OF IT.

    The preview is a photograph; the stock kept moving after it was taken. So
    nothing staged is trusted to still be true: each batch goes through
    `receive_units`, which checks every serial and IMEI again against what is
    in stock NOW. The first row that can no longer enter aborts the transaction
    and names its row, and the job stays previewed.

    Returns `(job, changed)`. A second apply of an applied job changes nothing.
    """
    job = BulkImportJob.objects.select_for_update().get(pk=job.pk)
    if job.status == BulkImportJob.APPLIED:
        return job, False
    if job.import_type != BulkImportJob.UNITS:
        raise UnitImportError('Este trabajo no es una carga de equipos.')
    if job.rows_error:
        raise UnitImportError(
            f'Hay {job.rows_error} fila(s) con error. Corrige el archivo y vuelve a subirlo.')

    rows = list(job.rows.filter(action=BulkImportRow.CREATE).order_by('row_number', 'pk'))
    if not rows:
        raise UnitImportError('El archivo no tiene equipos que registrar.')

    def data(row):
        return row.normalized_data or {}

    branches = {b.pk: b for b in Branch.objects.filter(
        company=job.company, is_active=True, pk__in={data(r).get('branch_id') for r in rows})}
    products = {p.pk: p for p in Product.objects.filter(
        company=job.company, pk__in={data(r).get('product_id') for r in rows})}

    # One reception per (branch, product, document): a Kardex line a person can
    # read. Sorted, so two uploads at once take their locks in the same order.
    groups: OrderedDict[tuple, list] = OrderedDict()
    for row in sorted(rows, key=lambda r: (
        data(r).get('branch_id') or 0, data(r).get('product_id') or 0, data(r).get('reason') or '', r.row_number,
    )):
        values = data(row)
        if branches.get(values.get('branch_id')) is None or products.get(values.get('product_id')) is None:
            raise UnitImportError(
                f'El catálogo o las sucursales cambiaron desde la previsualización (fila {row.row_number}). '
                f'Vuelve a subir el archivo.'
            )
        groups.setdefault(
            (values['branch_id'], values['product_id'], values.get('reason') or ''), [],
        ).append(row)

    registered, movements = 0, 0
    for (branch_id, product_id, reason), members in groups.items():
        for start in range(0, len(members), stock_unit_services.MAX_BATCH):
            chunk = members[start:start + stock_unit_services.MAX_BATCH]
            try:
                created = stock_unit_services.receive_units(
                    branch=branches[branch_id], product=products[product_id],
                    units=[{
                        'serial_number': data(row).get('serial_number'), 'imei': data(row).get('imei'),
                        'imei2': data(row).get('imei2'), 'condition': data(row).get('condition'),
                        'cost': data(row).get('cost'),
                    } for row in chunk],
                    actor=actor, reason=reason, movement_type=StockMovement.PURCHASE_ENTRY, request=request,
                )
            except stock_unit_services.StockUnitError as exc:
                where = (
                    f'fila {chunk[exc.line - 1].row_number}' if exc.line and exc.line <= len(chunk)
                    else f'filas {chunk[0].row_number}–{chunk[-1].row_number}'
                )
                raise UnitImportError(
                    f'No se cargó ningún equipo: cambió algo desde la previsualización ({where}). {exc}'
                ) from None
            registered += len(created)
            movements += 1

    job.status = BulkImportJob.APPLIED
    job.applied_by = actor
    job.applied_at = timezone.now()
    job.summary = {**(job.summary or {}), 'applied': {'units': registered, 'movements': movements}}
    job.save(update_fields=['status', 'applied_by', 'applied_at', 'summary'])

    AdminAuditLog.log(
        actor=actor, action='stock_units_import_applied', target_type='bulk_import_job', target_id=job.pk,
        # How many and where. The identifiers are on the units themselves.
        metadata={'company_id': job.company_id, 'job_id': job.pk, 'units': registered,
                  'movements': movements, 'branch_ids': sorted({key[0] for key in groups}),
                  'file_sha256': job.file_sha256},
        request=request, company=job.company,
    )
    return job, True
