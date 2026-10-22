"""
INVENTORY-UNIFIED-01 · parte 2 — descargas del inventario.

    UN ARCHIVO PARA VOLVER A CARGAR NO ES UN REPORTE, Y EL NOMBRE LO DICE.

Tres libros, una hoja cada uno: «Productos» (el catálogo), «Inventario» (las
cantidades por sucursal) y «Equipos» (una fila por unidad física). En modo
`reimport` son exactamente lo que los importadores piden, y cargarlos sobre la
MISMA tienda no cambia nada: ni un producto, ni una cantidad, ni un equipo dos
veces. En modo `report` llevan columnas de lectura y los importadores los
rechazan.

Lo que estas pruebas fijan:

  * empresa, permisos y sucursal en cada descarga y en cada filtro;
  * el contenido de cada libro y sus filtros;
  * la ida y vuelta por cada importador sin cambios y sin stock doble;
  * que ninguna celda pueda ejecutarse como fórmula, y que al volver se lea igual;
  * el costo sólo para quien puede ver costos (`costs.view`, DC1);
  * que la descarga de siempre (`/api/admin/inventory/export/`) responde igual.
"""
import io
import json
import zipfile
from decimal import Decimal
from unittest.mock import patch

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.models import F
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from store import (
    device_catalog, import_exports, inventory_exports, stock_unit_services, unit_import_services, xlsx_reader,
)
from store.models import (
    AdminAuditLog, BranchStock, BulkImportJob, Category, CosmeticGrade, Order, Product, ProductBarcode,
    ProductVariantOption, StockMovement, StockUnit,
)
from store.test_stock_units import IMEI_A, IMEI_B, IMEI_C, IMEI_D, imei
from store.tests import _p2d_branch, _p2d_member, _p2d_stock, _p3_company

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
PICTURE = 'https://imagenes.example.test/producto.png'

ALL = ['company.view', 'products.view', 'products.manage', 'inventory.view', 'inventory.adjust', 'costs.view']
#: The product template's width. DC2 grew it («Modelo», «Capacidad», «Color»,
#: «Línea»): «Estado» and «Identificador» come right after it.
T = len(import_exports.PRODUCT_TEMPLATE_HEADERS)
#: «Modelo», «Capacidad», «Color», «Línea» of a product that is no variant.
NO_VARIANT = [None, None, None, None]


def url(kind, **query):
    text = '&'.join(f'{key}={value}' for key, value in query.items())
    return f'/api/admin/inventory/exports/{kind}/' + (f'?{text}' if text else '')


def _client(user):
    client = APIClient()
    client.force_authenticate(user=user)
    return client


def upload(content, name='archivo.xlsx'):
    return SimpleUploadedFile(name, content, content_type=XLSX)


class _Base(TestCase):
    """
    Two branches, four products and four devices in one shop; another shop next door.

      Funda          — by quantity, on the web, EAN and internal code
      Cable          — by quantity, internal stock only, a code with leading zeros
      Batería vieja  — inactive, no code at all
      Teléfono X     — by serial number, with IMEI
    """

    def setUp(self):
        cache.clear()
        self.company = _p3_company('exp-a', 'Tienda A')
        self.other = _p3_company('exp-b', 'Tienda B')
        self.centro = self.branch(self.company, 'Centro')
        self.norte = _p2d_branch(self.company, 'Norte')
        self.foreign_branch = self.branch(self.other, 'Ajena')
        self.accessories = Category.objects.create(company=self.company, name='Accesorios', slug='accesorios')
        self.phones = Category.objects.create(company=self.company, name='Teléfonos', slug='telefonos')
        self.foreign_category = Category.objects.create(company=self.other, name='Ajena', slug='ajena')

        self.case = self.product('Funda', 'funda', '59.90', self.accessories, image_url=PICTURE,
                                 description='Funda de silicona.')
        self.cable = self.product('Cable', 'cable', '25.00', self.accessories, is_published_online=False)
        self.old = self.product('Batería vieja', 'bateria-vieja', '80.00', None, is_active=False)
        self.phone = self.product('Teléfono X', 'telefono-x', '4500.00', self.phones, image_url=PICTURE,
                                  is_serialized=True, requires_imei=True)
        self.foreign = Product.objects.create(
            company=self.other, name='Producto ajeno', slug='ajeno', price=Decimal('10.00'))
        self.code(self.case, 'FUN-001')
        self.code(self.case, '7751234567892', ProductBarcode.EAN13)
        self.code(self.cable, '00123')
        self.code(self.phone, 'TEL-X')
        ProductBarcode.objects.create(company=self.other, product=self.foreign, code='AJENO-1',
                                      symbology=ProductBarcode.INTERNAL)

        _p2d_stock(self.centro, self.case, 4)
        _p2d_stock(self.norte, self.case, 1)
        _p2d_stock(self.norte, self.cable, 7)
        _p2d_stock(self.centro, self.old, 2)
        _p2d_stock(self.foreign_branch, self.foreign, 99)

        self.admin, _ = _p2d_member(self.company, 'exp_admin', ALL)
        self.api = _client(self.admin)
        self.receive(self.centro, {'serial_number': 'SER-0001', 'imei': IMEI_A, 'cost': '3800.00'},
                     {'serial_number': 'SER-0002', 'imei': IMEI_B, 'condition': 'used'})
        self.receive(self.norte, {'serial_number': 'SER-0003', 'imei': IMEI_C, 'imei2': IMEI_D,
                                  'condition': 'open_box', 'cost': '3650.50'})

    @staticmethod
    def branch(company, name):
        """The branch a new company is born with, under the name this test uses."""
        branch = company.branches.order_by('pk').first()
        if branch is None:
            return _p2d_branch(company, name)
        branch.name = name
        branch.save(update_fields=['name'])
        return branch

    def product(self, name, slug, price, category, **fields):
        return Product.objects.create(
            company=self.company, name=name, slug=slug, price=Decimal(price), category=category, **fields)

    def code(self, product, code, symbology=None):
        return ProductBarcode.objects.create(
            company=product.company, product=product, code=code,
            symbology=symbology or ProductBarcode.INTERNAL)

    def receive(self, branch, *rows, product=None):
        return stock_unit_services.receive_units(
            branch=branch, product=product or self.phone, units=list(rows), actor=self.admin,
            reason='Compra de prueba')

    def member(self, name, *capabilities, branches=None, company=None):
        user, _ = _p2d_member(company or self.company, name, ['company.view', *capabilities], branches=branches)
        return _client(user)

    def get(self, kind, client=None, **query):
        # Downloads have their own small budget per minute (see the throttle
        # test, which does not come through here): a test that reads ten files
        # is not what that budget is about.
        cache.clear()
        return (client or self.api).get(url(kind, **query))

    def sheet(self, response):
        """The ONE sheet of a workbook response: (title, headers, rows as lists)."""
        self.assertEqual(response.status_code, 200, response.content[:300])
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(len(book.sheetnames), 1, book.sheetnames)
        page = book[book.sheetnames[0]]
        rows = [list(row) for row in page.iter_rows(values_only=True)]
        return page.title, rows[0], rows[1:]

    def records(self, response):
        _title, headers, rows = self.sheet(response)
        return [dict(zip(headers, row)) for row in rows]

    def exported(self):
        return AdminAuditLog.objects.filter(action='inventory_exported')


# =============================================================================
# Who may download what
# =============================================================================

class ExportAccessTest(_Base):
    def test_each_download_needs_its_own_capability(self):
        catalogue = self.member('exp_catalogo', 'products.view')
        warehouse = self.member('exp_almacen', 'inventory.view')
        both = self.member('exp_ambos', 'products.view', 'inventory.view')
        nobody = self.member('exp_nadie')
        expected = {
            #            catalogue  warehouse  both  nobody
            'products': (200, 403, 200, 403),
            'stock':    (403, 200, 200, 403),
            'units':    (403, 200, 200, 403),
            'full':     (403, 403, 200, 403),
        }
        for kind, codes in expected.items():
            for client, code in zip((catalogue, warehouse, both, nobody), codes):
                with self.subTest(kind=kind, code=code):
                    self.assertEqual(self.get(kind, client).status_code, code)

    def test_without_a_session_nothing_downloads(self):
        for kind in inventory_exports.KINDS:
            self.assertIn(self.get(kind, APIClient()).status_code, (401, 403))

    def test_managing_or_adjusting_is_not_what_opens_a_download(self):
        """`products.manage` no lee el catálogo; `inventory.adjust` no lee el stock."""
        writer = self.member('exp_escritora', 'products.manage', 'inventory.adjust')
        for kind in inventory_exports.KINDS:
            self.assertEqual(self.get(kind, writer).status_code, 403)

    def test_another_company_downloads_its_own_and_never_this_one(self):
        stranger = self.member('exp_ajena', *ALL[1:], company=self.other)
        names = [row['Nombre'] for row in self.records(self.get('products', stranger))]
        self.assertEqual(names, ['Producto ajeno'])
        _title, headers, rows = self.sheet(self.get('stock', stranger))
        self.assertEqual(headers[3:], ['ALMACEN Ajena'])
        self.assertEqual([row[2] for row in rows], ['Producto ajeno'])
        self.assertEqual(self.records(self.get('units', stranger)), [])

    def test_the_company_never_comes_from_the_request(self):
        stranger = self.member('exp_ajena', *ALL[1:], company=self.other)
        for kind in inventory_exports.KINDS:
            with self.subTest(kind):
                self.assertIn(stranger.get(url(kind, company=self.company.pk)).status_code, (403, 404))

    def test_a_branch_limited_person_downloads_only_their_branches(self):
        limited = self.member('exp_centro', 'inventory.view', 'products.view', branches=[self.centro])

        _title, headers, _rows = self.sheet(self.get('stock', limited))
        self.assertEqual(headers[3:], ['ALMACEN Centro'])
        self.assertEqual({row['Sucursal'] for row in self.records(self.get('stock', limited, mode='report'))}, {'Centro'})
        self.assertEqual(
            [row['Número de serie'] for row in self.records(self.get('units', limited))], ['SER-0001', 'SER-0002'])

        book = zipfile.ZipFile(io.BytesIO(self.get('full', limited).content))
        devices = load_workbook(io.BytesIO(book.read('Equipos.xlsx'))).active
        self.assertEqual({row[2] for row in devices.iter_rows(min_row=2, values_only=True)}, {'Centro'})

    def test_a_branch_out_of_reach_answers_like_one_that_does_not_exist(self):
        limited = self.member('exp_centro', 'inventory.view', 'products.view', branches=[self.centro])
        for kind in ('stock', 'units', 'full'):
            with self.subTest(kind):
                out_of_reach = self.get(kind, limited, branches=self.norte.pk)
                foreign = self.get(kind, limited, branches=self.foreign_branch.pk)
                unknown = self.get(kind, limited, branches=987654)
                mixed = self.get(kind, limited, branches=f'{self.centro.pk},{self.norte.pk}')
                self.assertEqual([r.status_code for r in (out_of_reach, foreign, unknown, mixed)], [404] * 4)
                self.assertEqual(out_of_reach.content, foreign.content)
                self.assertEqual(out_of_reach.content, unknown.content)
                self.assertEqual(self.get(kind, limited, branches='centro').status_code, 400)
                self.assertEqual(self.get(kind, limited, branches=self.centro.pk).status_code, 200)

    def test_another_company_s_category_answers_like_one_that_does_not_exist(self):
        for kind in inventory_exports.KINDS:
            with self.subTest(kind):
                foreign = self.get(kind, category=self.foreign_category.pk)
                unknown = self.get(kind, category=987654)
                self.assertEqual((foreign.status_code, unknown.status_code), (404, 404))
                self.assertEqual(foreign.content, unknown.content)
                self.assertEqual(self.get(kind, category='accesorios').status_code, 400)

    def test_a_filter_value_that_means_nothing_is_refused(self):
        for query in ({'mode': 'todo'}, {'type': 'telefono'}, {'stock': 'low_stock'}, {'condition': 'roto'},
                      {'visibility': 'web'}, {'is_active': 'si'}):
            with self.subTest(query=query):
                self.assertEqual(self.get('products', **query).status_code, 400)
                self.assertEqual(self.get('units', **query).status_code, 400)

    def test_the_full_download_is_only_for_reloading(self):
        self.assertEqual(self.get('full', mode='report').status_code, 400)
        self.assertEqual(self.get('full', mode='reimport').status_code, 200)

    def test_an_unknown_kind_is_not_found_and_only_get_is_allowed(self):
        self.assertEqual(self.api.get('/api/admin/inventory/exports/kardex/').status_code, 404)
        for method in ('post', 'put', 'patch', 'delete'):
            self.assertEqual(getattr(self.api, method)(url('products')).status_code, 405)

    def test_the_file_names_are_fixed_and_nothing_is_cached(self):
        names = {
            ('products', 'reimport'): 'productos-para-cargar.xlsx', ('products', 'report'): 'reporte-productos.xlsx',
            ('stock', 'reimport'): 'stock-para-cargar.xlsx', ('stock', 'report'): 'reporte-stock.xlsx',
            ('units', 'reimport'): 'equipos-para-cargar.xlsx', ('units', 'report'): 'reporte-equipos.xlsx',
            ('full', 'reimport'): 'inventario-completo.zip',
        }
        for (kind, mode), name in names.items():
            with self.subTest(kind=kind, mode=mode):
                response = self.get(kind, mode=mode)
                self.assertEqual(response['Content-Disposition'], f'attachment; filename="{name}"')
                self.assertEqual(response['Cache-Control'], 'no-store')
                self.assertEqual(response['X-Content-Type-Options'], 'nosniff')
                self.assertEqual(response['Content-Type'], 'application/zip' if kind == 'full' else XLSX)

    def test_downloads_have_their_own_budget_and_do_not_spend_the_one_of_the_loads(self):
        """
        Review N4. A download builds a workbook in the one server process, and it
        shared 30 a minute with the loads: thirty downloads left nobody able to
        preview an import. Six a minute, of their own.
        """
        from django.conf import settings

        from store.inventory_export_views import AdminInventoryExportsView
        from store.throttles import AdminExportThrottle

        self.assertEqual(AdminInventoryExportsView.throttle_classes, [AdminExportThrottle])
        self.assertEqual(settings.REST_FRAMEWORK['DEFAULT_THROTTLE_RATES']['admin_export'], '6/min')

        cache.clear()
        codes = [self.api.get(url(kind)).status_code for kind in ('products', 'stock', 'units', 'full', 'products', 'stock', 'units')]
        self.assertEqual(codes, [200] * 6 + [429])
        self.assertEqual(self.exported().count(), 6)
        # Somebody else is not held back by this person's downloads…
        self.assertEqual(self.member('exp_otra', 'products.view').get(url('products')).status_code, 200)
        # …and neither is this person's next load.
        staged = self.api.post(
            '/api/admin/products/import/preview/',
            {'file': upload(self.api.get('/api/admin/products/import/template/').content)}, format='multipart')
        self.assertEqual(staged.status_code, 201, staged.content)

    def test_the_five_cards_one_after_another_fit_in_the_default_budget(self):
        """Owner's decision D4: the limit must not get in the way of a legitimate download of everything."""
        from django.conf import settings

        self.assertEqual(settings.REST_FRAMEWORK['DEFAULT_THROTTLE_RATES']['admin_export'], '6/min')
        cache.clear()
        cards = [url('products'), url('products', type='quantity'), url('stock'), url('units'), url('full')]
        self.assertEqual([self.api.get(card).status_code for card in cards], [200] * 5)
        self.assertEqual(self.exported().count(), 5)

    def test_past_the_budget_the_answer_says_it_in_spanish_and_when_to_come_back(self):
        cache.clear()
        for _n in range(6):
            self.assertEqual(self.api.get(url('products')).status_code, 200)

        response = self.api.get(url('products'))

        self.assertEqual(response.status_code, 429)
        seconds = int(response['Retry-After'])
        self.assertTrue(1 <= seconds <= 60, seconds)
        body = response.json()
        self.assertEqual(body['code'], 'throttled')
        self.assertEqual(body['retry_after'], seconds)
        self.assertEqual(
            body['detail'],
            'Las descargas del inventario están limitadas para no saturar el servidor. '
            f'Vuelve a intentarlo en {seconds} segundos.')
        self.assertNotIn('throttled.', body['detail'])
        self.assertEqual(self.exported().count(), 6)                    # the refused one built nothing

    def test_the_budget_is_the_one_configured(self):
        from store.throttles import AdminExportThrottle

        with patch.dict(AdminExportThrottle.THROTTLE_RATES, {'admin_export': '2/min'}):
            cache.clear()
            codes = [self.api.get(url('products')).status_code for _n in range(3)]
        self.assertEqual(codes, [200, 200, 429])

    def test_each_download_leaves_one_audit_line_without_any_cell(self):
        self.get('units', mode='report', branches=self.centro.pk, condition='used', category=self.phones.pk)

        entry = self.exported().get()
        self.assertEqual(entry.company, self.company)
        self.assertEqual(entry.actor, self.admin)
        self.assertEqual(entry.target_type, 'inventory_export')
        self.assertEqual(entry.metadata, {
            'kind': 'units', 'mode': 'report',
            'filters': {
                'branches': [self.centro.pk], 'category': self.phones.pk, 'type': '', 'stock': '',
                'condition': 'used', 'visibility': '', 'is_active': '', 'lines': [],  # INV-EXP-01
            },
            'rows': {'units': 1}, 'cost_included': True,
        })
        text = json.dumps(entry.metadata)
        for value in ('SER-0002', IMEI_B, 'Teléfono X'):
            self.assertNotIn(value, text)

    def test_the_full_download_counts_its_three_files_in_one_line(self):
        self.get('full')
        self.assertEqual(self.exported().get().metadata['rows'], {'products': 4, 'stock': 3, 'units': 3})

    def test_a_refused_download_leaves_no_audit_line(self):
        self.get('stock', branches=self.foreign_branch.pk)
        self.get('products', self.member('exp_nadie'))
        self.assertEqual(self.exported().count(), 0)


class ExportRateSettingTest(TestCase):
    """
    `ADMIN_EXPORT_RATE`: the downloads' budget, from the environment. Read ONCE,
    at start-up; a value DRF could not parse would otherwise be a 500 on every
    download, so it falls back to the default and says so in the log.
    """

    def rate(self, value):
        from backend.rates import rate_from_env

        return rate_from_env('ADMIN_EXPORT_RATE', value, '6/min')

    def test_absent_or_empty_is_the_default(self):
        for value in (None, '', '   '):
            with self.subTest(value=value):
                self.assertEqual(self.rate(value), '6/min')

    def test_a_valid_rate_is_taken_as_written(self):
        for value in ('12/min', '1/s', '300/hour', '20/m', ' 30/min ', '5000/day'):
            with self.subTest(value=value):
                self.assertEqual(self.rate(value), value.strip())

    def test_what_is_not_a_rate_falls_back_to_the_default_and_is_logged(self):
        for value in ('seis', '6', '6/', '/min', '6/fortnight', '0/min', '-3/min', '6 por minuto', '1e3/min', '6/min/2'):
            with self.subTest(value=value):
                with self.assertLogs('backend.rates', level='WARNING') as logs:
                    self.assertEqual(self.rate(value), '6/min')
                self.assertIn('ADMIN_EXPORT_RATE', logs.output[0])
                self.assertIn('6/min', logs.output[0])

    def test_every_rate_it_accepts_is_one_the_throttle_can_use(self):
        from rest_framework.throttling import SimpleRateThrottle

        for value in ('12/min', '1/s', '300/hour', '20/m', '5000/day'):
            with self.subTest(value=value):
                requests, seconds = SimpleRateThrottle.parse_rate(None, self.rate(value))
                self.assertTrue(requests > 0 and seconds > 0)

    def configured(self, value):
        """What a fresh process ends up with: the settings are read once, at start-up."""
        import os
        import pathlib
        import subprocess
        import sys

        backend = pathlib.Path(__file__).resolve().parents[1]
        env = {**os.environ, 'DJANGO_SETTINGS_MODULE': 'backend.settings'}
        env.pop('ADMIN_EXPORT_RATE', None)
        if value is not None:
            env['ADMIN_EXPORT_RATE'] = value
        return subprocess.run(
            [sys.executable, '-c',
             'import django; django.setup(); from django.conf import settings; '
             "print('RATE=' + settings.REST_FRAMEWORK['DEFAULT_THROTTLE_RATES']['admin_export'])"],
            cwd=backend, env=env, capture_output=True, text=True, timeout=120)

    def test_the_server_starts_with_the_rate_of_the_environment(self):
        result = self.configured('9/min')
        self.assertEqual(result.returncode, 0, result.stderr[-600:])
        self.assertIn('RATE=9/min', result.stdout)
        self.assertIn('RATE=6/min', self.configured(None).stdout)

    def test_the_server_starts_with_the_default_when_the_value_is_not_a_rate_and_says_so(self):
        result = self.configured('muchas')
        self.assertEqual(result.returncode, 0, result.stderr[-600:])
        self.assertIn('RATE=6/min', result.stdout)
        self.assertIn('ADMIN_EXPORT_RATE', result.stderr)

    def test_the_deployment_files_name_it_without_a_value_that_is_anybodys(self):
        import pathlib

        root = pathlib.Path(__file__).resolve().parents[2]
        for name in ('deploy/.env.production.example', '.env.example'):
            with self.subTest(name):
                self.assertIn('# ADMIN_EXPORT_RATE=6/min', (root / name).read_text(encoding='utf-8'))
        self.assertIn("'ADMIN_EXPORT_RATE'", (root / 'deploy' / 'preflight.py').read_text(encoding='utf-8'))


# =============================================================================
# The catalogue
# =============================================================================

class ProductsExportTest(_Base):
    def test_to_reload_it_is_the_product_template_plus_the_state_one_row_per_product(self):
        title, headers, rows = self.sheet(self.get('products'))

        self.assertEqual(title, 'Productos')
        # DC1: «Costo de referencia» is the template's tenth column (empty here: none set).
        # DC2: then «Modelo», «Capacidad», «Color», «Línea» (empty here: no variants).
        self.assertEqual(headers[:T + 1], import_exports.PRODUCT_TEMPLATE_HEADERS + ['Estado'])
        self.assertEqual(headers[T + 1:], ['Identificador'])
        self.assertEqual([row[:T + 1] for row in rows], [
            [None, None, 'Batería vieja', None, 80, None, None, None, 'Publicar en e-commerce', None,
             *NO_VARIANT, 'Inactivo'],
            [None, '00123', 'Cable', None, 25, 'Accesorios', None, None, 'Solo stock interno', None,
             *NO_VARIANT, 'Activo'],
            ['7751234567892', 'FUN-001', 'Funda', 'Funda de silicona.', 59.9, 'Accesorios', None, None,
             'Publicar en e-commerce', None, *NO_VARIANT, 'Activo'],
            [None, 'TEL-X', 'Teléfono X', None, 4500, 'Teléfonos', None, None, 'Publicar en e-commerce', None,
             *NO_VARIANT, 'Activo'],
        ])

    def test_every_row_carries_the_identifier_of_its_product_as_text(self):
        """Owner's decision D5: a stable identity, bound to the company, that survives a rename."""
        response = self.get('products')
        _title, headers, _rows = self.sheet(response)
        self.assertEqual(headers, import_exports.PRODUCT_TEMPLATE_HEADERS + ['Estado', 'Identificador'])
        self.assertEqual(
            {row['Nombre']: row['Identificador'] for row in self.records(response)},
            {product.name: f'{self.company.pk}-{product.pk}' for product in (self.old, self.cable, self.case, self.phone)})
        page = load_workbook(io.BytesIO(response.content)).active
        cell = page.cell(row=2, column=headers.index('Identificador') + 1)
        self.assertEqual((cell.data_type, cell.number_format), ('s', '@'))
        # Only the file to reload: the report is read, not loaded.
        self.assertNotIn('Identificador', self.sheet(self.get('products', mode='report'))[1])

    def test_a_codeless_product_renamed_after_the_download_is_updated_not_created_again(self):
        """The reviewer's R1, for an inactive and for an active product."""
        loose = self.product('Mica suelta', 'mica-suelta', '12.00', self.accessories)        # active, no code
        content = self.get('products').content
        Product.objects.filter(pk=self.old.pk).update(name='Batería vieja (renombrada)')
        Product.objects.filter(pk=loose.pk).update(name='Mica suelta (renombrada)')

        staged, applied = self.reload(content)

        self.assertEqual((staged['counts']['create'], staged['counts']['update'], staged['counts']['total']), (0, 5, 5))
        self.assertEqual((applied['created'], applied['updated']), (0, 2))
        self.assertEqual(Product.objects.filter(company=self.company).count(), 5)
        self.old.refresh_from_db()
        loose.refresh_from_db()
        # The file says the old name: it is an update of the name, of the same product.
        self.assertEqual((self.old.name, self.old.is_active), ('Batería vieja', False))
        self.assertEqual((loose.name, loose.is_active), ('Mica suelta', True))
        by_name = {row['data']['name']: row['data'] for row in staged['rows']}
        self.assertEqual(by_name['Batería vieja']['matched_by'], 'identifier')
        self.assertEqual(by_name['Batería vieja']['product_id'], self.old.pk)

    def test_loaded_into_another_shop_the_identifiers_are_ignored_and_say_so(self):
        content = self.get('products').content
        Category.objects.create(company=self.other, name='Accesorios', slug='accesorios')
        Category.objects.create(company=self.other, name='Teléfonos', slug='telefonos')
        there, _ = _p2d_member(self.other, 'exp_otra_id', ALL)
        before = list(Product.objects.filter(company=self.company).order_by('pk').values())

        staged, applied = self.reload(content, _client(there))

        self.assertEqual((staged['counts']['create'], applied['created']), (4, 4))
        for row in staged['rows']:
            self.assertIn('no es de un producto de esta empresa', ' '.join(row['warnings']))
            self.assertNotEqual(row['data']['matched_by'], 'identifier')
        self.assertEqual(list(Product.objects.filter(company=self.company).order_by('pk').values()), before)

    def test_the_owners_template_does_not_gain_the_state_column(self):
        self.assertNotIn('Estado', import_exports.PRODUCT_TEMPLATE_HEADERS)
        self.assertNotIn('Identificador', import_exports.PRODUCT_TEMPLATE_HEADERS)
        template = load_workbook(io.BytesIO(self.api.get('/api/admin/products/import/template/').content)).active
        self.assertEqual([cell.value for cell in template[1]], import_exports.PRODUCT_TEMPLATE_HEADERS)

    def test_a_file_to_reload_carries_active_and_inactive_unless_a_filter_says_otherwise(self):
        """Owner's decision D3 (2026-10-13): both by default; the filter is honoured again."""
        def names(**query):
            return [row['Nombre'] for row in self.records(self.get('products', **query))]

        self.assertEqual(names(), ['Batería vieja', 'Cable', 'Funda', 'Teléfono X'])
        self.assertEqual(names(is_active='false'), ['Batería vieja'])
        self.assertEqual(names(mode='reimport', is_active='true'), ['Cable', 'Funda', 'Teléfono X'])
        self.assertEqual(names(mode='report'), ['Batería vieja', 'Cable', 'Funda', 'Teléfono X'])
        self.assertEqual(names(mode='report', is_active='false'), ['Batería vieja'])
        # The audit line says the filter that was asked, which is the one that ran.
        self.get('products', is_active='false')
        self.assertEqual(self.exported().order_by('-pk').first().metadata['filters']['is_active'], 'false')
        self.get('products')
        self.assertEqual(self.exported().order_by('-pk').first().metadata['filters']['is_active'], '')

    def reload(self, content, client=None):
        client = client or self.api
        staged = client.post('/api/admin/products/import/preview/', {'file': upload(content)}, format='multipart')
        self.assertEqual(staged.status_code, 201, staged.content)
        self.assertEqual(staged.json()['counts']['error'], 0, staged.json()['rows'])
        applied = client.post(f'/api/admin/products/import/{staged.json()["id"]}/apply/')
        self.assertEqual(applied.status_code, 200, applied.content)
        return staged.json(), applied.json()['summary']['applied']

    def test_an_inactive_product_reloaded_into_its_own_shop_is_not_reactivated(self):
        content = self.get('products').content
        self.code(self.old, 'BAT-OLD')                    # so the row finds it whatever else changes
        content = self.get('products').content

        staged, applied = self.reload(content)

        self.assertEqual((staged['counts']['create'], staged['counts']['total']), (0, 4))
        self.assertEqual((applied['created'], applied['updated']), (0, 0))
        self.old.refresh_from_db()
        self.assertFalse(self.old.is_active)

    def test_a_file_with_inactive_products_reloaded_twice_changes_nothing(self):
        before = list(Product.objects.filter(company=self.company).order_by('pk').values())
        codes = list(ProductBarcode.objects.order_by('pk').values())
        content = self.get('products').content

        for _time in range(2):
            staged, applied = self.reload(content)
            self.assertEqual((staged['counts']['create'], staged['counts']['total']), (0, 4))
            self.assertEqual(applied, {
                'created': 0, 'updated': 0, 'images_added': 0, 'destination_changed': 0, 'publication_refused': 0})

        self.assertEqual(list(Product.objects.filter(company=self.company).order_by('pk').values()), before)
        self.assertEqual(list(ProductBarcode.objects.order_by('pk').values()), codes)
        self.assertEqual(Product.objects.filter(company=self.company, is_active=False).count(), 1)

    def test_loaded_into_another_shop_an_inactive_product_is_created_inactive_and_internal(self):
        content = self.get('products').content
        Category.objects.create(company=self.other, name='Accesorios', slug='accesorios')
        Category.objects.create(company=self.other, name='Teléfonos', slug='telefonos')
        there, _ = _p2d_member(self.other, 'exp_otra', ALL)

        staged, applied = self.reload(content, _client(there))

        self.assertEqual((staged['counts']['create'], applied['created']), (4, 4))
        old = Product.objects.get(company=self.other, name='Batería vieja')
        self.assertEqual((old.is_active, old.is_published_online), (False, False))
        row = next(row for row in staged['rows'] if row['data']['name'] == 'Batería vieja')
        self.assertIn('El producto está inactivo.', ' '.join(row['warnings']))
        funda = Product.objects.get(company=self.other, name='Funda')
        self.assertEqual((funda.is_active, funda.is_published_online), (True, True))
        # Nothing of this shop moved.
        self.old.refresh_from_db()
        self.assertEqual((self.old.is_active, self.old.company_id), (False, self.company.pk))

    def test_codes_are_text_so_no_zero_is_lost(self):
        page = load_workbook(io.BytesIO(self.get('products').content)).active
        cable = next(row for row in page.iter_rows(min_row=2) if row[2].value == 'Cable')
        funda = next(row for row in page.iter_rows(min_row=2) if row[2].value == 'Funda')
        self.assertEqual((cable[1].value, cable[1].data_type, cable[1].number_format), ('00123', 's', '@'))
        self.assertEqual((funda[0].value, funda[0].data_type, funda[0].number_format), ('7751234567892', 's', '@'))

    def test_the_filters(self):
        def names(**query):
            # The REPORT: it is the mode that lists every product, active or not.
            return [row['Nombre'] for row in self.records(self.get('products', mode='report', **query))]

        self.assertEqual(names(category=self.accessories.pk), ['Cable', 'Funda'])
        self.assertEqual(names(type='serialized'), ['Teléfono X'])
        self.assertEqual(names(type='quantity'), ['Batería vieja', 'Cable', 'Funda'])
        self.assertEqual(names(visibility='internal'), ['Cable'])
        self.assertEqual(names(visibility='published'), ['Batería vieja', 'Funda', 'Teléfono X'])
        self.assertEqual(names(is_active='false'), ['Batería vieja'])
        self.assertEqual(names(is_active='true'), ['Cable', 'Funda', 'Teléfono X'])
        # Existence, in the catalogue, is the company's total.
        Product.objects.filter(pk=self.cable.pk).update(inventory=0)
        self.assertEqual(names(stock='in_stock'), ['Batería vieja', 'Funda', 'Teléfono X'])
        self.assertEqual(names(stock='out_of_stock'), ['Cable'])
        self.assertEqual(names(stock='in_stock', type='quantity', is_active='true'), ['Funda'])

    def test_the_report_says_it_is_one_and_adds_what_is_only_read(self):
        response = self.get('products', mode='report')
        title, headers, _rows = self.sheet(response)

        self.assertEqual(title, 'Reporte de productos')
        self.assertEqual(headers, import_exports.PRODUCT_TEMPLATE_HEADERS + [
            'Activo', 'Tipo', 'Stock total', 'Publicado', 'Problemas para publicar'])
        rows = {row['Nombre']: row for row in self.records(response)}
        self.assertEqual(
            [rows['Funda'][key] for key in headers[T:]], ['Sí', 'Por cantidad', 5, 'Sí', None])
        self.assertEqual(
            [rows['Teléfono X'][key] for key in headers[T:]], ['Sí', 'Por serie', 3, 'Sí', None])
        self.assertEqual(
            [rows['Batería vieja'][key] for key in headers[T:]],
            ['No', 'Por cantidad', 2, 'Sí', 'El producto está inactivo. No tiene categoría.'])
        self.assertEqual(rows['Cable']['Publicado'], 'No')

    def test_the_report_asks_for_a_picture_only_when_the_company_does(self):
        from store.models import CompanySettings

        CompanySettings.objects.update_or_create(company=self.company, defaults={'publication_requires_image': True})
        rows = {row['Nombre']: row for row in self.records(self.get('products', mode='report'))}
        self.assertEqual(
            rows['Batería vieja']['Problemas para publicar'],
            'El producto está inactivo. No tiene imagen. No tiene categoría.')

    def test_reloading_it_into_the_same_shop_changes_nothing(self):
        before = list(Product.objects.filter(company=self.company).order_by('pk').values())
        codes = list(ProductBarcode.objects.order_by('pk').values())

        staged = self.api.post(
            '/api/admin/products/import/preview/', {'file': upload(self.get('products').content)}, format='multipart')

        self.assertEqual(staged.status_code, 201, staged.content)
        counts = staged.json()['counts']
        self.assertEqual((counts['error'], counts['create'], counts['total']), (0, 0, 4))
        applied = self.api.post(f'/api/admin/products/import/{staged.json()["id"]}/apply/')
        self.assertEqual(applied.status_code, 200, applied.content)
        self.assertEqual(
            applied.json()['summary']['applied'],
            {'created': 0, 'updated': 0, 'images_added': 0, 'destination_changed': 0, 'publication_refused': 0})
        self.assertEqual(list(Product.objects.filter(company=self.company).order_by('pk').values()), before)
        self.assertEqual(list(ProductBarcode.objects.order_by('pk').values()), codes)
        self.assertFalse(AdminAuditLog.objects.filter(action='product_destination_changed').exists())

    def test_the_product_importer_refuses_the_report(self):
        """Elegido y fijado: un reporte no se importa, ni en parte."""
        report = self.get('products', mode='report').content
        response = self.api.post('/api/admin/products/import/preview/', {'file': upload(report)}, format='multipart')

        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn('es un reporte', response.json()['detail'])
        self.assertIn('para volver a cargar', response.json()['detail'])
        self.assertFalse(BulkImportJob.objects.exists())


# =============================================================================
# Quantities by branch
# =============================================================================

class StockExportTest(_Base):
    MAPPING = json.dumps({'code': 0, 'barcode': 1, 'name': 2})

    def reload(self, content, columns=None):
        columns = columns or {3: self.centro.pk, 4: self.norte.pk}
        return self.api.post('/api/admin/inventory/import/preview/', {
            'file': upload(content), 'mapping': self.MAPPING, 'header_row': 1,
            'branch_map': json.dumps({str(index): branch for index, branch in columns.items()}),
        }, format='multipart')

    def test_to_reload_it_has_the_inventory_shape_without_the_serialized_products(self):
        title, headers, rows = self.sheet(self.get('stock'))

        self.assertEqual(title, 'Inventario')
        self.assertEqual(headers, ['CODIGO', 'CODIGO EAN', 'NOMBRE', 'ALMACEN Centro', 'ALMACEN Norte'])
        self.assertEqual(rows, [
            [None, None, 'Batería vieja', 2, 0],
            ['00123', None, 'Cable', 0, 7],
            ['FUN-001', '7751234567892', 'Funda', 4, 1],
        ])

    def test_reloading_it_has_no_error_row_and_moves_nothing(self):
        def held():
            # Absence and zero are the same shelf: the importer may create the
            # row of a pair that had none, with nothing in it.
            return {
                (branch, product): quantity
                for branch, product, quantity in BranchStock.objects.values_list('branch_id', 'product_id', 'quantity')
                if quantity
            }

        movements = StockMovement.objects.count()
        stock, totals = held(), dict(Product.objects.values_list('pk', 'inventory'))

        staged = self.reload(self.get('stock').content)

        self.assertEqual(staged.status_code, 201, staged.content)
        counts = staged.json()['counts']
        self.assertEqual((counts['error'], counts['update'], counts['create']), (0, 0, 0))
        self.assertEqual(counts['no_change'], 6)                      # three products, two branches
        applied = self.api.post(f'/api/admin/inventory/import/{staged.json()["id"]}/apply/')
        self.assertEqual(applied.status_code, 200, applied.content)
        self.assertEqual(StockMovement.objects.count(), movements)
        self.assertEqual(held(), stock)
        self.assertEqual(dict(Product.objects.values_list('pk', 'inventory')), totals)

    def test_the_mismatch_this_fixes(self):
        """
        La descarga de siempre incluye los productos con serie, con su número, y
        el importador de stock rechaza cualquier celda con valor de uno de ellos:
        el archivo que la plataforma entregaba no volvía a entrar.
        """
        old = self.api.get('/api/admin/inventory/export/')
        staged = self.reload(old.content)
        self.assertEqual(staged.json()['counts']['error'], 2)          # Teléfono X, en dos sucursales

    def test_the_filters(self):
        def names(**query):
            return [row[2] for row in self.sheet(self.get('stock', **query))[2]]

        self.assertEqual(self.sheet(self.get('stock', branches=self.norte.pk))[1][3:], ['ALMACEN Norte'])
        self.assertEqual(names(category=self.accessories.pk), ['Cable', 'Funda'])
        self.assertEqual(names(is_active='true'), ['Cable', 'Funda'])
        self.assertEqual(names(visibility='internal'), ['Cable'])
        self.assertEqual(names(type='serialized'), [])
        # Existence is counted in the branches being downloaded.
        self.assertEqual(names(stock='in_stock', branches=self.centro.pk), ['Batería vieja', 'Funda'])
        self.assertEqual(names(stock='out_of_stock', branches=self.centro.pk), ['Cable'])

    def test_the_report_is_one_row_per_product_and_branch_with_the_devices_counted(self):
        response = self.get('stock', mode='report')
        title, headers, _rows = self.sheet(response)

        self.assertEqual(title, 'Reporte de stock')
        self.assertEqual(headers, ['Código', 'Código de barras', 'Producto', 'Categoría', 'Tipo', 'Sucursal', 'Cantidad',
                                   'Costo promedio'])
        found = {(row['Producto'], row['Sucursal']): (row['Tipo'], row['Cantidad']) for row in self.records(response)}
        self.assertEqual(found, {
            ('Batería vieja', 'Centro'): ('Por cantidad', 2), ('Batería vieja', 'Norte'): ('Por cantidad', 0),
            ('Cable', 'Centro'): ('Por cantidad', 0), ('Cable', 'Norte'): ('Por cantidad', 7),
            ('Funda', 'Centro'): ('Por cantidad', 4), ('Funda', 'Norte'): ('Por cantidad', 1),
            ('Teléfono X', 'Centro'): ('Por serie', 2), ('Teléfono X', 'Norte'): ('Por serie', 1),
        })
        in_stock = self.records(self.get('stock', mode='report', stock='in_stock', type='quantity'))
        self.assertEqual(
            [(row['Producto'], row['Sucursal']) for row in in_stock],
            [('Batería vieja', 'Centro'), ('Cable', 'Norte'), ('Funda', 'Centro'), ('Funda', 'Norte')])

    def test_the_stock_importer_refuses_the_report(self):
        report = self.get('stock', mode='report').content
        response = self.api.post('/api/admin/inventory/import/preview/', {
            'file': upload(report), 'mapping': json.dumps({'code': 0, 'barcode': 1, 'name': 2}),
            'branch_map': json.dumps({'6': self.centro.pk}),
        }, format='multipart')
        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn('es un reporte', response.json()['detail'])
        self.assertFalse(BulkImportJob.objects.exists())


# =============================================================================
# Devices, one row each
# =============================================================================

class UnitsExportTest(_Base):
    def test_to_reload_it_is_the_device_template_one_row_per_unit_on_the_shelf(self):
        title, headers, rows = self.sheet(self.get('units'))

        self.assertEqual(title, unit_import_services.SHEET)
        self.assertEqual(headers, unit_import_services.HEADERS)
        # DC2: «Modelo» … «Precio de venta» after «Motivo / referencia», empty here
        # (no variant, grade, battery or price).
        dc2 = [None] * 7
        self.assertEqual(rows, [
            ['TEL-X', 'Teléfono X', 'Centro', 'SER-0001', IMEI_A, None, 'Nuevo', 3800, None, *dc2],
            ['TEL-X', 'Teléfono X', 'Centro', 'SER-0002', IMEI_B, None, 'Usado', None, None, *dc2],
            ['TEL-X', 'Teléfono X', 'Norte', 'SER-0003', IMEI_C, IMEI_D, 'Caja abierta', 3650.5, None, *dc2],
        ])

    def test_serials_and_imeis_are_text(self):
        page = load_workbook(io.BytesIO(self.get('units').content)).active
        first = next(page.iter_rows(min_row=2))
        for index in (0, 3, 4):
            self.assertEqual((first[index].data_type, first[index].number_format), ('s', '@'))
        self.assertEqual(first[4].value, IMEI_A)

    def test_the_cost_is_only_for_who_may_view_costs(self):
        # DC1: `inventory.adjust` no longer implies it.
        adjuster = self.member('exp_ajuste', 'inventory.view', 'inventory.adjust')
        self.assertEqual({row['Costo'] for row in self.records(self.get('units', adjuster))}, {None})
        reader = self.member('exp_lectora', 'inventory.view')
        for mode in ('reimport', 'report'):
            with self.subTest(mode):
                self.assertEqual({row['Costo'] for row in self.records(self.get('units', reader, mode=mode))}, {None})
                self.assertIn(3800, {row['Costo'] for row in self.records(self.get('units', mode=mode))})
        both = self.get('full', self.member('exp_dos', 'inventory.view', 'products.view'))
        devices = load_workbook(io.BytesIO(zipfile.ZipFile(io.BytesIO(both.content)).read('Equipos.xlsx'))).active
        self.assertEqual({row[7] for row in devices.iter_rows(min_row=2, values_only=True)}, {None})
        mine = load_workbook(io.BytesIO(zipfile.ZipFile(io.BytesIO(self.get('full').content)).read('Equipos.xlsx'))).active
        self.assertIn(3800, {row[7] for row in mine.iter_rows(min_row=2, values_only=True)})
        self.assertIs(self.exported().filter(actor__username='exp_lectora').first().metadata['cost_included'], False)

    def out_of_the_shelf(self):
        """One device written off and one sold, the second with its order."""
        gone, sold = self.receive(self.centro, {'serial_number': 'SER-0004'}, {'serial_number': 'SER-0005'},
                                  product=self.laptop())
        stock_unit_services.write_off(gone, reason='Pantalla rota', actor=self.admin)
        order = Order.objects.create(
            company=self.company, fulfillment_branch=self.centro, customer_email='compra@correo.test',
            total=Decimal('5200.00'), status=Order.Status.PAID, paid=True, paid_at=timezone.now())
        StockUnit.objects.filter(pk=sold.pk).update(
            status=StockUnit.Status.SOLD, sold_at=timezone.now(), order=order, notes='Entregado en tienda')
        BranchStock.objects.filter(branch=self.centro, product=sold.product).update(quantity=F('quantity') - 1)
        return order

    def laptop(self):
        laptop = self.product('Portátil', 'portatil', '5200.00', self.phones, is_serialized=True)
        self.code(laptop, 'PORT-1')
        return laptop

    def test_what_left_the_shelf_is_not_in_the_file_to_reload(self):
        self.out_of_the_shelf()
        serials = [row['Número de serie'] for row in self.records(self.get('units'))]
        self.assertEqual(serials, ['SER-0001', 'SER-0002', 'SER-0003'])
        # Asking for it does not bring it back: reloading it would put it on sale.
        self.assertEqual(self.records(self.get('units', stock='out_of_stock')), [])

    def test_the_report_has_every_device_with_its_state_date_order_and_notes(self):
        order = self.out_of_the_shelf()
        response = self.get('units', mode='report')
        title, headers, _rows = self.sheet(response)

        self.assertEqual(title, 'Reporte de equipos')
        # The device template without «Motivo / referencia» (a reception's reason), plus what happened.
        self.assertEqual(headers, [header for header in unit_import_services.HEADERS if header != 'Motivo / referencia']
                         + ['Estado', 'Fecha de ingreso', 'Pedido', 'Notas'])
        rows = {row['Número de serie']: row for row in self.records(response)}
        self.assertEqual(set(rows), {'SER-0001', 'SER-0002', 'SER-0003', 'SER-0004', 'SER-0005'})
        self.assertEqual((rows['SER-0001']['Estado'], rows['SER-0001']['Pedido']), ('Disponible', None))
        self.assertEqual(rows['SER-0004']['Estado'], 'Dado de baja')
        self.assertEqual(
            (rows['SER-0005']['Estado'], rows['SER-0005']['Pedido'], rows['SER-0005']['Notas']),
            ('Vendido', order.pk, 'Entregado en tienda'))
        self.assertRegex(rows['SER-0001']['Fecha de ingreso'], r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$')

        gone = self.records(self.get('units', mode='report', stock='out_of_stock'))
        self.assertEqual([row['Número de serie'] for row in gone], ['SER-0004', 'SER-0005'])
        there = self.records(self.get('units', mode='report', stock='in_stock'))
        self.assertEqual(len(there), 3)

    def test_the_filters(self):
        def serials(**query):
            return [row['Número de serie'] for row in self.records(self.get('units', **query))]

        self.receive(self.norte, {'serial_number': 'SER-0009'}, product=self.laptop())
        self.assertEqual(serials(branches=self.norte.pk), ['SER-0009', 'SER-0003'])
        self.assertEqual(serials(condition='used'), ['SER-0002'])
        self.assertEqual(serials(condition='open_box'), ['SER-0003'])
        self.assertEqual(serials(category=self.accessories.pk), [])
        self.assertEqual(serials(type='quantity'), [])
        Product.objects.filter(pk=self.phone.pk).update(is_published_online=False, is_active=False)
        self.assertEqual(serials(visibility='published'), ['SER-0009'])
        self.assertEqual(serials(is_active='false'), ['SER-0001', 'SER-0002', 'SER-0003'])

    def test_reloading_it_registers_no_device_twice(self):
        """Cada fila se rechaza porque ese equipo ya está: un rechazo limpio, no stock doble."""
        units = list(StockUnit.objects.order_by('pk').values())
        movements = StockMovement.objects.count()
        stock = list(BranchStock.objects.order_by('pk').values())

        staged = self.api.post(
            '/api/admin/inventory/units/import/preview/',
            {'file': upload(self.get('units').content), 'reason': 'Recarga de prueba'}, format='multipart')

        self.assertEqual(staged.status_code, 201, staged.content)
        body = staged.json()
        self.assertEqual((body['counts']['create'], body['counts']['error'], body['counts']['total']), (0, 3, 3))
        for row in body['rows']:
            self.assertIn('ya está registrado', ' '.join(row['errors']))
            self.assertNotIn('No existe', ' '.join(row['errors']))
        applied = self.api.post(f'/api/admin/inventory/units/import/{body["id"]}/apply/')
        self.assertEqual(applied.status_code, 400, applied.content)
        self.assertEqual(list(StockUnit.objects.order_by('pk').values()), units)
        self.assertEqual(StockMovement.objects.count(), movements)
        self.assertEqual(list(BranchStock.objects.order_by('pk').values()), stock)

    def test_the_device_importer_refuses_the_report(self):
        report = self.get('units', mode='report').content
        response = self.api.post(
            '/api/admin/inventory/units/import/preview/', {'file': upload(report), 'reason': 'x'}, format='multipart')
        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn('es un reporte', response.json()['detail'])
        self.assertFalse(BulkImportJob.objects.exists())


# =============================================================================
# DC2 — the files to reload carry the variant, and read back as themselves
# =============================================================================

class DeviceCommerceRoundTripTest(_Base):
    """
    DC2 grew both templates. A download is «the template plus the state and the
    identifier», every template column filled from the product or the device
    and every value under ITS header: a reload creates nothing and moves no
    model or variant link.

      Teléfono X       — a NEW variant of «Serie 16», 128 GB Negro
      Teléfono X usado — a USED variant of «Serie 16», 128 GB Azul, one device
                         with its grade, its battery and its own sale price
    """

    def setUp(self):
        super().setUp()
        device_catalog.seed_cosmetic_grades(self.company)
        self.grade = CosmeticGrade.objects.filter(company=self.company, is_active=True).order_by('sort_order').first()
        capacity = device_catalog.create_attribute(company=self.company, name='Capacidad')
        colour = device_catalog.create_attribute(company=self.company, name='Color', kind='color')
        self.gb128 = device_catalog.create_option(company=self.company, attribute_id=capacity.pk, label='128 GB')
        black = device_catalog.create_option(company=self.company, attribute_id=colour.pk, label='Negro',
                                             swatch_hex='#111111')
        blue = device_catalog.create_option(company=self.company, attribute_id=colour.pk, label='Azul',
                                            swatch_hex='#1d4ed8')
        self.model = device_catalog.create_model(company=self.company, name='Serie 16')
        device_catalog.set_variant(self.phone, model_id=self.model.pk, option_ids=[self.gb128.pk, black.pk])
        self.used = self.product('Teléfono X usado', 'telefono-x-usado', '1.00', self.phones,
                                 is_serialized=True, requires_imei=True)
        self.code(self.used, 'TEL-X-U')
        device_catalog.set_variant(self.used, model_id=self.model.pk, condition_line=Product.CONDITION_USED,
                                   option_ids=[self.gb128.pk, blue.pk])
        Product.objects.filter(pk=self.phone.pk).update(reference_cost=Decimal('3700.00'))
        self.receive(self.centro, {
            'serial_number': 'SER-U001', 'imei': imei('35693803564390'), 'condition': 'used', 'cost': '1200.00',
            'cosmetic_grade': self.grade.pk, 'battery_health': 88, 'sale_price': '1899.90',
        }, product=self.used)

    def links(self):
        return (
            list(Product.objects.filter(company=self.company).order_by('pk')
                 .values_list('pk', 'product_model_id', 'condition_line', 'variant_key')),
            sorted(ProductVariantOption.objects.filter(company=self.company)
                   .values_list('product_id', 'attribute_id', 'option_id')),
        )

    def test_every_column_of_both_templates_is_filled_under_its_own_header(self):
        """The drift guard: a column the template gains is written, or left empty ON PURPOSE."""
        _title, headers, rows = self.sheet(self.get('products'))
        self.assertEqual(headers, import_exports.PRODUCT_TEMPLATE_HEADERS + ['Estado', 'Identificador'])
        filled = {header for row in rows for header, value in zip(headers, row) if value is not None}
        self.assertEqual(set(headers) - filled, set(inventory_exports.PRODUCT_LEFT_EMPTY))

        _title, headers, rows = self.sheet(self.get('units'))
        self.assertEqual(headers, unit_import_services.HEADERS)
        filled = {header for row in rows for header, value in zip(headers, row) if value is not None}
        self.assertEqual(set(headers) - filled, set(inventory_exports.UNIT_LEFT_EMPTY))

    def test_the_variant_grade_battery_and_price_are_written_as_the_loads_read_them(self):
        products = {row['Nombre']: row for row in self.records(self.get('products'))}
        self.assertEqual(
            [products['Teléfono X'][key] for key in ('Modelo', 'Capacidad', 'Color', 'Línea', 'Estado', 'Identificador')],
            ['Serie 16', '128 GB', 'Negro', 'Nuevo', 'Activo', f'{self.company.pk}-{self.phone.pk}'])
        self.assertEqual(
            [products['Teléfono X usado'][key] for key in ('Modelo', 'Capacidad', 'Color', 'Línea')],
            ['Serie 16', '128 GB', 'Azul', 'Seminuevo'])
        self.assertEqual([products['Funda'][key] for key in ('Modelo', 'Línea', 'Estado')], [None, None, 'Activo'])

        units = {row['Número de serie']: row for row in self.records(self.get('units'))}
        self.assertEqual(
            [units['SER-U001'][key] for key in ('Modelo', 'Capacidad', 'Color', 'Línea', 'Estado estético',
                                                'Batería %', 'Precio de venta', 'Motivo / referencia')],
            ['Serie 16', '128 GB', 'Azul', 'Seminuevo', self.grade.label, 88, 1899.9, None])
        # A new device sells at its variant's price: none of its own in the file.
        self.assertEqual((units['SER-0001']['Línea'], units['SER-0001']['Precio de venta']), ('Nuevo', None))
        # The report reads the same, and says what happened to each device.
        report = {row['Número de serie']: row for row in self.records(self.get('units', mode='report'))}
        self.assertEqual(
            [report['SER-U001'][key] for key in ('Estado estético', 'Batería %', 'Precio de venta', 'Estado')],
            [self.grade.label, 88, 1899.9, 'Disponible'])
        self.assertNotIn('Motivo / referencia', report['SER-U001'])

    def test_the_sale_price_is_not_a_cost_and_the_costs_still_are(self):
        reader = self.member('exp_dc2_lectora', 'inventory.view', 'products.view')
        units = {row['Número de serie']: row for row in self.records(self.get('units', reader))}
        self.assertEqual((units['SER-U001']['Precio de venta'], units['SER-U001']['Costo']), (1899.9, None))
        products = {row['Nombre']: row for row in self.records(self.get('products', reader))}
        self.assertIsNone(products['Teléfono X']['Costo de referencia'])
        mine = {row['Nombre']: row for row in self.records(self.get('products'))}
        self.assertEqual(mine['Teléfono X']['Costo de referencia'], 3700)

    def test_reloading_the_catalogue_creates_nothing_and_moves_no_variant(self):
        before, links = list(Product.objects.filter(company=self.company).order_by('pk').values()), self.links()
        content = self.get('products').content

        staged = self.api.post('/api/admin/products/import/preview/', {'file': upload(content)}, format='multipart')

        self.assertEqual(staged.status_code, 201, staged.content)
        body = staged.json()
        self.assertEqual((body['counts']['error'], body['counts']['create'], body['counts']['total']), (0, 0, 5),
                         body['rows'])
        by_name = {row['data']['name']: row['data'] for row in body['rows']}
        # The variant columns were READ (not dropped), each as the variant the product already is.
        self.assertEqual(by_name['Teléfono X']['variant']['model_id'], self.model.pk)
        self.assertEqual(by_name['Teléfono X usado']['variant']['condition_line'], Product.CONDITION_USED)
        for name in ('Teléfono X', 'Teléfono X usado', 'Batería vieja'):
            self.assertEqual(by_name[name]['matched_by'], 'identifier')
        applied = self.api.post(f'/api/admin/products/import/{body["id"]}/apply/')
        self.assertEqual(applied.status_code, 200, applied.content)
        self.assertEqual((applied.json()['summary']['applied']['created'],
                          applied.json()['summary']['applied']['updated']), (0, 0))
        self.assertEqual(list(Product.objects.filter(company=self.company).order_by('pk').values()), before)
        self.assertEqual(self.links(), links)

    def test_reloading_the_devices_refuses_each_one_only_because_it_is_already_registered(self):
        units, links = list(StockUnit.objects.order_by('pk').values()), self.links()
        movements = StockMovement.objects.count()

        staged = self.api.post(
            '/api/admin/inventory/units/import/preview/',
            {'file': upload(self.get('units').content), 'reason': 'Recarga de prueba'}, format='multipart')

        self.assertEqual(staged.status_code, 201, staged.content)
        body = staged.json()
        self.assertEqual((body['counts']['create'], body['counts']['total']), (0, 4))
        for row in body['rows']:
            # Only «already registered» (its serial, its IMEI): model, line, grade,
            # battery and price all read back clean.
            self.assertTrue(row['errors'], row)
            for error in row['errors']:
                self.assertIn('ya está registrado', error, row)
        used = next(row['data'] for row in body['rows'] if row['data']['serial_number'] == 'SER-U001')
        self.assertEqual(used['product_id'], self.used.pk)
        self.assertEqual((used['cosmetic_grade'], used['battery_health'], used['sale_price']),
                         (self.grade.pk, 88, '1899.90'))
        applied = self.api.post(f'/api/admin/inventory/units/import/{body["id"]}/apply/')
        self.assertEqual(applied.status_code, 400, applied.content)
        self.assertEqual(list(StockUnit.objects.order_by('pk').values()), units)
        self.assertEqual(StockMovement.objects.count(), movements)
        self.assertEqual(self.links(), links)

    def test_a_model_name_two_models_share_is_left_empty_and_the_reload_keeps_the_variant(self):
        """`resolve_variant_spec` takes the FIRST model of that name: written, it could re-link the product."""
        device_catalog.create_model(company=self.company, name='SERIE 16')
        links = self.links()
        records = {row['Nombre']: row for row in self.records(self.get('products'))}
        self.assertEqual([records['Teléfono X'][key] for key in ('Modelo', 'Capacidad', 'Color', 'Línea')],
                         [None, None, None, None])
        staged = self.api.post('/api/admin/products/import/preview/',
                               {'file': upload(self.get('products').content)}, format='multipart')
        self.assertEqual((staged.json()['counts']['error'], staged.json()['counts']['create']), (0, 0))
        self.api.post(f'/api/admin/products/import/{staged.json()["id"]}/apply/')
        self.assertEqual(self.links(), links)

    def test_an_option_label_two_attributes_share_is_left_empty_too(self):
        storage = device_catalog.create_attribute(company=self.company, name='Almacenamiento')
        device_catalog.create_option(company=self.company, attribute_id=storage.pk, label='128 gb')
        records = {row['Nombre']: row for row in self.records(self.get('products'))}
        self.assertIsNone(records['Teléfono X']['Modelo'])
        units = {row['Número de serie']: row for row in self.records(self.get('units'))}
        self.assertIsNone(units['SER-U001']['Modelo'])
        # The rest of the device still travels.
        self.assertEqual(units['SER-U001']['Batería %'], 88)


# =============================================================================
# Everything, in one ZIP
# =============================================================================

class FullExportTest(_Base):
    def archive(self, **query):
        response = self.get('full', **query)
        self.assertEqual(response.status_code, 200, response.content[:300])
        return zipfile.ZipFile(io.BytesIO(response.content))

    def page(self, archive, name):
        book = load_workbook(io.BytesIO(archive.read(name)))
        self.assertEqual(len(book.sheetnames), 1, name)
        return book.active

    def test_three_workbooks_of_one_sheet_and_a_note(self):
        archive = self.archive()
        self.assertEqual(archive.namelist(), ['LEEME.txt', 'Productos.xlsx', 'Stock.xlsx', 'Equipos.xlsx'])
        self.assertIsNone(archive.testzip())
        self.assertEqual(self.page(archive, 'Productos.xlsx').title, 'Productos')
        self.assertEqual(self.page(archive, 'Stock.xlsx').title, 'Inventario')
        self.assertEqual(self.page(archive, 'Equipos.xlsx').title, 'Equipos')

    def test_they_are_the_files_to_reload_active_and_inactive_unless_filtered(self):
        archive = self.archive()
        products = self.page(archive, 'Productos.xlsx')
        stock = self.page(archive, 'Stock.xlsx')
        devices = self.page(archive, 'Equipos.xlsx')

        self.assertEqual(
            [cell.value for cell in products[1]][:T + 1], import_exports.PRODUCT_TEMPLATE_HEADERS + ['Estado'])
        self.assertEqual([(row[2], row[T]) for row in products.iter_rows(min_row=2, values_only=True)], [
            ('Batería vieja', 'Inactivo'), ('Cable', 'Activo'), ('Funda', 'Activo'), ('Teléfono X', 'Activo')])
        self.assertEqual([row[2] for row in stock.iter_rows(min_row=2, values_only=True)],
                         ['Batería vieja', 'Cable', 'Funda'])
        self.assertEqual([cell.value for cell in devices[1]], unit_import_services.HEADERS)
        self.assertEqual([row[3] for row in devices.iter_rows(min_row=2, values_only=True)],
                         ['SER-0001', 'SER-0002', 'SER-0003'])
        # The filter is honoured in the three files.
        asked = self.archive(is_active='true')
        self.assertEqual([row[2] for row in self.page(asked, 'Productos.xlsx').iter_rows(min_row=2, values_only=True)],
                         ['Cable', 'Funda', 'Teléfono X'])
        self.assertEqual([row[2] for row in self.page(asked, 'Stock.xlsx').iter_rows(min_row=2, values_only=True)],
                         ['Cable', 'Funda'])
        only = self.archive(is_active='false')
        self.assertEqual([row[2] for row in self.page(only, 'Productos.xlsx').iter_rows(min_row=2, values_only=True)],
                         ['Batería vieja'])

    def test_the_note_says_what_each_file_is_the_order_and_that_the_same_shop_does_not_change(self):
        note = self.archive().read('LEEME.txt').decode('utf-8')
        for name in ('Productos.xlsx', 'Stock.xlsx', 'Equipos.xlsx'):
            self.assertIn(name, note)
        self.assertLess(note.index('1. Productos.xlsx'), note.index('2. Stock.xlsx'))
        self.assertLess(note.index('2. Stock.xlsx'), note.index('3. Equipos.xlsx'))
        self.assertIn('tienda vacía', note)
        self.assertIn('MISMA tienda no cambia nada', note)
        # Review N15, N16: what «nothing changes» depends on, and what the files do not cover.
        self.assertIn('si el stock se movió desde la descarga', note)
        self.assertIn('vuelve a dejar cada cantidad como dice el archivo', note)
        self.assertIn('sucursales activas', note)
        # D3: active and inactive travel, and a load never activates or deactivates what exists.
        self.assertIn('productos activos e inactivos', note)
        self.assertIn('«Estado»', note)
        self.assertIn('no activa ni desactiva', note)
        self.assertNotIn('Sólo incluye productos ACTIVOS', note)
        self.assertIn('«Stock total» del reporte de productos', note)
        self.assertIn('con número de serie', note)
        # It names no shop: the note is the platform's, the data is the tenant's.
        self.assertNotIn('Tienda A', note)

    def test_an_archive_over_the_size_cap_is_refused(self):
        with patch.object(inventory_exports, 'MAX_ZIP_BYTES', 200):
            response = self.get('full')
        self.assertEqual(response.status_code, 400, response.content)
        self.assertEqual(response.json()['code'], 'too_large')
        self.assertEqual(self.exported().count(), 0)


# =============================================================================
# The files themselves
# =============================================================================

class FileSafetyTest(_Base):
    """
    Owner's decision D7 (2026-10-13): no visible apostrophe. A text that starts
    like a formula is written as what it is — a STRING cell, marked as text
    (`quotePrefix`) — with the very text it had. What makes it safe is the cell
    type, not a character in front: a string cell is never evaluated.
    """

    DANGEROUS = ['=HYPERLINK("http://malo.example","clic")', '+SUM(1,2)', '-2+3', '@SUM(1)', "'=YA-CON-COMILLA",
                 '=1+1', '-10% dto', '+51 999', '@casa']

    def setUp(self):
        super().setUp()
        self.risky = [
            self.product(name, f'riesgo-{index}', '10.00', self.accessories, description=f'={index}+1')
            for index, name in enumerate(self.DANGEROUS)
        ]
        for index, product in enumerate(self.risky):
            self.code(product, f'RIESGO-{index}')

    def cells(self, response):
        page = load_workbook(io.BytesIO(response.content)).active
        return [cell for row in page.iter_rows(min_row=2) for cell in row if cell.value is not None]

    def test_no_text_cell_can_run_as_a_formula_and_none_carries_an_apostrophe(self):
        for kind, mode in (('products', 'reimport'), ('products', 'report'), ('stock', 'reimport'), ('stock', 'report')):
            with self.subTest(kind=kind, mode=mode):
                cells = self.cells(self.get(kind, mode=mode))
                self.assertEqual([cell.coordinate for cell in cells if cell.data_type == 'f'], [])
                by_text = {cell.value: cell for cell in cells if isinstance(cell.value, str)}
                for name in self.DANGEROUS:
                    self.assertIn(name, by_text)                       # the text itself, nothing in front
                    self.assertEqual(by_text[name].data_type, 's', name)
                for name in self.DANGEROUS[:4]:
                    self.assertTrue(by_text[name].quotePrefix, name)   # «this is text», the way Excel marks it
                    self.assertNotIn("'" + name, by_text)
                self.assertFalse(by_text['Funda'].quotePrefix)         # ordinary text is not marked

    def test_the_sheet_itself_stores_them_as_strings_not_as_formulas(self):
        """Read from the file's own XML, not through a library's idea of it."""
        for kind, mode in (('products', 'reimport'), ('products', 'report'), ('stock', 'reimport')):
            with self.subTest(kind=kind, mode=mode):
                archive = zipfile.ZipFile(io.BytesIO(self.get(kind, mode=mode).content))
                sheets = [name for name in archive.namelist() if name.startswith('xl/worksheets/sheet')]
                self.assertEqual(len(sheets), 1)
                xml = archive.read(sheets[0]).decode('utf-8')
                self.assertNotIn('<f>', xml)
                self.assertNotIn('<f ', xml)
                self.assertIn('t="inlineStr"><is><t>=1+1</t></is></c>', xml)
                self.assertIn('<t>-10% dto</t>', xml)
                # Nothing put in front. (`'=YA-CON-COMILLA` is a name that HAS an apostrophe: it keeps its one.)
                self.assertNotIn("<t>'=1+1", xml)
                self.assertNotIn("<t>'-10% dto", xml)
                self.assertIn("<t>'=YA-CON-COMILLA</t>", xml)
                self.assertNotIn("<t>''=YA-CON-COMILLA", xml)

    def test_a_branch_a_category_and_a_note_are_text_too(self):
        self.centro.name = '+Centro'
        self.centro.save(update_fields=['name'])
        Category.objects.filter(pk=self.phones.pk).update(name='=Teléfonos')
        StockUnit.objects.filter(serial_number='SER-0001').update(notes='@nota')

        report = self.cells(self.get('units', mode='report'))
        self.assertEqual([cell.coordinate for cell in report if cell.data_type == 'f'], [])
        self.assertIn('+Centro', [cell.value for cell in report])
        self.assertIn('@nota', [cell.value for cell in report])
        catalogue = self.cells(self.get('products'))
        self.assertEqual([cell.coordinate for cell in catalogue if cell.data_type == 'f'], [])
        self.assertIn('=Teléfonos', [cell.value for cell in catalogue])
        # A header carries a branch name: it is a string as well.
        stock = load_workbook(io.BytesIO(self.get('stock').content)).active
        self.assertEqual([cell.data_type for cell in stock[1]], ['s'] * stock.max_column)

    def test_a_description_that_is_a_list_with_dashes_is_shown_as_it_was_written(self):
        """Review N14: «- Incluye cable» came out as «'- Incluye cable» in every report."""
        Product.objects.filter(pk=self.case.pk).update(description='- Incluye cable\n- Garantía 6 meses')
        for mode in ('report', 'reimport'):
            with self.subTest(mode):
                row = next(row for row in self.records(self.get('products', mode=mode)) if row['Nombre'] == 'Funda')
                self.assertEqual(row['Descripción'], '- Incluye cable\n- Garantía 6 meses')

    def test_reading_takes_nothing_off_a_text_that_begins_with_an_apostrophe(self):
        """
        Review N8. No workbook of ours carries a guard apostrophe any more, so
        the reader removes none: a stored code like `'-1234` is read as it is.
        """
        from openpyxl import Workbook

        self.assertFalse(hasattr(xlsx_reader, 'unguard_text'))
        book = Workbook()
        page = book.active
        page.title = 'Productos'
        page.append(import_exports.PRODUCT_TEMPLATE_HEADERS)
        page.append(['', "'-1234", "'=Nombre con comilla", '', 10, '', '', '', ''])
        out = io.BytesIO()
        book.save(out)
        body = self.api.post(
            '/api/admin/products/import/preview/', {'file': upload(out.getvalue())}, format='multipart').json()
        self.assertEqual(body['counts']['error'], 0, body['rows'])
        self.assertEqual((body['rows'][0]['data']['code'], body['rows'][0]['data']['name']),
                         ("'-1234", "'=Nombre con comilla"))

    def test_reloading_the_catalogue_reads_the_same_names_and_changes_nothing(self):
        Product.objects.filter(pk=self.case.pk).update(description='- Incluye cable\n- Garantía 6 meses')
        before = list(Product.objects.filter(company=self.company).order_by('pk').values())

        staged = self.api.post(
            '/api/admin/products/import/preview/', {'file': upload(self.get('products').content)}, format='multipart')

        self.assertEqual(staged.status_code, 201, staged.content)
        body = staged.json()
        self.assertEqual((body['counts']['error'], body['counts']['create']), (0, 0), body['rows'])
        read_back = {row['data']['name']: row['data'] for row in body['rows']}
        for name in self.DANGEROUS:
            self.assertIn(name, read_back)
        self.assertEqual(read_back['Funda']['description'], '- Incluye cable\n- Garantía 6 meses')
        self.assertEqual(read_back['=1+1']['description'], '=5+1')
        applied = self.api.post(f'/api/admin/products/import/{body["id"]}/apply/')
        self.assertEqual(applied.json()['summary']['applied']['updated'], 0, applied.content)
        self.assertEqual(list(Product.objects.filter(company=self.company).order_by('pk').values()), before)

    def test_a_code_that_starts_like_a_formula_or_with_an_apostrophe_comes_back_as_itself(self):
        odd = self.product('Rareza', 'rareza', '10.00', self.accessories)
        self.code(odd, "'-1234")
        minus = self.product('Negativo', 'negativo', '10.00', self.accessories)
        self.code(minus, '-9988')

        body = self.api.post(
            '/api/admin/products/import/preview/', {'file': upload(self.get('products').content)},
            format='multipart').json()

        self.assertEqual((body['counts']['error'], body['counts']['create']), (0, 0), body['rows'])
        by_name = {row['data']['name']: row['data'] for row in body['rows']}
        self.assertEqual((by_name['Rareza']['code'], by_name['Rareza']['product_id']), ("'-1234", odd.pk))
        self.assertEqual((by_name['Negativo']['code'], by_name['Negativo']['product_id']), ('-9988', minus.pk))

    def test_reloading_the_stock_finds_every_product(self):
        staged = self.api.post('/api/admin/inventory/import/preview/', {
            'file': upload(self.get('stock').content), 'mapping': StockExportTest.MAPPING, 'header_row': 1,
            'branch_map': json.dumps({'3': self.centro.pk, '4': self.norte.pk}),
        }, format='multipart')
        self.assertEqual(staged.status_code, 201, staged.content)
        self.assertEqual((staged.json()['counts']['error'], staged.json()['counts']['update']), (0, 0))

    def test_a_real_formula_in_an_upload_is_still_refused(self):
        """Leer texto como texto no abre la puerta: una celda que ES una fórmula sigue siendo un error de la fila."""
        from openpyxl import Workbook

        book = Workbook()
        page = book.active
        page.title = 'Productos'
        page.append(import_exports.PRODUCT_TEMPLATE_HEADERS)
        page.append(['', 'NUEVO-1', '=1+1', '', 10, '', '', '', ''])
        out = io.BytesIO()
        book.save(out)
        staged = self.api.post('/api/admin/products/import/preview/', {'file': upload(out.getvalue())}, format='multipart')
        self.assertEqual(staged.json()['counts']['error'], 1, staged.content)

    def test_the_error_report_of_a_load_cannot_run_as_a_formula_either(self):
        """Review O1: errors.csv is opened in a spreadsheet, and it wrote what the upload carried."""
        import csv

        from store.models import BulkImportRow

        job = BulkImportJob.objects.create(company=self.company, import_type=BulkImportJob.PRODUCTS, rows_error=3)
        BulkImportRow.objects.create(
            job=job, sheet_name='=Hoja', row_number=2, action=BulkImportRow.ERROR, match_key='=HYPERLINK("http://malo.example","x")',
            normalized_data={'name': '+SUM(1,2)'}, errors=['@SUM(A1) no es un precio.'], warnings=['-2+3'])
        BulkImportRow.objects.create(
            job=job, sheet_name='Productos', row_number=3, action=BulkImportRow.ERROR, match_key='COD-1',
            normalized_data={'name': 'Funda'}, errors=['El precio no es un número.'], warnings=[])
        BulkImportRow.objects.create(
            job=job, sheet_name='Productos', row_number=4, action=BulkImportRow.ERROR, match_key='\tTAB',
            normalized_data={'name': "'=YA"}, errors=['x'], warnings=[])

        response = self.api.get(f'/api/admin/imports/{job.pk}/errors.csv/')

        self.assertEqual(response.status_code, 200)
        rows = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual(rows[0], ['Hoja', 'Fila', 'Identificador', 'Producto', 'Error', 'Advertencia'])
        self.assertEqual(rows[1], [
            "'=Hoja", '2', '\'=HYPERLINK("http://malo.example","x")', "'+SUM(1,2)", "'@SUM(A1) no es un precio.", "'-2+3"])
        # What is honest stays exactly as it was.
        self.assertEqual(rows[2], ['Productos', '3', 'COD-1', 'Funda', 'El precio no es un número.', ''])
        self.assertEqual(rows[3][2:4], ["'\tTAB", "''=YA"])
        for row in rows[1:]:
            self.assertEqual([cell for cell in row if cell[:1] in ('=', '+', '-', '@', '\t', '\r')], [])
        # A CSV has no cell types: there the apostrophe is the only mark there is, and it stays (D7).

    def test_a_character_a_workbook_cannot_hold_does_not_break_the_download(self):
        Product.objects.filter(pk=self.case.pk).update(name='Funda\x07 rota', description='línea\x0b dos')
        names = [row['Nombre'] for row in self.records(self.get('products'))]
        self.assertIn('Funda rota', names)

    def test_a_branch_name_a_workbook_cannot_hold_does_not_break_the_download(self):
        """Review N10: a control character in a branch name is in a HEADER of the stock file, and it answered 500."""
        from store.models import Branch

        Branch.objects.filter(pk=self.centro.pk).update(name='Cen\x07tro\x0b')
        for kind, mode in (('stock', 'reimport'), ('stock', 'report'), ('units', 'reimport'), ('units', 'report'), ('full', 'reimport')):
            with self.subTest(kind=kind, mode=mode):
                self.assertEqual(self.get(kind, mode=mode).status_code, 200)
        _title, headers, _rows = self.sheet(self.get('stock'))
        self.assertEqual(headers[3:], ['ALMACEN Centro', 'ALMACEN Norte'])
        self.assertEqual({row['Sucursal'] for row in self.records(self.get('units'))}, {'Centro', 'Norte'})

    def test_more_rows_than_the_cap_are_refused_with_the_count(self):
        with patch.object(inventory_exports, 'MAX_EXPORT_ROWS', 2):
            answers = {
                ('products', 'reimport'): self.get('products'),
                ('stock', 'report'): self.get('stock', mode='report'),
                ('units', 'reimport'): self.get('units'),
                ('full', 'reimport'): self.get('full'),
            }
        for key, response in answers.items():
            with self.subTest(key=key):
                self.assertEqual(response.status_code, 400, response.content[:200])
                self.assertEqual((response.json()['code'], response.json()['limit']), ('too_many_rows', 2))
                self.assertGreater(response.json()['count'], 2)
        self.assertEqual(self.exported().count(), 0)

    def test_the_workbooks_are_written_in_write_only_mode(self):
        import openpyxl

        seen = []
        real = openpyxl.Workbook

        def spy(*args, **kwargs):
            seen.append(kwargs.get('write_only', False))
            return real(*args, **kwargs)

        with patch.object(openpyxl, 'Workbook', side_effect=spy):
            for kind, mode in (('products', 'reimport'), ('products', 'report'), ('stock', 'reimport'),
                               ('stock', 'report'), ('units', 'reimport'), ('units', 'report'), ('full', 'reimport')):
                self.assertEqual(self.get(kind, mode=mode).status_code, 200)
        self.assertEqual(seen, [True] * 9)


# =============================================================================
# The download that was already there
# =============================================================================

class LegacyInventoryExportTest(_Base):
    """`GET /api/admin/inventory/export/` answers exactly as it did."""

    URL = '/api/admin/inventory/export/'

    def test_same_sheet_same_columns_same_rows(self):
        response = self.api.get(self.URL)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Disposition'], 'attachment; filename="inventario-actual.xlsx"')
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, ['Inventario'])
        rows = [list(row) for row in book.active.iter_rows(values_only=True)]
        self.assertEqual(rows[0], ['CODIGO', 'CODIGO EAN', 'NOMBRE', 'ALMACEN Centro', 'ALMACEN Norte'])
        # Active products only, the serialized one INCLUDED with its count — as before.
        self.assertEqual(rows[1:], [
            ['00123', None, 'Cable', 0, 7],
            ['FUN-001', '7751234567892', 'Funda', 4, 1],
            ['TEL-X', None, 'Teléfono X', 2, 1],
        ])
        page = book.active
        self.assertEqual((page['A2'].number_format, page['B2'].number_format), ('@', '@'))
        self.assertTrue(page['A1'].font.bold)
        self.assertEqual(page.freeze_panes, 'A2')

    def test_blank_for_counting_and_one_branch(self):
        response = self.api.get(f'{self.URL}?branches={self.norte.pk}&quantities=blank')
        self.assertEqual(response['Content-Disposition'], 'attachment; filename="inventario-para-conteo.xlsx"')
        rows = [list(row) for row in load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True)]
        self.assertEqual(rows[0], ['CODIGO', 'CODIGO EAN', 'NOMBRE', 'ALMACEN Norte'])
        self.assertEqual(rows[1], ['00123', None, 'Cable', None])

    def test_it_still_asks_for_the_capability_to_adjust_and_stays_inside_the_caller_s_branches(self):
        self.assertEqual(self.member('exp_lectora', 'inventory.view').get(self.URL).status_code, 403)
        limited = self.member('exp_centro', 'inventory.adjust', branches=[self.centro])
        self.assertEqual(limited.get(f'{self.URL}?branches={self.norte.pk}').status_code, 404)
        self.assertEqual(limited.get(f'{self.URL}?branches={self.foreign_branch.pk}').status_code, 404)

    def test_a_name_or_a_code_that_starts_like_a_formula_is_written_as_text(self):
        """Review N1: the old download wrote them as live formula cells."""
        evil = self.product('=HYPERLINK("http://malo.example","clic")', 'formula-1', '10.00', self.accessories)
        self.code(evil, '=COD-1')
        self.code(evil, '+775000000001', ProductBarcode.EAN13)
        plus = self.product('+SUM(1,2)', 'formula-2', '10.00', self.accessories)
        _p2d_stock(self.centro, evil, 3)

        response = self.api.get(self.URL)

        page = load_workbook(io.BytesIO(response.content)).active
        cells = [cell for row in page.iter_rows(min_row=2) for cell in row if cell.value is not None]
        self.assertEqual([cell.coordinate for cell in cells if cell.data_type == 'f'], [])
        by_text = {cell.value: cell for cell in cells if isinstance(cell.value, str)}
        # The text itself, as a string cell marked as text: no apostrophe in the value (D7).
        for written in ('=HYPERLINK("http://malo.example","clic")', '=COD-1', '+775000000001', '+SUM(1,2)'):
            self.assertIn(written, by_text)
            self.assertEqual((by_text[written].data_type, by_text[written].quotePrefix), ('s', True), written)
        xml = zipfile.ZipFile(io.BytesIO(response.content)).read('xl/worksheets/sheet1.xml').decode('utf-8')
        self.assertNotIn('<f>', xml)
        self.assertEqual(plus.name, '+SUM(1,2)')                       # the catalogue itself is untouched

    def test_the_download_still_reloads_and_finds_every_product(self):
        evil = self.product('=HYPERLINK("http://malo.example","clic")', 'formula-1', '10.00', self.accessories)
        self.code(evil, '=COD-1')
        _p2d_stock(self.centro, evil, 3)
        Product.objects.filter(pk=self.phone.pk).update(is_active=False)   # a serialized one would be refused: not this test

        staged = self.api.post('/api/admin/inventory/import/preview/', {
            'file': upload(self.api.get(self.URL).content), 'mapping': StockExportTest.MAPPING, 'header_row': 1,
            'branch_map': json.dumps({'3': self.centro.pk, '4': self.norte.pk}),
        }, format='multipart')

        self.assertEqual(staged.status_code, 201, staged.content)
        body = staged.json()
        self.assertEqual((body['counts']['error'], body['counts']['update']), (0, 0), body['rows'])
        self.assertIn(evil.pk, {row['data']['product_id'] for row in body['rows']})

    def test_it_leaves_no_export_audit_line_as_before(self):
        self.api.get(self.URL)
        self.assertEqual(self.exported().count(), 0)
