"""
DEVICE-COMMERCE-01 · DC3 — proveedores y venta bajo pedido.

Lo que estas pruebas fijan:

  * stock 0 NO es vendible: sólo con el modo «Bajo pedido» del producto Y una
    oferta de proveedor vigente, activa, confirmada y con saldo que cubra la
    línea entera (nunca mezcla stock propio y pedido en una línea);
  * la oferta se elige siempre igual (precio comercial más bajo, luego costo
    pactado, luego la que vence antes, luego id);
  * el pago de una línea bajo pedido NO mueve stock: compromete la oferta bajo
    bloqueo; si la oferta venció o se agotó, el pago se registra igual y la
    línea queda «abastecimiento por confirmar»;
  * recibir del proveedor escribe la entrada (al costo congelado) y la salida
    de venta de esa línea: el estante queda igual y el Kardex lo cuenta;
  * las instantáneas (origen, oferta, plazo, costo) no cambian después;
  * todo es por empresa y el costo sólo lo ve `costs.view`.
"""
from datetime import date, timedelta
from decimal import Decimal
from types import SimpleNamespace

from django.core.cache import cache
from django.db import transaction
from django.test import TestCase, TransactionTestCase
from django.utils import timezone

from store import checkout_services as checkout
from store import inventory_services as inventory
from store import supply_services as supply
from store.models import (
    AdminAuditLog, CartItem, Order, OrderItem, PaymentTransaction, Product, StockMovement, StockUnit,
    Supplier, SupplierOffer,
)
from store.tests import _p2d_member, _p3_company, _prod

D = Decimal


def today():
    return timezone.localdate()


class SupplyShop:
    """A shop with an accessory sold on order, a supplier and an offer for it."""

    def build(self, slug='dc3-tienda'):
        cache.clear()
        self.company = _p3_company(slug, 'Tienda DC3')
        self.branch = self.company.default_inventory_branch
        self.host = f'{slug}.tienda.pe'
        self.staff, _ = _p2d_member(self.company, f'{slug}_staff', [
            'company.view', 'products.view', 'products.manage', 'inventory.view', 'inventory.adjust',
            'sales.orders.view', 'sales.orders.manage', 'costs.view'])
        self.case = _prod(self.company, 'Cargador 20W', f'{slug}-cargador', price='120.00', inventory=0)
        Product.objects.filter(pk=self.case.pk).update(sale_mode=Product.SALE_MODE_STOCK_OR_ORDER)
        self.case.refresh_from_db()
        self.supplier = supply.create_supplier(company=self.company, name='Mayorista Lima', actor=self.staff)
        self.offer = self.make_offer()

    def make_offer(self, product=None, **extra):
        values = dict(agreed_cost='70.00', commercial_price='110.00', confirmed_quantity=5,
                      valid_from=today() - timedelta(days=1), valid_until=today() + timedelta(days=6),
                      lead_time_days=4)
        values.update(extra)
        return supply.create_offer(company=self.company, supplier_id=self.supplier.pk,
                                   product_id=(product or self.case).pk, actor=self.staff,
                                   may_view_costs=True, **values)

    def stock(self, product, n):
        inventory.create_stock_movement(branch=self.branch, product_id=product.pk,
                                        movement_type=StockMovement.PURCHASE_ENTRY, quantity=n, reason='Compra',
                                        unit_cost=D('60'))

    def shelf(self, product):
        return inventory.branch_quantity(self.branch, product)

    def add_on_order(self, product=None, quantity=1, session='carrito-a'):
        return CartItem.objects.create(session_key=session, product=product or self.case, quantity=quantity,
                                       supply_origin=CartItem.SUPPLY_SUPPLIER)

    def checkout(self, session='carrito-a'):
        items = list(CartItem.objects.filter(session_key=session).select_related('product'))
        lines = checkout.lines_from_cart(self.company, self.branch, session, items)
        subtotal = checkout.validate_lines_and_subtotal(self.branch, lines)
        pricing = checkout.price_checkout(self.company, subtotal)
        return checkout.create_pending_order(
            company=self.company, branch=self.branch, lines=lines, pricing=pricing,
            details=checkout.CustomerDetails(
                name='Ana', email='ana@example.com', phone='987654321', document_type='DNI',
                document_number='12345678', delivery_method='pickup', receipt_type='boleta',
                accepted_terms=True, accepted_warranty_policy=True),
            cart_session_key=session,
        )

    def pay(self, order):
        """The real `_confirm`, with a gateway answer that passed its checks."""
        from store.views import _SignedNotificationMixin

        attempt = PaymentTransaction.objects.create(
            order=order, provider='izipay', transaction_id=f'T{order.pk:019d}', amount=order.total, currency='PEN')
        result = SimpleNamespace(signed=True, response_code='00', pay_method='CARD', authorization_code='A1',
                                 reference_number='R1', unique_id=f'U{order.pk}')
        outcome = _SignedNotificationMixin()._confirm(attempt, result)
        order.refresh_from_db()
        return outcome


# ---------------------------------------------------------------------------
# Availability
# ---------------------------------------------------------------------------

class AvailabilityMatrixTest(SupplyShop, TestCase):
    """stock / no stock × mode × offer state."""

    def setUp(self):
        self.build()

    def answer(self, quantity=1):
        return supply.availability(self.case, self.branch, quantity)

    def test_own_stock_always_sells_from_the_shelf(self):
        self.stock(self.case, 3)
        result = self.answer(3)
        self.assertEqual((result.origin, result.offer, result.price), ('own', None, D('120.00')))

    def test_no_stock_with_mode_and_a_valid_offer_is_on_order(self):
        result = self.answer(2)
        self.assertEqual((result.origin, result.offer, result.lead_time_days, result.price),
                         ('supplier', self.offer, 4, D('110.00')))

    def test_no_stock_without_the_mode_is_not_sellable(self):
        Product.objects.filter(pk=self.case.pk).update(sale_mode=Product.SALE_MODE_STOCK)
        self.case.refresh_from_db()
        self.assertIsNone(self.answer().origin)

    def test_offers_that_do_not_count(self):
        cases = {
            'expired': dict(valid_until=today() - timedelta(days=1), valid_from=today() - timedelta(days=9)),
            'not_yet': dict(valid_from=today() + timedelta(days=1)),
            'unconfirmed': dict(confirmed_quantity=None),
            'inactive': dict(is_active=False),
        }
        for label, change in cases.items():
            with self.subTest(label):
                SupplierOffer.objects.filter(pk=self.offer.pk).update(**{
                    'valid_from': today() - timedelta(days=1), 'valid_until': today() + timedelta(days=6),
                    'confirmed_quantity': 5, 'is_active': True, **change})
                self.assertIsNone(self.answer().origin)
        SupplierOffer.objects.filter(pk=self.offer.pk).update(
            valid_from=today() - timedelta(days=1), valid_until=today(), confirmed_quantity=5, is_active=True)
        self.assertEqual(self.answer().origin, 'supplier', 'valid through its last day')

    def test_an_exhausted_offer_or_one_that_does_not_cover_the_line_does_not_count(self):
        SupplierOffer.objects.filter(pk=self.offer.pk).update(committed_quantity=5)
        self.assertIsNone(self.answer().origin)
        SupplierOffer.objects.filter(pk=self.offer.pk).update(committed_quantity=3)
        self.assertEqual(self.answer(2).origin, 'supplier')
        self.assertIsNone(self.answer(3).origin)

    def test_an_inactive_supplier_sells_nothing(self):
        Supplier.objects.filter(pk=self.supplier.pk).update(is_active=False)
        self.assertIsNone(self.answer().origin)

    def test_partial_own_stock_is_never_mixed_the_whole_line_goes_on_order(self):
        self.stock(self.case, 1)
        result = self.answer(3)
        self.assertEqual(result.origin, 'supplier')
        self.assertEqual(result.own_stock, 1)

    def test_the_offer_is_chosen_the_same_way_every_time(self):
        # Distinct windows (one offer per key); `valid_from` plays no part in the choice.
        cheap = self.make_offer(commercial_price='100.00', agreed_cost='80.00', lead_time_days=9,
                                valid_from=today() - timedelta(days=2))
        tie_cheaper_cost = self.make_offer(commercial_price='100.00', agreed_cost='75.00', lead_time_days=7,
                                           valid_from=today() - timedelta(days=3))
        self.assertEqual(self.answer().offer, tie_cheaper_cost)
        SupplierOffer.objects.filter(pk=tie_cheaper_cost.pk).update(agreed_cost=D('80.00'),
                                                                   valid_until=today() + timedelta(days=9))
        self.assertEqual(self.answer().offer, cheap)          # same price and cost: the one that ends first
        SupplierOffer.objects.filter(pk=tie_cheaper_cost.pk).update(valid_until=today() + timedelta(days=6))
        self.assertEqual(self.answer().offer, cheap)          # all equal: the lowest id

    def test_another_companys_offer_never_counts(self):
        other = _p3_company('dc3-otra', 'Otra')
        foreign = _prod(other, 'Cargador 20W', 'dc3-otra-cargador', price='120.00', inventory=0)
        with self.assertRaises(supply.SupplyError):
            supply.create_offer(company=other, supplier_id=self.supplier.pk, product_id=foreign.pk,
                                agreed_cost='1', commercial_price='1', confirmed_quantity=1,
                                valid_from=today(), valid_until=today(), lead_time_days=1, actor=None)
        with self.assertRaises(supply.SupplyError):
            supply.create_offer(company=self.company, supplier_id=self.supplier.pk, product_id=foreign.pk,
                                agreed_cost='1', commercial_price='1', confirmed_quantity=1,
                                valid_from=today(), valid_until=today(), lead_time_days=1, actor=None)


class OfferRulesTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()

    def test_bad_offers_are_refused(self):
        bad = [
            dict(valid_until=today() - timedelta(days=2)),            # ends before it starts
            dict(commercial_price='0'),
            dict(commercial_price='-1'),
            dict(agreed_cost='-5'),
            dict(confirmed_quantity=-1),
            dict(lead_time_days=-1),
            dict(lead_time_days=400),
        ]
        for change in bad:
            # Its own window: a duplicate key (DC3 review P2) must not be what refuses it.
            with self.subTest(change), self.assertRaises(supply.SupplyError) as caught:
                self.make_offer(**{'valid_until': today() + timedelta(days=7), **change})
            self.assertNotEqual(caught.exception.code, 'duplicate')

    def test_without_costs_view_the_cost_is_not_written(self):
        offer = supply.create_offer(
            company=self.company, supplier_id=self.supplier.pk, product_id=self.case.pk, agreed_cost='50.00',
            commercial_price='100.00', confirmed_quantity=1, valid_from=today(), valid_until=today(),
            lead_time_days=2, actor=self.staff, may_view_costs=False)
        self.assertIsNone(offer.agreed_cost)

    def test_the_confirmed_quantity_cannot_drop_below_what_was_sold(self):
        SupplierOffer.objects.filter(pk=self.offer.pk).update(committed_quantity=3)
        self.offer.refresh_from_db()
        with self.assertRaises(supply.SupplyError):
            supply.update_offer(self.offer, actor=self.staff, confirmed_quantity=2, may_view_costs=True)
        supply.update_offer(self.offer, actor=self.staff, confirmed_quantity=3, may_view_costs=True)

    def test_history_is_audited_and_a_used_offer_is_never_deleted(self):
        supply.update_offer(self.offer, actor=self.staff, commercial_price='105.00', may_view_costs=True)
        supply.deactivate_offer(self.offer, actor=self.staff, reason='Proveedor sin stock')
        actions = list(AdminAuditLog.objects.filter(target_type='supplier_offer', target_id=str(self.offer.pk))
                       .order_by('pk').values_list('action', flat=True))
        self.assertEqual(actions, ['supplier_offer_created', 'supplier_offer_updated', 'supplier_offer_deactivated'])
        self.offer.refresh_from_db()
        self.assertFalse(self.offer.is_active)

    def test_a_used_variant_is_never_sold_on_order(self):
        Product.objects.filter(pk=self.case.pk).update(condition_line=Product.CONDITION_USED)
        self.case.refresh_from_db()
        with self.assertRaises(supply.SupplyError):
            supply.set_sale_mode(self.case, Product.SALE_MODE_STOCK_OR_ORDER, actor=self.staff)


# ---------------------------------------------------------------------------
# Checkout, payment, reception
# ---------------------------------------------------------------------------

class OnOrderSaleTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()

    def test_checkout_freezes_origin_offer_lead_time_cost_and_price_without_committing(self):
        self.add_on_order(quantity=2)
        order = self.checkout()
        [line] = order.items.all()
        self.assertEqual((line.supply_origin, line.supplier_offer_id, line.lead_time_days),
                         ('supplier', self.offer.pk, 4))
        self.assertEqual((line.price, line.unit_cost, line.cost_source), (D('110.00'), D('70.0000'), 'supplier_offer'))
        self.assertEqual(order.total, D('220.00'))
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 0)        # a checkout commits nothing

    def test_checkout_refuses_an_on_order_line_without_a_valid_offer(self):
        self.add_on_order(quantity=6)                             # the offer has 5
        with self.assertRaises(checkout.CheckoutError):
            self.checkout()
        CartItem.objects.all().delete()
        self.add_on_order(quantity=1)
        SupplierOffer.objects.filter(pk=self.offer.pk).update(is_active=False)
        with self.assertRaises(checkout.CheckoutError):
            self.checkout()
        self.assertFalse(Order.objects.exists())

    def test_a_shelf_line_without_stock_is_refused_even_with_an_offer(self):
        """Stock 0 does not mean sellable: the buyer has to take the line on order."""
        CartItem.objects.create(session_key='carrito-a', product=self.case, quantity=1)
        with self.assertRaises(checkout.CheckoutError):
            self.checkout()

    def test_payment_commits_the_offer_and_moves_no_stock(self):
        self.add_on_order(quantity=2)
        order = self.checkout()
        self.assertEqual(self.pay(order), 'paid')
        self.assertEqual(order.status, Order.Status.PAID)
        [line] = order.items.all()
        self.assertEqual(line.supply_status, 'pending')
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 2)
        self.assertFalse(StockMovement.objects.filter(product=self.case).exists())
        self.assertEqual(self.shelf(self.case), 0)
        # Not a shortfall: nothing was owed from the shelf.
        self.assertEqual(inventory.order_stock_shortfall(order), [])
        # A replayed notification commits nothing twice.
        from store.views import _SignedNotificationMixin
        attempt = order.payment_transactions.get()
        _SignedNotificationMixin()._confirm(attempt, SimpleNamespace(
            signed=True, response_code='00', pay_method='CARD', authorization_code='A1', reference_number='R1',
            unique_id='U'))
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 2)

    def test_an_offer_that_expired_before_the_payment_keeps_the_money_and_asks_to_confirm(self):
        self.add_on_order(quantity=1)
        order = self.checkout()
        SupplierOffer.objects.filter(pk=self.offer.pk).update(valid_until=today() - timedelta(days=1),
                                                              valid_from=today() - timedelta(days=5))
        self.assertEqual(self.pay(order), 'paid')
        self.assertTrue(order.paid)
        [line] = order.items.all()
        self.assertEqual(line.supply_status, 'to_confirm')
        self.assertIn('abastecimiento por confirmar', order.payment_error.lower())
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 0)
        self.assertTrue(AdminAuditLog.objects.filter(action='order_supply_to_confirm', target_id=str(order.pk)).exists())
        self.assertFalse(StockMovement.objects.filter(product=self.case).exists())

    def test_an_exhausted_offer_at_payment_asks_to_confirm(self):
        self.add_on_order(quantity=2)
        order = self.checkout()
        SupplierOffer.objects.filter(pk=self.offer.pk).update(committed_quantity=4)
        self.pay(order)
        self.assertEqual(order.items.get().supply_status, 'to_confirm')
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 4)

    def test_staff_reconfirm_a_line_against_a_valid_offer(self):
        self.add_on_order(quantity=1)
        order = self.checkout()
        SupplierOffer.objects.filter(pk=self.offer.pk).update(is_active=False)
        self.pay(order)
        line = order.items.get()
        fresh = self.make_offer(agreed_cost='72.00', commercial_price='115.00', lead_time_days=6,
                                valid_from=today() - timedelta(days=2))
        supply.confirm_line(line, actor=self.staff)
        line.refresh_from_db()
        fresh.refresh_from_db()
        self.assertEqual((line.supply_status, line.supplier_offer_id, line.unit_cost, line.lead_time_days),
                         ('pending', fresh.pk, D('72.0000'), 6))
        self.assertEqual(line.price, D('110.00'), 'the buyer paid what was paid')
        self.assertEqual(fresh.committed_quantity, 1)

    def test_a_line_already_committed_is_not_confirmed_again(self):
        self.add_on_order(quantity=2)
        order = self.checkout()
        self.pay(order)
        with self.assertRaises(supply.SupplyError):
            supply.confirm_line(order.items.get(), actor=self.staff)
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 2)

    def test_checkout_validation_itself_refuses_an_offer_that_no_longer_covers(self):
        """Not only the cart reader: the validator any surface calls asks the offer again."""
        line = checkout.CheckoutLine(product=self.case, quantity=2, offer=self.offer)
        self.assertEqual(checkout.validate_lines_and_subtotal(self.branch, [line]), D('220.00'))
        SupplierOffer.objects.filter(pk=self.offer.pk).update(committed_quantity=4)
        with self.assertRaises(checkout.CheckoutError):
            checkout.validate_lines_and_subtotal(self.branch, [line])

    def test_receiving_writes_the_entry_at_the_frozen_cost_and_the_sale_exit(self):
        self.stock(self.case, 2)                      # the shelf has other units of its own
        CartItem.objects.all().delete()
        Product.objects.filter(pk=self.case.pk).update()
        self.add_on_order(quantity=3)                 # 2 on the shelf do not cover 3: the line is on order
        order = self.checkout()
        self.pay(order)
        SupplierOffer.objects.filter(pk=self.offer.pk).update(agreed_cost=D('999'))   # the offer changes later
        line = order.items.get()
        supply.receive_line(line, actor=self.staff)
        line.refresh_from_db()
        self.assertEqual((line.supply_status, line.unit_cost), ('received', D('70.0000')))
        entry = StockMovement.objects.get(product=self.case, order=order, movement_type='purchase_entry')
        exit_ = StockMovement.objects.get(product=self.case, order=order, movement_type='sale_exit')
        self.assertEqual((entry.quantity, entry.unit_cost, exit_.quantity), (3, D('70.0000'), 3))
        self.assertEqual(self.shelf(self.case), 2, 'net zero on the shelf')
        self.assertEqual(inventory.order_stock_shortfall(order), [])
        with self.assertRaises(supply.SupplyError):
            supply.receive_line(line, actor=self.staff)            # once
        # A reprocess of the order's exits does not sell the line again from the shelf.
        inventory.reprocess_order_stock_exit(order, actor=self.staff)
        self.assertEqual(StockMovement.objects.filter(order=order, movement_type='sale_exit').count(), 1)
        self.assertEqual(self.shelf(self.case), 2)

    def test_a_failed_supply_releases_the_offer_and_flags_the_order(self):
        self.add_on_order(quantity=2)
        order = self.checkout()
        self.pay(order)
        line = order.items.get()
        with self.assertRaises(supply.SupplyError):
            supply.fail_line(line, actor=self.staff, reason='')
        supply.fail_line(line, actor=self.staff, reason='El proveedor no tiene stock')
        line.refresh_from_db()
        order.refresh_from_db()
        self.assertEqual(line.supply_status, 'failed')
        self.assertIn('El proveedor no tiene stock', line.supply_note)
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 0)
        self.assertEqual(order.status, Order.Status.PAID, 'the refund goes through the existing flow')
        self.assertIn('reembols', order.payment_error.lower())
        with self.assertRaises(supply.SupplyError):
            supply.receive_line(line, actor=self.staff)
        self.assertFalse(StockMovement.objects.filter(product=self.case).exists())

    def test_the_snapshot_does_not_follow_the_offer(self):
        self.add_on_order(quantity=1)
        order = self.checkout()
        self.pay(order)
        supply.update_offer(self.offer, actor=self.staff, agreed_cost='10.00', commercial_price='10.00',
                            lead_time_days=30, may_view_costs=True)
        line = order.items.get()
        self.assertEqual((line.unit_cost, line.price, line.lead_time_days), (D('70.0000'), D('110.00'), 4))

    def test_a_line_of_another_company_cannot_be_received_by_this_one(self):
        self.add_on_order(quantity=1)
        order = self.checkout()
        self.pay(order)
        other = _p3_company('dc3-ajena', 'Ajena')
        stranger, _ = _p2d_member(other, 'dc3_ajeno', ['inventory.adjust', 'sales.orders.manage'])
        with self.assertRaises(supply.SupplyError):
            supply.receive_line(order.items.get(), actor=stranger, company=other)


class MarginWithSupplierCostTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()

    def test_the_margin_counts_the_frozen_supplier_cost(self):
        from rest_framework.test import APIClient
        self.add_on_order(quantity=2)
        order = self.checkout()
        self.pay(order)
        Order.objects.filter(pk=order.pk).update(paid_at=timezone.now())
        client = APIClient()
        admin, _ = _p2d_member(self.company, 'dc3_margen', ['company.view', 'sales.analytics.view', 'costs.view'])
        client.force_authenticate(user=admin)
        res = client.get('/api/admin/sales/margin/?days=7')
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(D(body['cost_of_goods_sold']), D('140.00'))
        self.assertEqual(body['pending_cost']['lines'], 0)


class LastUnitRaceTest(TransactionTestCase):
    """Two payments at the same instant for the last confirmed unit: one commits, the other asks."""

    def test_one_commits_and_the_other_is_kept_to_confirm(self):
        from threading import Barrier, Thread

        from django.db import connection
        from store.views import _SignedNotificationMixin

        shop = SupplyShop()
        shop.build('dc3-carrera')
        SupplierOffer.objects.filter(pk=shop.offer.pk).update(confirmed_quantity=1)
        orders = []
        for session in ('carrito-a', 'carrito-b'):
            shop.add_on_order(quantity=1, session=session)
            orders.append(shop.checkout(session))
        attempts = [
            PaymentTransaction.objects.create(order=order, provider='izipay', transaction_id=f'R{order.pk:019d}',
                                              amount=order.total, currency='PEN')
            for order in orders
        ]
        barrier, outcomes = Barrier(2), []

        def run(attempt):
            try:
                barrier.wait(timeout=5)
                outcomes.append(_SignedNotificationMixin()._confirm(attempt, SimpleNamespace(
                    signed=True, response_code='00', pay_method='CARD', authorization_code='A',
                    reference_number='R', unique_id=f'U{attempt.pk}')))
            finally:
                connection.close()

        threads = [Thread(target=run, args=(attempt,)) for attempt in attempts]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=20)
        self.assertEqual(outcomes, ['paid', 'paid'])
        states = sorted(OrderItem.objects.filter(order__in=orders).values_list('supply_status', flat=True))
        self.assertEqual(states, ['pending', 'to_confirm'])
        shop.offer.refresh_from_db()
        self.assertEqual(shop.offer.committed_quantity, 1)
        self.assertEqual(Order.objects.filter(pk__in=[o.pk for o in orders], paid=True).count(), 2)
        self.assertFalse(StockMovement.objects.filter(product=shop.case).exists())


# ---------------------------------------------------------------------------
# The storefront: catalogue, cart
# ---------------------------------------------------------------------------

from django.test import override_settings  # noqa: E402
from rest_framework.test import APIClient  # noqa: E402

from store.test_device_commerce import DeviceShop  # noqa: E402


@override_settings(ALLOWED_HOSTS=['.tienda.pe', 'testserver'])
class StorefrontOnOrderTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()
        self.public = APIClient(HTTP_HOST=self.host)

    def add(self, quantity=1, on_order=None, session='carrito-a', product=None):
        body = {'session_key': session, 'product': (product or self.case).pk, 'quantity': quantity}
        if on_order is not None:
            body['on_order'] = on_order
        return self.public.post('/api/cart/add/', body, format='json')

    def test_the_catalogue_says_on_order_with_its_lead_time_and_price_never_cost_or_supplier(self):
        res = self.public.get(f'/api/products/?slug={self.case.slug}')
        body = res.json()
        row = (body['results'] if isinstance(body, dict) else body)[0]
        self.assertEqual(row['inventory'], 0)
        self.assertEqual(row['on_order'], {'lead_time_days': 4, 'price': '110.00'})
        text = res.content.decode()
        for secret in ('70.00', 'Mayorista', 'agreed_cost', 'supplier'):
            self.assertNotIn(secret, text)
        SupplierOffer.objects.all().update(is_active=False)
        res = self.public.get(f'/api/products/?slug={self.case.slug}').json()
        self.assertIsNone((res['results'] if isinstance(res, dict) else res)[0]['on_order'])

    def test_stock_zero_is_not_added_unless_the_buyer_takes_it_on_order(self):
        refused = self.add(2)
        self.assertEqual(refused.status_code, 400)
        self.assertEqual(refused.json()['code'], 'on_order_available')
        self.assertEqual(refused.json()['on_order']['lead_time_days'], 4)
        self.assertFalse(CartItem.objects.exists())
        res = self.add(2, on_order=True)
        self.assertEqual(res.status_code, 200, res.content)
        line = res.json()
        self.assertEqual((line['supply']['origin'], line['supply']['lead_time_days'], line['line_total']),
                         ('supplier', 4, '220.00'))
        self.assertEqual(CartItem.objects.get().supply_origin, 'supplier')
        self.assertEqual(self.add(4, on_order=True).json()['code'], 'not_on_order')      # 6 > 5 confirmed
        patched = self.public.patch(f'/api/cart/{CartItem.objects.get().pk}/?session_key=carrito-a',
                                    {'quantity': 5}, format='json')
        self.assertEqual(patched.status_code, 200, patched.content)
        too_many = self.public.patch(f'/api/cart/{CartItem.objects.get().pk}/?session_key=carrito-a',
                                     {'quantity': 6}, format='json')
        self.assertEqual(too_many.status_code, 400)

    def test_without_the_mode_or_an_offer_on_order_is_refused(self):
        Product.objects.filter(pk=self.case.pk).update(sale_mode=Product.SALE_MODE_STOCK)
        res = self.add(1, on_order=True)
        self.assertEqual((res.status_code, res.json()['code']), (409, 'not_on_order'))
        self.assertEqual(self.add(1).json().get('code'), None)       # plain «stock insuficiente»

    def test_a_line_is_never_half_shelf_half_order(self):
        self.stock(self.case, 1)
        self.assertEqual(self.add(1).status_code, 200)
        mixed = self.add(1, on_order=True)
        self.assertEqual((mixed.status_code, mixed.json()['code']), (409, 'mixed_supply'))
        CartItem.objects.all().delete()
        self.assertEqual(self.add(3, on_order=True).status_code, 200)
        self.assertEqual(self.add(1).json()['code'], 'mixed_supply')

    def test_web_checkout_and_payment_of_an_on_order_line(self):
        self.add(2, on_order=True)
        order = self.checkout()
        self.assertEqual(order.items.get().supply_origin, 'supplier')
        self.pay(order)
        self.assertEqual(order.items.get().supply_status, 'pending')

    def test_the_buyers_order_shows_the_lead_time_and_not_the_cost(self):
        from django.contrib.auth.models import User
        buyer = User.objects.create_user(username='dc3_comprador', password='x')
        self.add(1, on_order=True)
        order = self.checkout()
        Order.objects.filter(pk=order.pk).update(user=buyer)
        self.pay(order)
        client = APIClient(HTTP_HOST=self.host)
        client.force_authenticate(user=buyer)
        res = client.get('/api/orders/')
        body = res.json()
        row = (body['results'] if isinstance(body, dict) else body)[0]
        item = row['items'][0]
        self.assertEqual((item['supply_origin'], item['lead_time_days']), ('supplier', 4))
        for secret in ('unit_cost', 'supplier_offer', '70.00', 'Mayorista'):
            self.assertNotIn(secret, res.content.decode())


@override_settings(ALLOWED_HOSTS=['.tienda.pe', 'testserver'])
class DeviceOnOrderTest(DeviceShop, TestCase):
    """A new variant sold by unit, with no device on the shelf, sold on order."""

    def setUp(self):
        self.build('dc3-equipos')
        self.public = APIClient(HTTP_HOST=self.host)
        Product.objects.filter(pk=self.new256.pk).update(sale_mode=Product.SALE_MODE_STOCK_OR_ORDER)
        self.new256.refresh_from_db()
        self.supplier = supply.create_supplier(company=self.company, name='Importadora')
        self.offer = supply.create_offer(
            company=self.company, supplier_id=self.supplier.pk, product_id=self.new256.pk, agreed_cost='2800.00',
            commercial_price='3400.00', confirmed_quantity=2, valid_from=today(),
            valid_until=today() + timedelta(days=7), lead_time_days=10, may_view_costs=True)

    def test_the_model_page_offers_the_variant_on_order(self):
        detail = self.public.get(f'/api/catalog/models/{self.model.slug}/').json()
        row = next(v for v in detail['variants'] if v['product_id'] == self.new256.pk)
        self.assertEqual((row['available'], row['availability']), (True, 'supplier'))
        self.assertEqual(row['on_order'], {'lead_time_days': 10, 'price': '3400.00'})
        self.assertNotIn('2800', str(detail))
        own = next(v for v in detail['variants'] if v['product_id'] == self.new128.pk)
        self.assertEqual((own['available'], own['availability']), (False, 'none'))

    def test_a_device_on_order_holds_nothing_then_is_received_with_its_real_imei_and_sold(self):
        from store.test_device_commerce import IMEIS
        res = self.public.post('/api/cart/add/', {'session_key': 'c', 'product': self.new256.pk, 'on_order': True},
                               format='json')
        self.assertEqual(res.status_code, 200, res.content)
        self.assertFalse(StockUnit.objects.filter(status='reserved').exists())
        listed = self.public.get('/api/cart/?session_key=c').json()
        listed = listed['results'] if isinstance(listed, dict) else listed
        self.assertEqual(len(listed), 1, 'reading the cart keeps an on-order device line')
        # A device on the shelf cannot be held into the same line.
        self.receive(self.new256, 1)
        mixed = self.public.post('/api/cart/add/', {'session_key': 'c', 'product': self.new256.pk}, format='json')
        self.assertEqual(mixed.json()['code'], 'mixed_supply')
        shop = SupplyShop()
        shop.company, shop.branch = self.company, self.branch
        order = shop.checkout('c')
        shop.pay(order)
        line = order.items.get()
        self.assertEqual((line.supply_status, line.stock_unit_id), ('pending', None))
        with self.assertRaises(supply.SupplyError):
            supply.receive_line(line, actor=self.staff, units=[])
        supply.receive_line(line, actor=self.staff, units=[{'serial_number': 'SERDC3001', 'imei': IMEIS[30]}])
        unit = StockUnit.objects.get(serial_number='SERDC3001')
        self.assertEqual((unit.status, unit.order_id, unit.cost), ('sold', order.pk, D('2800.00')))
        self.assertEqual(self.shelf(self.new256), 1, 'the device of the shelf is untouched')
        from store import unit_identification
        self.assertEqual(unit_identification.pending_units(order), {})


# ---------------------------------------------------------------------------
# The panel
# ---------------------------------------------------------------------------

class AdminSupplierApiTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()
        self.client_ = APIClient()
        self.client_.force_authenticate(user=self.staff)
        self.viewer, _ = _p2d_member(self.company, 'dc3_viewer', ['company.view', 'products.view'])
        self.seller, _ = _p2d_member(self.company, 'dc3_manager_sin_costos',
                                     ['company.view', 'products.view', 'products.manage'])

    def as_(self, user):
        client = APIClient()
        client.force_authenticate(user=user)
        return client

    def test_suppliers_are_listed_created_and_edited_by_company(self):
        res = self.client_.post('/api/admin/suppliers/', {'name': 'Distribuidora Sur', 'document_type': 'RUC',
                                                         'document_number': '20100070970'}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        bad_ruc = self.client_.post('/api/admin/suppliers/', {'name': 'X', 'document_type': 'RUC',
                                                             'document_number': '20100070971'}, format='json')
        self.assertEqual(bad_ruc.status_code, 400)
        dup = self.client_.post('/api/admin/suppliers/', {'name': 'distribuidora sur'}, format='json')
        self.assertEqual(dup.status_code, 400)
        pk = res.json()['id']
        self.assertEqual(self.client_.patch(f'/api/admin/suppliers/{pk}/', {'is_active': False},
                                            format='json').json()['is_active'], False)
        names = [row['name'] for row in self.as_(self.viewer).get('/api/admin/suppliers/').json()['results']]
        self.assertEqual(sorted(names), ['Distribuidora Sur', 'Mayorista Lima'])
        self.assertEqual(self.as_(self.viewer).post('/api/admin/suppliers/', {'name': 'Y'},
                                                    format='json').status_code, 403)
        # Another company's supplier does not exist here.
        other = _p3_company('dc3-otra-api', 'Otra')
        foreign = supply.create_supplier(company=other, name='Ajeno')
        self.assertEqual(self.client_.get(f'/api/admin/suppliers/{foreign.pk}/').status_code, 404)
        self.assertEqual(self.client_.patch(f'/api/admin/suppliers/{foreign.pk}/', {'name': 'Mío'},
                                            format='json').status_code, 404)

    def test_offers_show_the_cost_only_to_costs_view(self):
        with_cost = self.client_.get('/api/admin/supplier-offers/').json()['results'][0]
        self.assertEqual(with_cost['agreed_cost'], '70.00')
        res = self.as_(self.seller).get('/api/admin/supplier-offers/')
        self.assertNotIn('agreed_cost', res.json()['results'][0])
        self.assertNotIn('70.00', res.content.decode())
        detail = self.as_(self.seller).get(f'/api/admin/supplier-offers/{self.offer.pk}/')
        self.assertNotIn('agreed_cost', detail.json())
        # Writing a cost without the permission is ignored, not stored.
        made = self.as_(self.seller).post('/api/admin/supplier-offers/', {
            'supplier_id': self.supplier.pk, 'product_id': self.case.pk, 'agreed_cost': '1.00',
            'commercial_price': '100.00', 'confirmed_quantity': 2, 'valid_from': str(today()),
            'valid_until': str(today()), 'lead_time_days': 3}, format='json')
        self.assertEqual(made.status_code, 201, made.content)
        self.assertIsNone(SupplierOffer.objects.get(pk=made.json()['id']).agreed_cost)
        self.as_(self.seller).patch(f'/api/admin/supplier-offers/{self.offer.pk}/', {'agreed_cost': '1.00'},
                                    format='json')
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.agreed_cost, D('70.00'))

    def test_offer_filters_history_and_deactivation(self):
        old = self.make_offer(valid_from=today() - timedelta(days=20), valid_until=today() - timedelta(days=10))
        rows = self.client_.get('/api/admin/supplier-offers/?state=expired').json()['results']
        self.assertEqual([row['id'] for row in rows], [old.pk])
        rows = self.client_.get('/api/admin/supplier-offers/?state=active').json()['results']
        self.assertEqual([row['id'] for row in rows], [self.offer.pk])
        self.assertEqual(rows[0]['state'], 'active')
        rows = self.client_.get(f'/api/admin/supplier-offers/?product={self.case.pk}&supplier={self.supplier.pk}').json()
        self.assertEqual(rows['count'], 2)
        self.client_.patch(f'/api/admin/supplier-offers/{self.offer.pk}/', {'commercial_price': '105.00'},
                           format='json')
        res = self.client_.post(f'/api/admin/supplier-offers/{self.offer.pk}/deactivate/', {'reason': 'Sin stock'},
                                format='json')
        self.assertEqual(res.json()['state'], 'inactive')
        history = self.client_.get(f'/api/admin/supplier-offers/{self.offer.pk}/history/').json()['results']
        self.assertEqual([row['action'] for row in history],
                         ['supplier_offer_deactivated', 'supplier_offer_updated', 'supplier_offer_created'])
        self.assertNotIn('70.00', str(history))

    def test_the_sale_mode_of_a_product(self):
        url = f'/api/admin/products/{self.case.pk}/sale-mode/'
        self.assertEqual(self.as_(self.viewer).patch(url, {'sale_mode': 'stock'}, format='json').status_code, 403)
        res = self.client_.patch(url, {'sale_mode': 'stock'}, format='json')
        self.assertEqual((res.status_code, res.json()['sale_mode']), (200, 'stock'))
        self.assertEqual(self.client_.patch(url, {'sale_mode': 'magic'}, format='json').status_code, 400)
        self.assertEqual(self.client_.get(f'/api/admin/products/{self.case.pk}/').json()['sale_mode'], 'stock')

    def test_the_settlement_summary_reads_received_lines_at_frozen_cost(self):
        self.add_on_order(quantity=2)
        order = self.checkout()
        self.pay(order)
        supply.receive_line(order.items.get(), actor=self.staff)
        res = self.client_.get('/api/admin/suppliers/settlement/')
        self.assertEqual(res.json()['results'], [{'supplier_id': self.supplier.pk, 'supplier_name': 'Mayorista Lima',
                                                  'units': 2, 'amount': '140.00', 'units_without_cost': 0}])
        self.assertEqual(self.as_(self.seller).get('/api/admin/suppliers/settlement/').status_code, 403)


class AdminOrderSupplyApiTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()
        self.client_ = APIClient()
        self.client_.force_authenticate(user=self.staff)
        self.add_on_order(quantity=1)
        self.order = self.checkout()
        SupplierOffer.objects.filter(pk=self.offer.pk).update(is_active=False)
        self.pay(self.order)
        self.line = self.order.items.get()

    def test_the_panel_shows_the_supply_and_filters_by_it(self):
        detail = self.client_.get(f'/api/admin/orders/{self.order.pk}/').json()
        self.assertEqual(detail['supply_status'], 'to_confirm')
        item = detail['items'][0]
        self.assertEqual((item['supply_origin'], item['supply_status'], item['lead_time_days'], item['supplier_name']),
                         ('supplier', 'to_confirm', 4, 'Mayorista Lima'))
        self.assertNotIn('unit_cost', item)
        self.assertIs(item['product_is_serialized'], False)       # the receive form asks for devices when true
        self.assertTrue(detail['can_manage_supply'])
        self.assertTrue(detail['can_receive_supply'])
        rows = self.client_.get('/api/admin/orders/?supply=to_confirm').json()['results']
        self.assertEqual([row['id'] for row in rows], [self.order.pk])
        self.assertEqual(rows[0]['supply_status'], 'to_confirm')
        self.assertEqual(self.client_.get('/api/admin/orders/?supply=pending').json()['results'], [])

    def test_confirm_receive_and_fail_through_the_panel(self):
        url = f'/api/admin/orders/{self.order.pk}/supply/{self.line.pk}/'
        self.assertEqual(self.client_.post(url + 'confirm/', {}, format='json').status_code, 400)   # no valid offer
        fresh = self.make_offer(valid_from=today() - timedelta(days=2))
        res = self.client_.post(url + 'confirm/', {'offer_id': fresh.pk}, format='json')
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()['supply_status'], 'pending')
        clerk, _ = _p2d_member(self.company, 'dc3_clerk', ['company.view', 'sales.orders.view', 'sales.orders.manage'])
        as_clerk = APIClient()
        as_clerk.force_authenticate(user=clerk)
        self.assertEqual(as_clerk.post(url + 'receive/', {}, format='json').status_code, 403)   # inventory.adjust
        res = self.client_.post(url + 'receive/', {}, format='json')
        self.assertEqual((res.status_code, res.json()['supply_status']), (200, 'received'))
        self.assertEqual(self.client_.post(url + 'fail/', {'reason': 'x'}, format='json').status_code, 400)

    def test_another_company_cannot_touch_the_line(self):
        other = _p3_company('dc3-otra-orden', 'Otra')
        stranger, _ = _p2d_member(other, 'dc3_extrano', ['company.view', 'sales.orders.view', 'sales.orders.manage',
                                                          'inventory.adjust'])
        client = APIClient()
        client.force_authenticate(user=stranger)
        url = f'/api/admin/orders/{self.order.pk}/supply/{self.line.pk}/fail/'
        self.assertEqual(client.post(url, {'reason': 'x'}, format='json').status_code, 404)
        self.line.refresh_from_db()
        self.assertEqual(self.line.supply_status, 'to_confirm')


# ---------------------------------------------------------------------------
# The till
# ---------------------------------------------------------------------------

class PosOnOrderTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()
        self.cashier, _ = _p2d_member(self.company, 'dc3_cajero', ['company.view', 'sales.pos.use'])

    def sell(self, items):
        from store.tests import _c1_sale
        order, _created = _c1_sale(actor=self.cashier, company=self.company, branch=self.branch, items=items)
        return order

    def test_a_till_sale_on_order_commits_the_offer_and_hands_nothing_over(self):
        order = self.sell([{'product': self.case.pk, 'quantity': 2, 'on_order': True}])
        self.assertEqual((order.status, order.total), (Order.Status.PAID, D('220.00')))
        self.assertNotEqual(order.fulfillment_status, Order.FulfillmentStatus.DELIVERED)
        line = order.items.get()
        self.assertEqual((line.supply_origin, line.supply_status, line.lead_time_days, line.price, line.unit_cost),
                         ('supplier', 'pending', 4, D('110.00'), D('70.0000')))
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 2)
        self.assertFalse(StockMovement.objects.filter(product=self.case).exists())

    def test_the_till_cannot_sell_on_order_without_an_offer_or_out_of_stock_without_asking(self):
        from store.pos_services import PosValidationError
        with self.assertRaises(inventory.InsufficientStockError):
            self.sell([{'product': self.case.pk, 'quantity': 1}])                    # stock 0, not on order
        with self.assertRaises(PosValidationError):
            self.sell([{'product': self.case.pk, 'quantity': 6, 'on_order': True}])  # offer has 5
        SupplierOffer.objects.filter(pk=self.offer.pk).update(valid_until=today() - timedelta(days=1),
                                                              valid_from=today() - timedelta(days=3))
        with self.assertRaises(PosValidationError):
            self.sell([{'product': self.case.pk, 'quantity': 1, 'on_order': True}])
        self.assertFalse(Order.objects.exists())
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 0)

    def test_an_offer_taken_between_pricing_and_charging_refuses_the_whole_sale(self):
        """Nothing is captured at the till yet: an offer that no longer covers the line refuses the sale."""
        from unittest import mock
        from store.pos_services import PosValidationError
        with mock.patch.object(supply, '_offer_is_valid', return_value=False), \
                self.assertRaises(PosValidationError):
            self.sell([{'product': self.case.pk, 'quantity': 1, 'on_order': True}])
        self.assertFalse(Order.objects.exists())
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 0)

    def test_one_product_is_never_half_shelf_half_order_at_the_till(self):
        from store.pos_services import PosValidationError
        self.stock(self.case, 1)
        with self.assertRaises(PosValidationError):
            self.sell([{'product': self.case.pk, 'quantity': 1},
                       {'product': self.case.pk, 'quantity': 1, 'on_order': True}])

    def test_the_till_shows_lead_time_and_on_order_price(self):
        from store import pos_payloads, pos_services
        self.assertEqual(pos_payloads.product_payload(self.case, self.branch)['on_order'],
                         {'lead_time_days': 4, 'price': '110.00'})
        priced = pos_services.build_pos_sale(
            operator=self.cashier, company=self.company, branch=self.branch,
            items=[{'product': self.case.pk, 'quantity': 1, 'on_order': True}], validate_cash=False)
        line = pos_payloads.preview_payload(priced, may_see_commission=False)['lines'][0]
        self.assertEqual((line['price'], line['on_order'], line['lead_time_days']), ('110.00', True, 4))


# ---------------------------------------------------------------------------
# The weekly Excel
# ---------------------------------------------------------------------------

OFFER_HEADERS = ['Proveedor', 'Código/SKU del producto', 'Modelo', 'Capacidad', 'Color', 'Línea', 'Costo pactado',
                 'Precio comercial', 'Cantidad confirmada', 'Vigente desde', 'Vigente hasta', 'Plazo (días)']


class OfferExcelTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()
        from store.models import ProductBarcode  # noqa: F401
        self.client_ = APIClient()
        self.client_.force_authenticate(user=self.staff)
        self.start, self.end = today(), today() + timedelta(days=6)

    def file(self, rows, headers=OFFER_HEADERS):
        from store.test_cost_imports import upload
        return upload(headers, rows, title='Ofertas de proveedor')

    def row(self, **change):
        values = {'Proveedor': 'Mayorista Lima', 'Código/SKU del producto': self.case.slug, 'Modelo': '',
                  'Capacidad': '', 'Color': '', 'Línea': '', 'Costo pactado': '68', 'Precio comercial': '109.90',
                  'Cantidad confirmada': 8, 'Vigente desde': self.start.isoformat(),
                  'Vigente hasta': self.end.isoformat(), 'Plazo (días)': 5}
        values.update(change)
        return [values[h] for h in OFFER_HEADERS]

    def preview(self, rows, client=None, headers=OFFER_HEADERS):
        return (client or self.client_).post('/api/admin/supplier-offers/import/preview/',
                                             {'file': self.file(rows, headers)}, format='multipart')

    def apply(self, job_id, client=None):
        return (client or self.client_).post(f'/api/admin/supplier-offers/import/{job_id}/apply/', {}, format='json')

    def test_template_has_one_sheet_with_the_columns(self):
        import io
        import openpyxl
        res = self.client_.get('/api/admin/supplier-offers/import/template/')
        self.assertEqual(res.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertEqual(book.sheetnames, ['Ofertas de proveedor'])
        # «Activa (Sí/No)» is optional (DC3 review P3.3): the template offers it, files without it still load.
        self.assertEqual([c.value for c in book.active[1]], OFFER_HEADERS + ['Activa (Sí/No)'])

    def test_preview_marks_errors_per_row_and_writes_nothing(self):
        res = self.preview([
            self.row(),
            self.row(**{'Proveedor': 'Nadie'}),
            self.row(**{'Código/SKU del producto': 'no-existe'}),
            self.row(**{'Precio comercial': 'gratis'}),
            self.row(**{'Vigente hasta': (self.start - timedelta(days=1)).isoformat()}),
            self.row(**{'Cantidad confirmada': -2}),
        ])
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual((body['rows_create'], body['rows_error']), (1, 5))
        self.assertEqual([r['action'] for r in body['rows']], ['create'] + ['error'] * 5)
        self.assertTrue(all(r['errors'] for r in body['rows'][1:]))
        self.assertEqual(SupplierOffer.objects.count(), 1)
        self.assertEqual(self.apply(body['id']).status_code, 400)       # a file with errors is not applied

    def test_apply_is_idempotent_and_updates_the_same_window(self):
        body = self.preview([self.row()]).json()
        res = self.apply(body['id'])
        self.assertEqual((res.status_code, res.json()['created']), (200, 1), res.content)
        made = SupplierOffer.objects.get(valid_from=self.start, valid_until=self.end)
        self.assertEqual((made.agreed_cost, made.commercial_price, made.confirmed_quantity, made.lead_time_days),
                         (D('68.00'), D('109.90'), 8, 5))
        self.assertEqual(self.apply(body['id']).json(), {'status': 'applied', 'created': 1, 'updated': 0,
                                                          'unchanged': 0})    # the same job twice: nothing
        again = self.preview([self.row()]).json()
        self.assertEqual(again['rows_no_change'], 1)
        self.apply(again['id'])
        self.assertEqual(SupplierOffer.objects.filter(valid_from=self.start, valid_until=self.end).count(), 1)
        # Next week's file: a new price for the same window updates it; a new window is a new offer.
        week = self.preview([self.row(**{'Precio comercial': '105'}),
                             self.row(**{'Vigente desde': (self.end + timedelta(days=1)).isoformat(),
                                         'Vigente hasta': (self.end + timedelta(days=7)).isoformat()})]).json()
        self.assertEqual((week['rows_update'], week['rows_create']), (1, 1))
        self.apply(week['id'])
        made.refresh_from_db()
        self.assertEqual(made.commercial_price, D('105.00'))
        self.assertEqual(SupplierOffer.objects.count(), 3)
        self.assertTrue(AdminAuditLog.objects.filter(action='supplier_offer_updated', target_id=str(made.pk)).exists())

    def test_the_cost_column_needs_costs_view(self):
        manager, _ = _p2d_member(self.company, 'dc3_sin_costo', ['company.view', 'products.view', 'products.manage'])
        client = APIClient()
        client.force_authenticate(user=manager)
        body = self.preview([self.row()], client=client).json()
        self.assertTrue(body['cost_column_ignored'])
        self.assertNotIn('agreed_cost', body['rows'][0]['data'])
        from store.models import BulkImportRow
        # Not read at all: the staged row does not keep it either.
        self.assertNotIn('agreed_cost', BulkImportRow.objects.get(job_id=body['id']).normalized_data)
        self.assertNotIn('68', str(body['rows'][0]['data']))
        self.apply(body['id'], client=client)
        made = SupplierOffer.objects.get(valid_from=self.start, valid_until=self.end)
        self.assertIsNone(made.agreed_cost)
        export = client.get('/api/admin/supplier-offers/export/')
        import io
        import openpyxl
        sheet = openpyxl.load_workbook(io.BytesIO(export.content)).active
        self.assertNotIn('Costo pactado', [c.value for c in sheet[1]])
        viewer, _ = _p2d_member(self.company, 'dc3_solo_ver', ['company.view', 'products.view'])
        as_viewer = APIClient()
        as_viewer.force_authenticate(user=viewer)
        self.assertEqual(self.preview([self.row()], client=as_viewer).status_code, 403)

    def test_a_variant_is_found_by_model_and_options_when_there_is_no_code(self):
        shop = DeviceShop()
        shop.company, shop._imei = self.company, iter([])
        from store import device_catalog
        capacity = device_catalog.create_attribute(company=self.company, name='Capacidad')
        gb = device_catalog.create_option(company=self.company, attribute_id=capacity.pk, label='512 GB')
        model = device_catalog.create_model(company=self.company, name='Tableta Z')
        variant = _prod(self.company, 'Tableta Z 512', 'tableta-z-512', price='5000.00', inventory=0)
        device_catalog.set_variant(variant, model_id=model.pk, option_ids=[gb.pk])
        body = self.preview([self.row(**{'Código/SKU del producto': '', 'Modelo': 'Tableta Z', 'Capacidad': '512 GB',
                                         'Línea': 'Nuevo'})]).json()
        self.assertEqual(body['rows'][0]['action'], 'create', body['rows'][0])
        self.apply(body['id'])
        self.assertTrue(SupplierOffer.objects.filter(product=variant).exists())

    def test_the_export_lists_current_offers_and_loads_back_unchanged(self):
        import io
        import openpyxl
        res = self.client_.get('/api/admin/supplier-offers/export/')
        self.assertEqual(res.status_code, 200)
        sheet = openpyxl.load_workbook(io.BytesIO(res.content)).active
        headers = [c.value for c in sheet[1]]
        for column in OFFER_HEADERS:
            self.assertIn(column, headers)
        from django.core.files.uploadedfile import SimpleUploadedFile
        upload = SimpleUploadedFile('ofertas.xlsx', res.content,
                                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        body = self.client_.post('/api/admin/supplier-offers/import/preview/', {'file': upload},
                                 format='multipart').json()
        self.assertEqual((body['rows_no_change'], body['rows_error']), (1, 0), body['rows'])

    def test_confirmed_quantity_below_what_was_sold_is_a_row_error(self):
        SupplierOffer.objects.filter(pk=self.offer.pk).update(committed_quantity=4)
        body = self.preview([self.row(**{'Vigente desde': self.offer.valid_from.isoformat(),
                                         'Vigente hasta': self.offer.valid_until.isoformat(),
                                         'Cantidad confirmada': 3})]).json()
        self.assertEqual(body['rows'][0]['action'], 'error')


class ExpiryCommandTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()

    def run_command(self):
        import io
        from django.core.management import call_command
        out = io.StringIO()
        call_command('expire_supplier_offers', stdout=out)
        return out.getvalue()

    def test_expired_offers_are_switched_off_once_audited_and_commitments_stay(self):
        self.add_on_order(quantity=2)
        order = self.checkout()
        self.pay(order)
        SupplierOffer.objects.filter(pk=self.offer.pk).update(valid_from=today() - timedelta(days=9),
                                                              valid_until=today() - timedelta(days=1))
        current = self.make_offer()
        self.assertIn('1', self.run_command())
        self.offer.refresh_from_db()
        current.refresh_from_db()
        self.assertEqual((self.offer.is_active, self.offer.committed_quantity, current.is_active), (False, 2, True))
        self.assertEqual(order.items.get().supply_status, 'pending')
        self.assertIn('0', self.run_command())                        # idempotent
        self.assertEqual(AdminAuditLog.objects.filter(action='supplier_offer_expired').count(), 1)


class SupplyFollowUpSurfacesTest(SupplyShop, TestCase):
    """What has not arrived is not handed over, and a paid order waiting for a person is said."""

    def setUp(self):
        self.build()
        self.add_on_order(quantity=1)
        self.order = self.checkout()
        self.pay(self.order)

    def test_an_order_does_not_leave_or_get_ready_before_its_goods_arrive(self):
        from store import order_fulfillment_services as fulfillment
        for state in ('ready_for_pickup', 'shipped', 'delivered'):
            with self.subTest(state), self.assertRaises(fulfillment.FulfillmentBlocked) as caught:
                fulfillment.change_fulfillment_status(order=self.order, new_status=state, actor=self.staff,
                                                      company=self.company)
            self.assertEqual(caught.exception.code, 'supply_pending')
        fulfillment.change_fulfillment_status(order=self.order, new_status='preparing', actor=self.staff,
                                              company=self.company)
        supply.receive_line(self.order.items.get(), actor=self.staff)
        self.order.refresh_from_db()
        fulfillment.change_fulfillment_status(order=self.order, new_status='delivered', actor=self.staff,
                                              company=self.company)

    def test_the_status_check_says_when_a_paid_order_waits_for_its_supply(self):
        import io
        from django.core.management import call_command
        out = io.StringIO()
        try:
            call_command('ops_status', stdout=out)
        except SystemExit:
            pass                      # other areas of an empty test installation may need somebody
        self.assertIn('OK    abastecimiento', out.getvalue())
        self.order.items.update(supply_status='to_confirm')
        out = io.StringIO()
        try:
            call_command('ops_status', stdout=out)
        except SystemExit:
            pass
        text = out.getvalue()
        self.assertIn('ATENCIÓN abastecimiento', text)
        self.assertNotIn('ana@example.com', text)


class RollbackOnOrderTest(SupplyShop, TestCase):
    """Before going back to a release before DC3: no payable on-order order is left."""

    def setUp(self):
        self.build()

    def test_pending_on_order_orders_are_listed_then_cancelled(self):
        import io
        from django.core.management import call_command
        self.add_on_order(quantity=1)
        pending = self.checkout()
        self.add_on_order(quantity=1, session='carrito-b')
        paid = self.checkout('carrito-b')
        self.pay(paid)
        out = io.StringIO()
        call_command('cancel_pending_on_order_orders', stdout=out)
        self.assertIn(f'#{pending.pk}', out.getvalue())
        self.assertNotIn(f'#{paid.pk}', out.getvalue())
        pending.refresh_from_db()
        self.assertEqual(pending.status, Order.Status.PENDING_PAYMENT)        # a simulation changes nothing
        call_command('cancel_pending_on_order_orders', '--apply', stdout=io.StringIO())
        pending.refresh_from_db()
        paid.refresh_from_db()
        self.assertEqual((pending.status, paid.status), (Order.Status.CANCELLED, Order.Status.PAID))
        self.assertTrue(AdminAuditLog.objects.filter(action='order_cancelled_before_rollback',
                                                     target_id=str(pending.pk)).exists())


# ---------------------------------------------------------------------------
# DC3 independent review (P2, P3.1–P3.10)
# ---------------------------------------------------------------------------

STALE = 'El archivo cambió: vuelve a previsualizarlo'


class OfferFile(SupplyShop):
    """The Excel helpers of `OfferExcelTest`, without its tests."""

    file, row, preview, apply = OfferExcelTest.file, OfferExcelTest.row, OfferExcelTest.preview, OfferExcelTest.apply

    def build_excel(self, slug='dc3-tienda'):
        self.build(slug)
        self.client_ = APIClient()
        self.client_.force_authenticate(user=self.staff)
        self.start, self.end = today(), today() + timedelta(days=6)


class ReviewExcelTest(OfferFile, TestCase):
    """The weekly Excel after the review: one offer per key, the cost by who applies, nothing reactivated."""

    def setUp(self):
        self.build_excel()

    def test_the_same_file_previewed_twice_creates_one_offer_and_the_second_apply_is_refused(self):
        first = self.preview([self.row()]).json()
        second = self.preview([self.row()]).json()
        self.assertEqual((first['rows_create'], second['rows_create']), (1, 1))
        self.assertEqual(self.apply(first['id']).status_code, 200)
        refused = self.apply(second['id'])
        self.assertEqual(refused.status_code, 400, refused.content)
        self.assertIn(STALE, refused.json()['detail'])
        self.assertEqual(SupplierOffer.objects.filter(valid_from=self.start, valid_until=self.end).count(), 1)
        # The sellable quantity on order is the offer's 8, not 16.
        self.assertEqual(sum(o.remaining for o in supply.valid_offers(self.case).filter(valid_from=self.start)), 8)

    def test_the_database_refuses_a_second_offer_with_the_same_key(self):
        from django.db import IntegrityError
        with self.assertRaises(supply.SupplyError) as caught:
            self.make_offer(valid_from=self.offer.valid_from, valid_until=self.offer.valid_until)
        self.assertEqual(caught.exception.code, 'duplicate')
        with self.assertRaises(IntegrityError), transaction.atomic():
            SupplierOffer.objects.create(
                company=self.company, supplier=self.supplier, product=self.case, commercial_price=D('1'),
                valid_from=self.offer.valid_from, valid_until=self.offer.valid_until, lead_time_days=1)

    def test_an_integrity_error_at_apply_is_the_same_refusal(self):
        from unittest import mock
        from django.db import IntegrityError
        body = self.preview([self.row()]).json()
        with mock.patch.object(SupplierOffer.objects, 'create', side_effect=IntegrityError('duplicate key')):
            refused = self.apply(body['id'])
        self.assertEqual(refused.status_code, 400, refused.content)
        self.assertIn(STALE, refused.json()['detail'])

    def test_the_cost_is_written_only_if_who_applies_may_see_costs(self):
        """P3.1 — previewed by `costs.view`, applied by somebody without it: no cost is written."""
        manager, _ = _p2d_member(self.company, 'dc3_aplica_sin_costo',
                                 ['company.view', 'products.view', 'products.manage'])
        other = APIClient()
        other.force_authenticate(user=manager)
        SupplierOffer.objects.filter(pk=self.offer.pk).update(agreed_cost=D('70.00'))
        body = self.preview([
            self.row(),
            self.row(**{'Vigente desde': self.offer.valid_from.isoformat(),
                        'Vigente hasta': self.offer.valid_until.isoformat(), 'Costo pactado': '50',
                        'Precio comercial': '111'}),
        ]).json()
        self.assertEqual((body['rows_create'], body['rows_update']), (1, 1), body['rows'])
        self.assertEqual(self.apply(body['id'], client=other).status_code, 200)
        made = SupplierOffer.objects.get(valid_from=self.start, valid_until=self.end)
        self.assertIsNone(made.agreed_cost)
        self.offer.refresh_from_db()
        self.assertEqual((self.offer.agreed_cost, self.offer.commercial_price), (D('70.00'), D('111.00')))

    def test_an_empty_cost_cell_changes_nothing(self):
        """P3.2 — empty means «no change», like the other imports."""
        key = {'Vigente desde': self.offer.valid_from.isoformat(), 'Vigente hasta': self.offer.valid_until.isoformat(),
               'Precio comercial': '110', 'Cantidad confirmada': 5, 'Plazo (días)': 4}
        body = self.preview([self.row(**key, **{'Costo pactado': ''})]).json()
        self.assertEqual(body['rows'][0]['action'], 'no_change', body['rows'])
        body = self.preview([self.row(**{**key, 'Precio comercial': '112'}, **{'Costo pactado': ''})]).json()
        self.assertEqual(body['rows'][0]['action'], 'update')
        self.apply(body['id'])
        self.offer.refresh_from_db()
        self.assertEqual((self.offer.agreed_cost, self.offer.commercial_price), (D('70.00'), D('112.00')))

    def test_a_manually_deactivated_offer_is_not_reactivated_by_a_reupload(self):
        """P3.3 — off stays off unless the row says «Activa: Sí»; the preview says so."""
        supply.deactivate_offer(self.offer, actor=self.staff, reason='Proveedor sin stock')
        key = {'Vigente desde': self.offer.valid_from.isoformat(), 'Vigente hasta': self.offer.valid_until.isoformat(),
               'Precio comercial': '110', 'Cantidad confirmada': 5, 'Plazo (días)': 4, 'Costo pactado': '70'}
        body = self.preview([self.row(**key)]).json()
        self.assertEqual(body['rows'][0]['action'], 'no_change', body['rows'])
        self.assertTrue(any('desactivada a mano: no se reactiva' in w for w in body['rows'][0]['warnings']))
        body = self.preview([self.row(**{**key, 'Precio comercial': '115'})]).json()
        self.assertEqual(body['rows'][0]['action'], 'update')
        self.apply(body['id'])
        self.offer.refresh_from_db()
        self.assertEqual((self.offer.is_active, self.offer.commercial_price), (False, D('115.00')))
        # Asked for explicitly, it is reactivated.
        headers = OFFER_HEADERS + ['Activa (Sí/No)']
        body = self.preview([self.row(**key) + ['Sí']], headers=headers).json()
        self.assertEqual(body['rows'][0]['action'], 'update', body['rows'])
        self.apply(body['id'])
        self.offer.refresh_from_db()
        self.assertTrue(self.offer.is_active)
        bad = self.preview([self.row(**key) + ['quizás']], headers=headers).json()
        self.assertEqual(bad['rows'][0]['action'], 'error')

    def test_apply_locks_the_jobs_offers_in_id_order_before_writing(self):
        """P3.8 — the file's row order is not the lock order: one statement, by id, first."""
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        later = self.make_offer(valid_from=self.start + timedelta(days=10), valid_until=self.end + timedelta(days=10))
        rows = [self.row(**{'Vigente desde': o.valid_from.isoformat(), 'Vigente hasta': o.valid_until.isoformat(),
                            'Precio comercial': '99'}) for o in (later, self.offer)]
        body = self.preview(rows).json()
        self.assertEqual(body['rows_update'], 2, body['rows'])
        with CaptureQueriesContext(connection) as queries:
            self.assertEqual(self.apply(body['id']).status_code, 200)
        sql = [q['sql'] for q in queries.captured_queries]
        locks = [i for i, s in enumerate(sql) if 'FOR UPDATE' in s and '"store_supplieroffer"' in s.split('FROM')[1][:40]]
        writes = [i for i, s in enumerate(sql) if s.startswith('UPDATE "store_supplieroffer"')]
        self.assertTrue(locks and writes)
        first = sql[locks[0]]
        self.assertLess(locks[0], writes[0])
        self.assertIn(str(later.pk), first)
        self.assertIn(str(self.offer.pk), first)
        self.assertIn('ORDER BY "store_supplieroffer"."id" ASC', first)


class ReviewOfferApplyRaceTest(TransactionTestCase):
    """P2 — two people apply the same file at the same instant: one offer, one refusal."""

    def test_one_applies_and_the_other_is_refused(self):
        from threading import Barrier, Thread

        from django.db import connection
        from store import supplier_offer_import as offer_import
        from store.models import BulkImportJob

        shop = OfferFile()
        shop.build_excel('dc3-dos-cargas')
        jobs = [shop.preview([shop.row()]).json()['id'] for _ in range(2)]
        barrier, outcomes = Barrier(2), []

        def run(job_id):
            try:
                barrier.wait(timeout=5)
                try:
                    offer_import.apply(job=BulkImportJob.objects.get(pk=job_id), actor=shop.staff)
                    outcomes.append('applied')
                except offer_import.OfferImportError as exc:
                    outcomes.append(str(exc))
            finally:
                connection.close()

        threads = [Thread(target=run, args=(job,)) for job in jobs]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=20)
        self.assertEqual(len(outcomes), 2, outcomes)
        self.assertIn('applied', outcomes)
        self.assertTrue(any(STALE in o for o in outcomes), outcomes)
        self.assertEqual(SupplierOffer.objects.filter(valid_from=shop.start, valid_until=shop.end).count(), 1)


class ReviewSupplyTest(SupplyShop, TestCase):
    def setUp(self):
        self.build()

    def test_a_refunded_or_cancelled_order_releases_what_it_committed(self):
        """P3.4 — the offer gets back the quantity, never below zero, audited."""
        for status_ in (Order.Status.REFUNDED, Order.Status.CANCELLED):
            with self.subTest(status_):
                SupplierOffer.objects.filter(pk=self.offer.pk).update(committed_quantity=0)
                CartItem.objects.all().delete()
                self.add_on_order(quantity=2)
                order = self.checkout()
                self.pay(order)
                self.offer.refresh_from_db()
                self.assertEqual(self.offer.committed_quantity, 2)
                Order.objects.filter(pk=order.pk).update(status=status_)
                supply.fail_line(order.items.get(), actor=self.staff, reason='Reembolsado al cliente')
                self.offer.refresh_from_db()
                self.assertEqual(self.offer.committed_quantity, 0)
                self.assertEqual(order.items.get().supply_status, 'failed')
                log = AdminAuditLog.objects.filter(action='order_supply_failed', target_id=str(order.pk)).get()
                self.assertEqual(log.metadata['released'], 2)
        # Never below zero: what the offer no longer holds is not taken twice.
        CartItem.objects.all().delete()
        self.add_on_order(quantity=2)
        order = self.checkout()
        self.pay(order)
        SupplierOffer.objects.filter(pk=self.offer.pk).update(committed_quantity=1)
        Order.objects.filter(pk=order.pk).update(status=Order.Status.REFUNDED)
        supply.fail_line(order.items.get(), actor=self.staff, reason='Reembolsado')
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 0)
        # Confirming or receiving still needs a paid order.
        CartItem.objects.all().delete()
        self.add_on_order(quantity=1)
        order = self.checkout()
        self.pay(order)
        Order.objects.filter(pk=order.pk).update(status=Order.Status.REFUNDED)
        with self.assertRaises(supply.SupplyError):
            supply.receive_line(order.items.get(), actor=self.staff)

    def test_payment_asks_to_confirm_when_the_product_left_the_on_order_mode(self):
        """P3.5 — switched to «sólo stock» between checkout and payment: nothing is committed."""
        self.add_on_order(quantity=1)
        order = self.checkout()
        Product.objects.filter(pk=self.case.pk).update(sale_mode=Product.SALE_MODE_STOCK)
        self.assertEqual(self.pay(order), 'paid')
        self.assertEqual(order.items.get().supply_status, 'to_confirm')
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 0)

    def test_an_order_without_any_branch_still_commits_its_on_order_lines(self):
        """P3.7 — no branch on the order nor in the company: the offer is committed anyway."""
        from unittest import mock
        self.add_on_order(quantity=2)
        order = self.checkout()
        Order.objects.filter(pk=order.pk).update(fulfillment_branch=None)
        order.refresh_from_db()
        with mock.patch('store.tenancy.company_fulfillment_branch', return_value=None):
            inventory.record_sale_stock_movements(order)
        self.assertEqual(order.items.get().supply_status, 'pending')
        self.offer.refresh_from_db()
        self.assertEqual(self.offer.committed_quantity, 2)

    def test_committed_requires_a_confirmed_quantity_in_the_database(self):
        """P3.10 — NULL confirmed never holds a commitment."""
        from django.db import IntegrityError
        with self.assertRaises(IntegrityError), transaction.atomic():
            SupplierOffer.objects.filter(pk=self.offer.pk).update(confirmed_quantity=None, committed_quantity=1)
        SupplierOffer.objects.filter(pk=self.offer.pk).update(confirmed_quantity=None, committed_quantity=0)



@override_settings(ALLOWED_HOSTS=['.tienda.pe', 'testserver'])
class ReviewShelfFirstWebTest(SupplyShop, TestCase):
    """P3.6 — «Bajo pedido» only when the shelf cannot cover the line (web)."""

    def setUp(self):
        self.build()
        self.public = APIClient(HTTP_HOST=self.host)

    def add(self, quantity, session='carrito-a'):
        return self.public.post('/api/cart/add/', {'session_key': session, 'product': self.case.pk,
                                                   'quantity': quantity, 'on_order': True}, format='json')

    def test_on_order_is_refused_while_the_shelf_covers_the_line(self):
        self.stock(self.case, 3)
        refused = self.add(2)
        self.assertEqual((refused.status_code, refused.json()['code']), (409, 'stock_available'), refused.content)
        self.assertIn('stock', refused.json()['detail'].lower())
        self.assertFalse(CartItem.objects.exists())
        self.assertEqual(self.add(3).json().get('code'), 'stock_available')
        taken = self.add(4)
        self.assertEqual(taken.status_code, 200, taken.content)
        # Lowering an on-order line to what the shelf covers is refused the same way.
        patched = self.public.patch(f'/api/cart/{CartItem.objects.get().pk}/?session_key=carrito-a',
                                    {'quantity': 3}, format='json')
        self.assertEqual((patched.status_code, patched.json()['code']), (409, 'stock_available'))


class ReviewShelfFirstPosTest(SupplyShop, TestCase):
    """P3.6 — the same rule at the till."""

    def setUp(self):
        self.build()
        self.cashier, _ = _p2d_member(self.company, 'dc3_cajero_r', ['company.view', 'sales.pos.use'])

    def test_the_till_refuses_on_order_while_the_shelf_covers_it(self):
        from store import pos_services
        from store.tests import _c1_sale
        self.stock(self.case, 3)
        with self.assertRaises(pos_services.PosStockAvailableError):
            _c1_sale(actor=self.cashier, company=self.company, branch=self.branch,
                     items=[{'product': self.case.pk, 'quantity': 2, 'on_order': True}])
        client = APIClient()
        client.force_authenticate(user=self.cashier)
        res = client.post('/api/admin/pos/preview/', {
            'branch': self.branch.pk, 'payment_method': 'cash',
            'items': [{'product': self.case.pk, 'quantity': 3, 'on_order': True}]}, format='json')
        self.assertEqual((res.status_code, res.json().get('code')), (409, 'stock_available'), res.content)
        order, _ = _c1_sale(actor=self.cashier, company=self.company, branch=self.branch,
                            items=[{'product': self.case.pk, 'quantity': 4, 'on_order': True}])
        self.assertEqual(order.items.get().supply_origin, 'supplier')
