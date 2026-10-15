"""
UNIT-IMPORT — la carga de equipos: muchos equipos de una vez.

    UNA FILA = UNA UNIDAD FÍSICA.

La carga masiva de inventario escribe cantidades, y un producto con serie no
tiene una cantidad que alguien escribe: tiene equipos. Por eso un archivo de
stock rechaza esos productos, y hasta ahora la única forma de registrarlos era
uno por uno en pantalla.

Esta plantilla no es otro camino al stock. Cada fila termina en
`stock_unit_services.receive_units`, el mismo escritor que usa «Registrar
equipo»: mismas reglas de serie e IMEI, misma línea de Kardex, misma invariante
(cantidad == equipos disponibles). Lo que añade es el ensayo: se ve qué va a
pasar con cada fila antes de que pase, y entra todo o no entra nada.
"""
import io

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from store import unit_import_services as unit_import
from store.models import (
    AdminAuditLog, BulkImportJob, BulkImportRow, Membership, Product, ProductBarcode,
    StockMovement, StockUnit,
)
from store.test_stock_units import IMEI_A, IMEI_B, IMEI_C, IMEI_D, UnitsBase, imei
from store.tests import _m7_user

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
TEMPLATE_URL = '/api/admin/inventory/units/import/template/'
PREVIEW_URL = '/api/admin/inventory/units/import/preview/'
HEADERS = ['Código', 'Producto', 'Sucursal', 'Número de serie', 'IMEI', 'IMEI 2', 'Condición', 'Costo',
           'Motivo / referencia']
#: DC2 added these, optional, at the end of the template.
DC2_HEADERS = ['Modelo', 'Capacidad', 'Color', 'Línea', 'Estado estético', 'Batería %', 'Precio de venta']


def apply_url(job_id):
    return f'/api/admin/inventory/units/import/{job_id}/apply/'


def workbook(rows, headers=HEADERS, sheet='Equipos', name='equipos.xlsx'):
    book = Workbook()
    page = book.active
    page.title = sheet
    page.append(headers)
    for row in rows:
        page.append(row)
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile(name, out.getvalue(), content_type=XLSX)


class ImportBase(UnitsBase):
    def setUp(self):
        super().setUp()
        ProductBarcode.objects.create(company=self.company, product=self.phone, code='IP16-256-N')
        ProductBarcode.objects.create(company=self.company, product=self.laptop, code='MBA-13-M3')
        self.api = APIClient()
        self.api.force_authenticate(user=self.staff)

    def phone_row(self, serial, imei, **over):
        row = {'Código': 'IP16-256-N', 'Producto': '', 'Sucursal': self.branch_a.name, 'Número de serie': serial,
               'IMEI': imei, 'IMEI 2': '', 'Condición': 'Nuevo', 'Costo': '3800.00',
               'Motivo / referencia': 'Compra F001-204'}
        row.update(over)
        return [row[h] for h in HEADERS]

    def preview(self, rows, client=None, **form):
        return (client or self.api).post(PREVIEW_URL, {'file': workbook(rows), **form}, format='multipart')

    def staged(self, rows, **form):
        response = self.preview(rows, **form)
        self.assertEqual(response.status_code, 201, response.content)
        return response.json()

    def errors(self, payload):
        return {row['row']: ' '.join(row['errors']) for row in payload['rows'] if row['errors']}


def legacy_template():
    """
    The template as it was before IMPORT-SPLIT-01: an example row under the
    headers and a second sheet of instructions. People still have copies of it.
    """
    book = Workbook()
    page = book.active
    page.title = 'Equipos'
    page.append(HEADERS)
    page.append(['IP16PM-256-W', 'iPhone 16 Pro Max 256GB White', 'Tienda principal', 'EJEMPLO-F2LXK0PQN7',
                 '490154203237518', '', 'Nuevo', 3999.00, 'Compra factura F001-123'])
    notes = book.create_sheet('Instrucciones')
    notes.append([None, 'Cómo cargar equipos con número de serie'])
    notes.append([None, 'UNA FILA = UNA UNIDAD FÍSICA'])
    return book


def saved(book, name='equipos.xlsx'):
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile(name, out.getvalue(), content_type=XLSX)


class TemplateTest(ImportBase):
    """IMPORT-SPLIT-01 — ONE sheet, «Equipos»: the headers and nothing under them."""

    def book(self, client=None):
        response = (client or self.api).get(TEMPLATE_URL)
        self.assertEqual(response.status_code, 200)
        return response, load_workbook(io.BytesIO(response.content))

    def test_one_sheet_with_the_headers_and_nothing_else(self):
        response, book = self.book()

        self.assertIn('carga-masiva-equipos.xlsx', response['Content-Disposition'])
        self.assertEqual(book.sheetnames, ['Equipos'])
        page = book['Equipos']
        self.assertEqual([c.value for c in page[1]], [*HEADERS, *DC2_HEADERS])
        self.assertEqual(page.max_row, 1)                                # no example row, no help row
        self.assertEqual(page.freeze_panes, 'A2')
        # No hay columna de cantidad: una fila es un equipo.
        self.assertFalse([h for h in HEADERS if 'antidad' in h or 'tock' in h])

    def test_identifiers_stay_text(self):
        """Un IMEI de 15 dígitos guardado como número lo estropea Excel."""
        _response, book = self.book()
        page = book['Equipos']
        for column in ('A', 'D', 'E', 'F'):
            self.assertEqual(page.column_dimensions[column].number_format, '@', column)

    def test_the_condition_is_chosen_from_the_list_the_system_knows(self):
        _response, book = self.book()
        rules = {str(rule.sqref): rule for rule in book['Equipos'].data_validations.dataValidation}

        condition = rules['G2:G501']
        self.assertEqual(condition.type, 'list')
        self.assertEqual(condition.formula1, '"Nuevo,Caja abierta,Reacondicionado,Usado"')
        # As in the owner's file: the list offers and does not block. Excel compares a
        # typed «nuevo» letter by letter; the upload is what accepts or refuses.
        self.assertFalse(condition.showErrorMessage)
        self.assertEqual(
            condition.formula1.strip('"').split(','), [label for _value, label in StockUnit.Condition.choices])

    def test_the_branch_list_is_the_branches_this_person_can_load_into(self):
        _response, book = self.book()
        rules = {str(rule.sqref): rule for rule in book['Equipos'].data_validations.dataValidation}
        branches = rules['C2:C501']
        self.assertEqual(branches.type, 'list')
        self.assertEqual(sorted(branches.formula1.strip('"').split(',')), sorted([self.branch_a.name, self.branch_b.name]))
        self.assertNotIn(self.foreign_branch.name, branches.formula1)
        self.assertFalse(branches.showErrorMessage)                        # a help, not a wall: the server decides

        # Quien sólo opera una sucursal no ve el nombre de las demás.
        self.restrict_to_branch_a()
        _response, theirs = self.book()
        only = {str(rule.sqref): rule for rule in theirs['Equipos'].data_validations.dataValidation}['C2:C501']
        self.assertEqual(only.formula1, f'"{self.branch_a.name}"')

    def test_the_cost_is_a_number_with_two_decimals(self):
        _response, book = self.book()
        self.assertEqual(book['Equipos'].column_dimensions['H'].number_format, '0.00')

    def test_a_branch_name_a_cell_list_cannot_hold_leaves_the_list_out_and_the_template_whole(self):
        """A comma, a formula sign, a control character somebody pasted: no list, never a broken download."""
        from store.models import Branch

        for name in ('Lima, Norte', '=1+1', '+Sur', 'Norte\x01', 'Centro\x0b'):
            with self.subTest(name=name):
                Branch.objects.filter(pk=self.branch_b.pk).update(name=name)
                response, book = self.book()
                rules = {str(rule.sqref) for rule in book['Equipos'].data_validations.dataValidation}
                self.assertNotIn('C2:C501', rules)
                self.assertIn('G2:G501', rules)
                self.assertEqual(book.sheetnames, ['Equipos'])

    def test_the_template_of_a_company_that_is_not_theirs_is_not_there(self):
        response = self.api.get(f'{TEMPLATE_URL}?company={self.other.pk}')
        self.assertIn(response.status_code, (403, 404))
        self.assertNotIn('spreadsheetml', response.get('Content-Type', ''))

    def test_each_header_says_what_goes_under_it(self):
        _response, book = self.book()
        notes = {cell.value: cell.comment.text for cell in book['Equipos'][1] if cell.comment}

        self.assertEqual(set(notes), {*HEADERS, *DC2_HEADERS})
        self.assertIn('Obligatorio', notes['Número de serie'])
        self.assertIn('15 dígitos', notes['IMEI'])
        self.assertIn('lleva IMEI', notes['IMEI'])                         # según lo que el producto lleva
        self.assertIn('dos líneas', notes['IMEI 2'])
        self.assertIn('Una fila', notes['Código'])

    def test_uploading_the_template_as_downloaded_registers_nothing(self):
        response, _book = self.book()
        template = SimpleUploadedFile('carga-masiva-equipos.xlsx', response.content, content_type=XLSX)

        staged = self.api.post(PREVIEW_URL, {'file': template}, format='multipart')

        self.assertEqual(staged.status_code, 201, staged.content)
        payload = staged.json()
        self.assertEqual((payload['counts']['total'], payload['counts']['create'], payload['counts']['error']), (0, 0, 0))
        self.assertFalse(StockUnit.objects.exists())

    def test_the_row_right_under_the_headers_is_a_device(self):
        response, book = self.book()
        book['Equipos'].append(self.phone_row('F2LXK1ABC1', IMEI_A))

        payload = self.api.post(PREVIEW_URL, {'file': saved(book)}, format='multipart').json()

        self.assertEqual((payload['counts']['create'], payload['counts']['error']), (1, 0), payload['rows'])

    def test_a_file_made_with_the_previous_template_still_loads(self):
        """Its example row is not a device and its instructions are not a sheet of devices."""
        book = legacy_template()
        book['Equipos'].append(self.phone_row('F2LXK1ABC1', IMEI_A))

        staged = self.api.post(PREVIEW_URL, {'file': saved(book, 'Equipos serializados.xlsx')}, format='multipart')

        self.assertEqual(staged.status_code, 201, staged.content)
        payload = staged.json()
        self.assertEqual((payload['counts']['create'], payload['counts']['error']), (1, 0), payload['rows'])
        applied = self.api.post(apply_url(payload['id']))
        self.assertEqual(applied.status_code, 200, applied.content)
        self.assertEqual([u.serial_number for u in StockUnit.objects.all()], ['F2LXK1ABC1'])


class WhatADeviceCarriesTest(ImportBase):
    """
    IMPORT-SPLIT-01 — serial always; IMEI and IMEI 2 by what the PRODUCT carries.

    Three shapes, and only three: serial alone, serial + IMEI, serial + two
    IMEIs. A product marked «lleva IMEI» needs the IMEI on every device.
    """

    def laptop_row(self, serial, imei='', imei2=''):
        return ['MBA-13-M3', '', self.branch_a.name, serial, imei, imei2, '', '', 'Compra F001-204']

    def test_the_three_shapes_are_accepted(self):
        payload = self.staged([
            self.laptop_row('C02ZK0ABMD6R'),                                        # serial alone
            self.phone_row('F2LXK1ABC1', IMEI_A),                                   # serial + IMEI
            self.phone_row('F2LXK1ABC2', IMEI_B, **{'IMEI 2': IMEI_C}),             # serial + two IMEIs
        ])

        self.assertEqual((payload['counts']['create'], payload['counts']['error']), (3, 0), payload['rows'])
        self.assertEqual(self.api.post(apply_url(payload['id'])).status_code, 200)
        self.assertEqual(
            sorted((u.serial_number, u.imei, u.imei2) for u in StockUnit.objects.all()),
            [('C02ZK0ABMD6R', None, None), ('F2LXK1ABC1', IMEI_A, None), ('F2LXK1ABC2', IMEI_B, IMEI_C)])

    def test_an_imei_under_a_product_not_marked_as_carrying_one_is_kept_and_said(self):
        """
        The rule of the screen stands (there, an IMEI is «opcional para este
        producto»). A file is where a row lands under the wrong product: the
        preview says it, and does not refuse it.
        """
        payload = self.staged([
            self.laptop_row('C02ZK0ABMD6R', IMEI_A),
            self.laptop_row('C02ZK0ABMD6T'),
        ])

        self.assertEqual((payload['counts']['create'], payload['counts']['error']), (2, 0), payload['rows'])
        said = {row['row']: ' '.join(row['warnings']) for row in payload['rows'] if row['warnings']}
        self.assertEqual(list(said), [2])
        self.assertIn('no está marcado «lleva IMEI»', said[2])
        self.assertIn(self.laptop.name, said[2])
        self.assertEqual(self.api.post(apply_url(payload['id'])).status_code, 200)
        self.assertEqual(StockUnit.objects.get(serial_number='C02ZK0ABMD6R').imei, IMEI_A)

    def test_it_is_written_down_that_a_load_carried_such_imeis(self):
        """The warning is on each row; the audit line of the load says how many, without naming a device."""
        payload = self.staged([
            self.laptop_row('C02ZK0ABMD6R', IMEI_A),
            self.laptop_row('C02ZK0ABMD6S', IMEI_B),
            self.phone_row('F2LXK1ABC1', IMEI_C),
        ])
        self.assertEqual(payload['summary']['imei_on_unmarked_products'], 2)
        kept = [row.row_number for row in BulkImportRow.objects.filter(job_id=payload['id']) if row.warnings]
        self.assertEqual(sorted(kept), [2, 3])                              # stays with the job, row by row

        self.assertEqual(self.api.post(apply_url(payload['id'])).status_code, 200)

        logged = AdminAuditLog.objects.get(action='stock_units_import_applied', target_id=payload['id'])
        self.assertEqual(logged.metadata['imei_on_unmarked_products'], 2)
        self.assertNotIn(IMEI_A, str(logged.metadata))
        # And a load without any says zero, rather than leaving the key out.
        clean = self.staged([self.phone_row('F2LXK1ABC2', IMEI_D)])
        self.assertEqual(clean['summary']['imei_on_unmarked_products'], 0)
        self.assertEqual(self.api.post(apply_url(clean['id'])).status_code, 200)
        self.assertEqual(
            AdminAuditLog.objects.get(action='stock_units_import_applied', target_id=clean['id']).metadata['imei_on_unmarked_products'], 0)

    def test_a_second_imei_alone_is_refused_for_any_product(self):
        payload = self.staged([self.laptop_row('C02ZK0ABMD6S', '', IMEI_B)])
        self.assertIn('acompaña al primero', self.errors(payload)[2])

    def test_a_second_imei_is_optional_and_checked_like_the_first_when_it_is_there(self):
        """IMEI 2: format, check digit and uniqueness — in the file, and against every device of the company, either column."""
        other_first, other_second = imei('35693803564391'), imei('35693803564392')
        self.receive(
            self.row('YAESTABA01', IMEI_A), self.row('YAESTABA02', IMEI_B, imei2=IMEI_C),
            self.row('YAESTABA03', other_first, imei2=other_second))
        wrong_digit = IMEI_D[:-1] + ('0' if IMEI_D[-1] != '0' else '1')
        fresh = [imei(f'3569380356440{n}') for n in range(6)]
        payload = self.staged([
            self.phone_row('F2LXK1ABC1', fresh[0], **{'IMEI 2': '12345'}),                    # 2 not 15 digits
            self.phone_row('F2LXK1ABC2', fresh[1], **{'IMEI 2': wrong_digit}),                 # 3 check digit
            self.phone_row('F2LXK1ABC3', fresh[2], **{'IMEI 2': 'N/A'}),                       # 4 a placeholder
            self.phone_row('F2LXK1ABC4', fresh[3], **{'IMEI 2': IMEI_A}),                      # 5 another device's IMEI
            self.phone_row('F2LXK1ABC5', fresh[4], **{'IMEI 2': IMEI_C}),                      # 6 another device's IMEI 2
            self.phone_row('F2LXK1ABC6', other_second),                                        # 7 IMEI that is somebody's IMEI 2
            self.phone_row('F2LXK1ABC7', fresh[5]),                                            # 8 fine, and no second one
        ])

        errors = self.errors(payload)
        self.assertEqual(sorted(errors), [2, 3, 4, 5, 6, 7], errors)
        self.assertIn('segundo IMEI', errors[2])
        self.assertIn('control', errors[3])
        self.assertIn('ya está registrado', errors[5])
        self.assertIn('ya está registrado', errors[6])
        self.assertIn('ya está registrado', errors[7])
        created = [row for row in payload['rows'] if row['action'] == 'create']
        self.assertEqual([row['row'] for row in created], [8])

    def test_a_product_with_imei_needs_it_and_a_second_one_only_after_the_first(self):
        payload = self.staged([
            self.phone_row('F2LXK1ABC1', ''),
            self.phone_row('F2LXK1ABC2', '', **{'IMEI 2': IMEI_B}),
            self.phone_row('F2LXK1ABC3', IMEI_C, **{'IMEI 2': IMEI_C}),
        ])

        self.assertEqual(sorted(self.errors(payload)), [2, 3, 4])

    def test_any_column_that_counts_is_refused_whole(self):
        for header in ('Cantidad', 'Cantidad inicial', 'Stock inicial', 'Cant.', 'Unidades', 'Stock - Tienda principal', 'QTY',
                       'Nº de unidades', 'N° unidades', 'Total unidades', 'Inventario', 'Saldo', 'Quantity', 'Existencias iniciales'):
            with self.subTest(header):
                response = self.api.post(
                    PREVIEW_URL,
                    {'file': workbook([[*self.phone_row('F2LXK1ABC1', IMEI_A), 2]], headers=[*HEADERS, header])},
                    format='multipart')
                self.assertEqual(response.status_code, 400, response.content)
                self.assertIn('una fila es un equipo', response.json()['detail'])
        self.assertFalse(BulkImportJob.objects.exists())

    def test_a_column_that_only_sounds_like_one_is_not_refused(self):
        """«Costo unitario» has «unit» in it and is not a count."""
        headers = [h if h != 'Costo' else 'Costo unitario' for h in HEADERS]
        response = self.api.post(
            PREVIEW_URL, {'file': workbook([self.phone_row('F2LXK1ABC1', IMEI_A)], headers=headers)}, format='multipart')
        self.assertEqual(response.status_code, 201, response.content)

    def test_a_column_that_counts_something_else_is_not_refused(self):
        """A reorder point, a unit of measure, how many SIMs: none of them says «two devices»."""
        for header in ('Stock mínimo', 'Unidades de medida', 'Cant. de SIM', 'Cantidad de SIM', 'Unidades por caja', 'Total'):
            with self.subTest(header):
                response = self.api.post(
                    PREVIEW_URL,
                    {'file': workbook([[*self.phone_row('F2LXK1ABC1', IMEI_A), 1]], headers=[*HEADERS, header])},
                    format='multipart')
                self.assertEqual(response.status_code, 201, response.content)

    def test_the_serial_column_is_found_as_people_head_it(self):
        for header in ('Nº de serie', 'N° de serie', 'No. de serie', 'Nro. de serie', 'S/N', 'Serial'):
            with self.subTest(header):
                headers = [h if h != 'Número de serie' else header for h in HEADERS]
                response = self.api.post(
                    PREVIEW_URL, {'file': workbook([self.phone_row('F2LXK1ABC1', IMEI_A)], headers=headers)}, format='multipart')
                self.assertEqual(response.status_code, 201, response.content)
                self.assertEqual(response.json()['counts']['create'], 1)

    def test_a_warned_row_is_not_buried_under_the_rows_that_fit_on_the_screen(self):
        """The screen shows the first rows of a long file: what is wrong first, what is warned next."""
        from store.import_views import MAX_PREVIEW_ROWS

        rows = [self.laptop_row(f'C02ZK{n:07d}') for n in range(MAX_PREVIEW_ROWS + 5)]
        rows.append(self.laptop_row('C02ZKWARNED1', IMEI_A))                 # the very last row of the file
        payload = self.staged(rows)

        self.assertTrue(payload['rows_truncated'])
        self.assertEqual(len(payload['rows']), MAX_PREVIEW_ROWS)
        self.assertEqual(payload['summary']['imei_on_unmarked_products'], 1)
        warned = [row['row'] for row in payload['rows'] if row['warnings']]
        self.assertEqual(warned, [len(rows) + 1])
        self.assertEqual(payload['rows'][0]['row'], len(rows) + 1)           # nothing wrong in the file: the warned one leads

    def test_the_audit_line_counts_what_was_stored_not_what_the_preview_saw(self):
        """«Lleva IMEI» can be switched between the preview and the load."""
        seen_marked = self.staged([self.phone_row('F2LXK1ABC1', IMEI_A), self.phone_row('F2LXK1ABC2', IMEI_B)])
        self.assertEqual(seen_marked['summary']['imei_on_unmarked_products'], 0)
        Product.objects.filter(pk=self.phone.pk).update(requires_imei=False)
        self.assertEqual(self.api.post(apply_url(seen_marked['id'])).status_code, 200)
        logged = AdminAuditLog.objects.get(action='stock_units_import_applied', target_id=seen_marked['id'])
        self.assertEqual(logged.metadata['imei_on_unmarked_products'], 2)

        seen_unmarked = self.staged([self.laptop_row('C02ZK0ABMD6R', IMEI_C)])
        self.assertEqual(seen_unmarked['summary']['imei_on_unmarked_products'], 1)
        Product.objects.filter(pk=self.laptop.pk).update(requires_imei=True)
        self.assertEqual(self.api.post(apply_url(seen_unmarked['id'])).status_code, 200)
        logged = AdminAuditLog.objects.get(action='stock_units_import_applied', target_id=seen_unmarked['id'])
        self.assertEqual(logged.metadata['imei_on_unmarked_products'], 0)


class PreviewTest(ImportBase):
    def test_a_preview_stages_every_row_and_registers_nothing(self):
        payload = self.staged([
            self.phone_row('F2LXK1ABC1', IMEI_A),
            self.phone_row('F2LXK1ABC2', IMEI_B, **{'IMEI 2': IMEI_C, 'Condición': 'Usado'}),
            ['MBA-13-M3', '', 'Sucursal Norte', 'C02ZK0ABMD6R', '', '', '', '', 'Compra F001-204'],
        ])

        self.assertEqual(payload['import_type'], 'units')
        self.assertEqual(payload['status'], 'previewed')
        self.assertEqual(payload['counts'], {**payload['counts'], 'total': 3, 'create': 3, 'error': 0})
        self.assertTrue(payload['is_applicable'])
        first = next(r for r in payload['rows'] if r['row'] == 2)['data']
        self.assertEqual(first['product_id'], self.phone.pk)
        self.assertEqual(first['branch_id'], self.branch_a.pk)
        self.assertEqual(first['serial_number'], 'F2LXK1ABC1')
        self.assertEqual(first['imei'], IMEI_A)
        self.assertEqual(next(r for r in payload['rows'] if r['row'] == 3)['data']['condition'], 'used')
        self.assertEqual(next(r for r in payload['rows'] if r['row'] == 4)['data']['branch_id'], self.branch_b.pk)
        # Ensayar no escribe.
        self.assertFalse(StockUnit.objects.exists())
        self.assertFalse(StockMovement.objects.filter(product__in=[self.phone, self.laptop]).exists())
        self.assertEqual(self.quantity(), 0)

    def test_an_identifier_excel_stored_as_a_number_is_read_whole(self):
        payload = self.staged([self.phone_row('F2LXK1ABC1', int(IMEI_A))])

        self.assertEqual(payload['counts']['error'], 0, payload['rows'])
        self.assertEqual(payload['rows'][0]['data']['imei'], IMEI_A)

    def test_each_bad_row_says_what_is_wrong_with_it(self):
        self.receive(self.row('YAESTABA01', IMEI_D))
        Product.objects.filter(pk=self.product.pk).update(name='Cargador sin serie')
        payload = self.staged([
            self.phone_row('', IMEI_A),                                          # 2 sin serie
            self.phone_row('F2LXK1ABC2', ''),                                    # 3 teléfono sin IMEI
            self.phone_row('F2LXK1ABC3', IMEI_A[:-1] + ('0' if IMEI_A[-1] != '0' else '1')),   # 4 control
            self.phone_row('F2LXK1ABC4', 'N/A'),                                 # 5 marcador
            self.phone_row('F2LXK1ABC5', IMEI_B, **{'Código': 'NO-EXISTE'}),     # 6 producto
            self.phone_row('F2LXK1ABC6', IMEI_B, **{'Código': '', 'Producto': 'Cargador sin serie'}),  # 7 sin serie
            self.phone_row('F2LXK1ABC7', IMEI_B, **{'Sucursal': 'Lima Sur'}),    # 8 sucursal
            self.phone_row('F2LXK1ABC8', IMEI_B, **{'Condición': 'Roto'}),       # 9 condición
            self.phone_row('F2LXK1ABC9', IMEI_B, **{'Costo': 'gratis'}),         # 10 costo
            self.phone_row('YAESTABA01', IMEI_B),                                # 11 serie ya registrada
            self.phone_row('F2LXK1ABD0', IMEI_D),                                # 12 IMEI ya registrado
            self.phone_row('F2LXK1ABD1', IMEI_B, **{'Motivo / referencia': ''}), # 13 sin motivo
            self.phone_row('F2LXK1ABD2', IMEI_C, **{'IMEI 2': IMEI_C}),          # 14 IMEI2 igual
        ])

        errors = self.errors(payload)
        self.assertEqual(sorted(errors), list(range(2, 15)), errors)
        self.assertIn('número de serie', errors[2])
        self.assertIn('lleva IMEI', errors[3])
        self.assertIn('control', errors[4])
        self.assertIn('IMEI', errors[5])
        self.assertIn('No existe un producto', errors[6])
        self.assertIn('no se controla por número de serie', errors[7])
        self.assertIn('Lima Sur', errors[8])
        self.assertIn('Condición', errors[9])
        self.assertIn('costo', errors[10].lower())
        self.assertIn('ya está registrado', errors[11])
        self.assertIn('ya está registrado', errors[12])
        self.assertIn('motivo', errors[13].lower())
        self.assertIn('distinto', errors[14])
        self.assertFalse(payload['is_applicable'])
        self.assertEqual(StockUnit.objects.count(), 1)

    def test_a_cost_that_is_not_a_number_is_a_row_error_not_a_crash(self):
        for cost in ('NaN', 'Infinity', '-Infinity', 'sNaN'):
            with self.subTest(cost):
                payload = self.staged([self.phone_row('F2LXK1ABC1', IMEI_A, **{'Costo': cost})])
                self.assertIn('costo', self.errors(payload)[2].lower())

    def test_the_same_device_twice_in_one_file_is_an_error_on_the_second_row(self):
        payload = self.staged([
            self.phone_row('F2LXK1ABC1', IMEI_A),
            self.phone_row('f2lxk1abc1', IMEI_B),                                # misma serie
            self.phone_row('F2LXK1ABC3', IMEI_C),
            self.phone_row('F2LXK1ABC4', IMEI_D, **{'IMEI 2': IMEI_C}),          # IMEI de la fila 4
        ])

        errors = self.errors(payload)
        self.assertEqual(sorted(errors), [3, 5], errors)
        self.assertIn('repetido', errors[3])
        self.assertIn('repetido', errors[5])

    def test_a_quantity_is_not_a_device(self):
        """Dos teléfonos son dos filas. Una columna «Cantidad» no los sustituye."""
        headers = HEADERS + ['Cantidad']
        response = self.api.post(PREVIEW_URL, {
            'file': workbook([self.phone_row('F2LXK1ABC1', IMEI_A) + [2]], headers=headers),
        }, format='multipart')

        self.assertEqual(response.status_code, 400)
        self.assertIn('Cantidad', response.json()['detail'])
        self.assertFalse(BulkImportJob.objects.exists())

    def test_a_default_branch_and_reason_fill_the_blanks(self):
        payload = self.staged(
            [self.phone_row('F2LXK1ABC1', IMEI_A, **{'Sucursal': '', 'Motivo / referencia': ''})],
            branch=self.branch_b.pk, reason='Inventario inicial de equipos',
        )

        self.assertEqual(payload['counts']['error'], 0, payload['rows'])
        data = payload['rows'][0]['data']
        self.assertEqual(data['branch_id'], self.branch_b.pk)
        self.assertEqual(data['reason'], 'Inventario inicial de equipos')

    def test_a_file_that_is_not_the_template_is_refused_whole(self):
        response = self.api.post(PREVIEW_URL, {
            'file': workbook([['x', 1]], headers=['Nombre', 'Stock']),
        }, format='multipart')
        self.assertEqual(response.status_code, 400)
        self.assertIn('Número de serie', response.json()['detail'])

        self.assertEqual(self.api.post(PREVIEW_URL, {}, format='multipart').status_code, 400)
        self.assertFalse(BulkImportJob.objects.exists())


class ApplyTest(ImportBase):
    def test_applying_registers_each_row_as_one_unit_through_the_kardex(self):
        job = self.staged([
            self.phone_row('F2LXK1ABC1', IMEI_A),
            self.phone_row('F2LXK1ABC2', IMEI_B, **{'IMEI 2': IMEI_C, 'Condición': 'Usado', 'Costo': ''}),
            ['MBA-13-M3', '', 'Sucursal Norte', 'C02ZK0ABMD6R', '', '', 'Reacondicionado', '4100', 'Compra F001-204'],
        ])

        response = self.api.post(apply_url(job['id']))

        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['status'], 'applied')
        self.assertEqual(response.json()['summary']['applied']['units'], 3)
        self.assertEqual(StockUnit.objects.count(), 3)
        self.assertEqual(self.quantity(), 2)
        self.assert_invariant()
        self.assertEqual(self.quantity(self.laptop, self.branch_b), 1)
        self.assert_invariant(self.laptop, self.branch_b)
        used = StockUnit.objects.get(serial_number='F2LXK1ABC2')
        self.assertEqual((used.imei, used.imei2, used.condition, used.cost), (IMEI_B, IMEI_C, 'used', None))
        self.assertEqual(used.status, 'available')
        self.assertEqual(used.company, self.company)
        self.assertEqual(str(StockUnit.objects.get(serial_number='C02ZK0ABMD6R').cost), '4100.00')
        movements = StockMovement.objects.filter(movement_type=StockMovement.PURCHASE_ENTRY,
                                                 product__in=[self.phone, self.laptop])
        self.assertEqual(sorted(m.quantity for m in movements), [1, 2])
        self.assertEqual({m.reason for m in movements}, {'Compra F001-204'})
        log = AdminAuditLog.objects.get(action='stock_units_import_applied')
        self.assertEqual(log.company, self.company)
        self.assertEqual(log.metadata['units'], 3)
        self.assertNotIn(IMEI_A, str(log.metadata))

    def test_applying_twice_registers_once(self):
        job = self.staged([self.phone_row('F2LXK1ABC1', IMEI_A)])

        self.assertEqual(self.api.post(apply_url(job['id'])).status_code, 200)
        self.assertEqual(self.api.post(apply_url(job['id'])).status_code, 200)

        self.assertEqual(StockUnit.objects.count(), 1)
        self.assertEqual(self.quantity(), 1)

    def test_a_preview_with_errors_cannot_be_applied(self):
        job = self.staged([self.phone_row('F2LXK1ABC1', IMEI_A), self.phone_row('', IMEI_B)])

        response = self.api.post(apply_url(job['id']))

        self.assertEqual(response.status_code, 400)
        self.assertFalse(StockUnit.objects.exists())

    def test_what_changed_since_the_preview_refuses_the_whole_file(self):
        """Todo o nada: si una fila ya no puede entrar, no entra ninguna."""
        job = self.staged([
            self.phone_row('F2LXK1ABC1', IMEI_A),
            self.phone_row('F2LXK1ABC2', IMEI_B),
            ['MBA-13-M3', '', self.branch_a.name, 'C02ZK0ABMD6R', '', '', '', '', 'Compra F001-204'],
        ])
        # Alguien registra a mano el segundo teléfono entre el ensayo y la carga.
        self.receive(self.row('OTRASERIE1', IMEI_B))

        response = self.api.post(apply_url(job['id']))

        self.assertEqual(response.status_code, 400)
        self.assertIn('fila 3', response.json()['detail'])
        self.assertEqual(StockUnit.objects.count(), 1)
        self.assertEqual(self.quantity(self.laptop), 0)
        self.assertEqual(BulkImportJob.objects.get(pk=job['id']).status, 'previewed')

    def test_a_product_that_stopped_being_serialized_refuses_the_file(self):
        job = self.staged([self.phone_row('F2LXK1ABC1', IMEI_A)])
        Product.objects.filter(pk=self.phone.pk).update(is_serialized=False)

        self.assertEqual(self.api.post(apply_url(job['id'])).status_code, 400)
        self.assertFalse(StockUnit.objects.exists())


class ScopeTest(ImportBase):
    def test_another_companys_product_and_branch_do_not_exist_here(self):
        ProductBarcode.objects.create(company=self.other, product=self.foreign_phone, code='AJENO-1')
        payload = self.staged([
            self.phone_row('F2LXK1ABC1', IMEI_A, **{'Código': 'AJENO-1'}),
            self.phone_row('F2LXK1ABC2', IMEI_B, **{'Sucursal': self.foreign_branch.name}),
        ])

        errors = self.errors(payload)
        self.assertIn('No existe un producto', errors[2])
        self.assertIn('No existe la sucursal', errors[3])

    def test_a_branch_the_caller_cannot_reach_is_an_error_at_preview_and_at_apply(self):
        job = self.staged([self.phone_row('F2LXK1ABC1', IMEI_A, **{'Sucursal': 'Sucursal Norte'})])
        self.restrict_to_branch_a()

        again = self.preview([self.phone_row('F2LXK1ABC1', IMEI_A, **{'Sucursal': 'Sucursal Norte'})])
        self.assertIn('No tienes acceso', self.errors(again.json())[2])
        # El ensayo hecho ANTES de perder el acceso tampoco se puede cargar ni leer.
        self.assertIn(self.api.post(apply_url(job['id'])).status_code, (403, 404))
        self.assertEqual(self.api.get(f'/api/admin/imports/{job["id"]}/').status_code, 404)
        self.assertFalse(StockUnit.objects.exists())

    def test_a_job_whose_only_rows_for_a_branch_are_errors_still_belongs_to_that_branch(self):
        """Una fila con error también lleva la serie y el IMEI: cuenta para saber quién puede leerla."""
        self.receive(self.row('YAESTABA01', IMEI_D), branch=self.branch_b)
        job = self.staged([
            self.phone_row('F2LXK1ABC1', IMEI_A),                                   # Centro, entra
            self.phone_row('F2LXK1ABC2', IMEI_D, **{'Sucursal': 'Sucursal Norte'}), # Norte, error
        ])
        self.assertEqual(self.errors(job).keys(), {3})
        self.restrict_to_branch_a()

        self.assertEqual(self.api.get(f'/api/admin/imports/{job["id"]}/').status_code, 404)
        self.assertEqual(self.api.get(f'/api/admin/imports/{job["id"]}/errors.csv/').status_code, 404)
        self.assertEqual(self.api.get('/api/admin/imports/?type=units').json()['results'], [])
        self.assertEqual(self.api.post(apply_url(job['id'])).status_code, 404)

    def test_it_needs_the_capability_to_adjust_inventory(self):
        viewer = _m7_user('solo-mira')
        membership = Membership.objects.create(user=viewer, company=self.company, role='inventory')
        from store.tests import _assign, _role
        _assign(membership, _role(self.company, 'Sólo ver', capabilities=['inventory.view'], slug='solo-ver'))
        client = APIClient()
        client.force_authenticate(user=viewer)

        self.assertEqual(client.get(TEMPLATE_URL).status_code, 403)
        self.assertEqual(self.preview([self.phone_row('F2LXK1ABC1', IMEI_A)], client=client).status_code, 403)
        self.assertEqual(APIClient().post(PREVIEW_URL, {}, format='multipart').status_code, 401)

    def test_a_job_of_another_company_is_not_found(self):
        job = self.staged([self.phone_row('F2LXK1ABC1', IMEI_A)])
        outsider = _m7_user('ajena')
        Membership.objects.create(user=outsider, company=self.other, role='admin')
        client = APIClient()
        client.force_authenticate(user=outsider)

        self.assertEqual(client.post(apply_url(job['id'])).status_code, 404)
        self.assertFalse(StockUnit.objects.exists())

    def test_the_history_lists_the_job_as_equipment(self):
        job = self.staged([self.phone_row('F2LXK1ABC1', IMEI_A)])

        listed = self.api.get('/api/admin/imports/?type=units').json()['results']

        self.assertEqual([j['id'] for j in listed], [job['id']])
        self.assertEqual(BulkImportRow.objects.filter(job_id=job['id']).count(), 1)
