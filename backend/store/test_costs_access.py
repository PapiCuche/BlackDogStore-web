"""
DEVICE-COMMERCE-01 · DC1 — el costo sólo para quien tiene `costs.view`.

EL HALLAZGO QUE ESTO CIERRA. `GET /api/admin/inventory/units/` pedía sólo
`inventory.view` y devolvía el `cost` de cada equipo, mientras la descarga lo
escondía salvo con `inventory.adjust`. Dos superficies, dos reglas, y la más
laxa era la de la tabla que cualquiera con acceso al inventario abre.

Ahora hay UNA regla: el costo viaja sólo hacia quien tiene `costs.view`. Sin
ella, la clave `cost` no está en la respuesta (no un `null`, que se lee como
«sin costo»). Y cambiar el costo de un equipo existente también la exige:
corregir una valorización a ciegas no es mover inventario.

El IMEI completo NO cambia: es la regla del inventario para el personal
(`stock_unit_views`, la descarga y el punto de venta lo muestran entero a
quien opera el stock); sólo se enmascara hacia el cliente.
"""
import io

from django.core.cache import cache
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from store.models import Membership, StockUnit
from store.test_stock_units import IMEI_A, IMEI_B, IMEI_D, UnitsBase
from store import stock_unit_services as units
from store.test_device_commerce import assert_no_amount
from store.tests import _assign, _role

LIST = '/api/admin/inventory/units/'
INVENTORY = ['company.view', 'inventory.view', 'inventory.adjust']


class CostOnTheUnitListTest(UnitsBase):
    def setUp(self):
        super().setUp()
        self.first, = self.receive(self.row('SERIE00001', IMEI_A, cost='3800.00'))
        units.receive_units(
            branch=self.foreign_branch, product=self.foreign_phone,
            units=[self.row('SERIEAJENA1', IMEI_D, cost='999.00')], actor=self.staff, reason='Compra')
        self.api = APIClient()
        self.api.force_authenticate(user=self.staff)

    def grant(self, *capabilities):
        _assign(self.membership, _role(self.company, 'Rol costos', capabilities=list(capabilities),
                                       slug='rol-costos'))

    def rows(self, query='?branch=all'):
        res = self.api.get(LIST + query)
        self.assertEqual(res.status_code, 200, res.content)
        return res.json()['results']

    def test_inventory_alone_does_not_see_the_cost_and_the_key_is_absent(self):
        # Legacy `inventory` role: inventory.view + inventory.adjust, no costs.view.
        [row] = self.rows()
        self.assertNotIn('cost', row)
        assert_no_amount(self, self.api.get(LIST + '?branch=all').json(), '3800')
        # The identifiers stay whole for staff who operate the shelf.
        self.assertEqual(row['imei'], IMEI_A)

    def test_costs_view_sees_the_cost(self):
        self.grant(*INVENTORY, 'costs.view')
        [row] = self.rows()
        self.assertEqual(row['cost'], '3800.00')

    def test_a_legacy_admin_sees_the_cost(self):
        Membership.objects.filter(pk=self.membership.pk).update(role='admin')
        [row] = self.rows()
        self.assertEqual(row['cost'], '3800.00')

    def test_costs_view_without_inventory_view_does_not_open_the_list(self):
        self.grant('company.view', 'costs.view')
        self.assertEqual(self.api.get(LIST).status_code, 403)

    def test_costs_view_in_another_company_shows_nothing_of_it_here(self):
        # Holding costs.view somewhere else is not holding it here, and the
        # other tenant's device is not in the list in any case.
        stranger = Membership.objects.create(user=self.staff, company=self.other, role='admin')
        self.addCleanup(stranger.delete)
        res = self.api.get(LIST + f'?branch=all&company={self.company.pk}')
        self.assertEqual(res.status_code, 200, res.content)
        [row] = res.json()['results']
        self.assertNotIn('cost', row)
        self.assertNotIn('SERIEAJENA1', res.content.decode())

    def test_receiving_answers_without_the_cost_for_who_cannot_see_it(self):
        res = self.api.post(LIST, {
            'product_id': self.phone.pk, 'branch': self.branch_a.pk, 'reason': 'Compra',
            'units': [self.row('SERIE00002', IMEI_B, cost='3700.00')],
        }, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        self.assertNotIn('cost', res.json()['results'][0])
        # Recording what the invoice says is part of receiving: it is stored.
        self.assertEqual(str(StockUnit.objects.get(serial_number='SERIE00002').cost), '3700.00')

    def test_editing_the_cost_needs_costs_view_and_the_answer_hides_it_otherwise(self):
        url = f'{LIST}{self.first.pk}/'
        res = self.api.patch(url, {'condition': 'open_box'}, format='json')
        self.assertEqual(res.status_code, 200, res.content)
        self.assertNotIn('cost', res.json())

        refused = self.api.patch(url, {'cost': '1.00'}, format='json')
        self.assertEqual(refused.status_code, 403, refused.content)
        self.first.refresh_from_db()
        self.assertEqual(str(self.first.cost), '3800.00')

        reserved = self.api.post(url + 'reserve/', {'reason': 'Apartado'}, format='json')
        self.assertNotIn('cost', reserved.json())

        self.grant(*INVENTORY, 'costs.view')
        # A1 — a correction from the panel says why (`test_costs_ux`).
        allowed = self.api.patch(url, {'cost': '3750.00', 'cost_reason': 'Factura'}, format='json')
        self.assertEqual(allowed.status_code, 200, allowed.content)
        self.assertEqual(allowed.json()['cost'], '3750.00')


class CostInTheExportsTest(UnitsBase):
    """The download follows the same rule: `costs.view`, not `inventory.adjust`."""

    def setUp(self):
        super().setUp()
        self.receive(self.row('SERIE00001', IMEI_A, cost='3800.00'))
        self.api = APIClient()
        self.api.force_authenticate(user=self.staff)

    def costs(self, client=None):
        cache.clear()
        res = (client or self.api).get('/api/admin/inventory/exports/units/?mode=report')
        self.assertEqual(res.status_code, 200, res.content)
        sheet = load_workbook(io.BytesIO(res.content)).active
        rows = list(sheet.iter_rows(values_only=True))
        column = rows[0].index('Costo')
        return {row[column] for row in rows[1:]}

    def test_inventory_adjust_alone_gets_no_cost(self):
        self.assertEqual(self.costs(), {None})

    def test_costs_view_gets_the_cost_even_without_inventory_adjust(self):
        _assign(self.membership, _role(self.company, 'Lectura con costos', slug='lectura-costos',
                                       capabilities=['company.view', 'inventory.view', 'costs.view']))
        self.assertEqual(self.costs(), {3800})


class DjangoAdminCannotEditCostsTest(TestCase):
    """DC1 review — the frozen snapshots and the averages are written by the Kardex only."""

    SNAPSHOT_FIELDS = {'unit_cost', 'cost_source', 'unit_price_net', 'discount_net', 'net_amount'}

    def request(self):
        from django.contrib.auth import get_user_model
        from django.test import RequestFactory

        request = RequestFactory().get('/admin/')
        request.user = get_user_model().objects.create_superuser('dc1_admin_root', 'r@example.com', 'Pass123!')
        return request

    def test_sale_line_snapshots_are_read_only(self):
        from django.contrib import admin as dj_admin
        from store.admin import OrderItemInline
        from store.models import Order

        inline = OrderItemInline(Order, dj_admin.site)
        request = self.request()
        self.assertLessEqual(self.SNAPSHOT_FIELDS, set(inline.get_readonly_fields(request)))
        form_fields = set(inline.get_formset(request).form.base_fields)
        self.assertEqual(form_fields & self.SNAPSHOT_FIELDS, set())

    def test_branch_averages_are_read_only(self):
        from django.contrib import admin as dj_admin
        from store.models import BranchStock

        model_admin = dj_admin.site._registry[BranchStock]
        request = self.request()
        self.assertIn('average_cost', model_admin.get_readonly_fields(request))
        self.assertIn('uncosted_quantity', model_admin.get_readonly_fields(request))
        form_fields = set(model_admin.get_form(request).base_fields)
        self.assertNotIn('average_cost', form_fields)
        self.assertNotIn('uncosted_quantity', form_fields)
